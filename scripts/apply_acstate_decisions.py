#!/usr/bin/env python3
"""apply_acstate_decisions.py — write the flight states picked on acstate_review.html into the wizard record.
    python scripts/apply_acstate_decisions.py [ACSTATE_DECISIONS.csv]            # dry run
    python scripts/apply_acstate_decisions.py [ACSTATE_DECISIONS.csv] --apply    # back up, write, re-split the batch files
Default decisions file: 0_labelling/inputs/review_decisions/ACSTATE_DECISIONS.csv (the page's export).
Rule (codebook R2-02 (a), 2026-09-19): a part drawn in two positions → the state drawn in solid lines; both drawn with the
same weight → Both (new state 2026-09-19); nothing readable → Other; lift units beside cruise units are not two states; Invariant never for a figure showing two positions.
Only rows whose new value differs from the one the page showed are written, and only if the record still holds that old
value (a figure re-labelled in the wizard since the page was built is listed, never overwritten). Values are written as
the wizard writes them ("id — Label"); a note typed on the page goes to the figure's stateNote only when it is empty.
After --apply: notebook 04 → canonicalize_record.py --apply → 04; the user reloads the record in the wizard.
"""
import sys, shutil, subprocess
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook, Workbook

ROOT = Path("/mnt/storage_11tb/Drive_files_to_syncronize/3 - Images DataSets & Labelling Outputs/1639_LABELLED/0_labelling")
RECORD = ROOT / "inputs" / "record" / "reviewed_patents_Batch_ALL.xlsx"
args = [a for a in sys.argv[1:] if not a.startswith("--")]
DEC = Path(args[0]) if args else ROOT / "inputs" / "review_decisions" / "ACSTATE_DECISIONS.csv"
APPLY = "--apply" in sys.argv[1:]
LABEL = {"Hover": "Hover — Hover (VTOL)", "Transition": "Transition — Transition", "Cruise": "Cruise — Cruise / Forward",
         "HoverCruise": "HoverCruise — Invariant", "Both": "Both — Both (two positions drawn)", "Other": "Other — Other"}
def s(v): return "" if v is None else str(v)
def sid(v): return s(v).split(" — ")[0].strip()

D = pd.read_csv(DEC, dtype=str, keep_default_na=False)
D = D[D.new != D.old]
assert D.new.isin(LABEL).all(), D[~D.new.isin(LABEL)]
rows = [list(r) for r in load_workbook(RECORD)["Review"].iter_rows(values_only=True)]
hdr, body = rows[0], rows[1:]
ix = {c: i for i, c in enumerate(hdr)}
base = lambda p: s(p).split("_arch")[0]
where = {}
for i, r in enumerate(body):
    if s(r[ix["Section"]]) == "T2": where[(base(r[ix["Patent_ID"]]), s(r[ix["Sub_Dimension"]]), s(r[ix["Field"]]))] = i
done, skipped, notes = [], [], []
for d in D.itertuples():
    k = (d.patent_id, f"Image: {d.image_file}", "acState"); i = where.get(k)
    if i is None: skipped.append((d, "no acState row for this figure in the record")); continue
    cur = sid(body[i][ix["Value"]]); cur = "HoverCruise" if cur == "Invariant" else cur
    if cur != d.old: skipped.append((d, f"the record now holds {cur!r} (changed in the wizard since the page was built)")); continue
    body[i][ix["Value"]] = LABEL[d.new]; done.append(d)
    j = where.get((d.patent_id, f"Image: {d.image_file}", "stateNote"))
    if d.note.strip() and j is not None and not s(body[j][ix["Value"]]).strip():
        body[j][ix["Value"]] = d.note.strip(); notes.append(d)
print(f"{DEC.name}: {len(D)} changed answers → {len(done)} written, {len(skipped)} skipped, {len(notes)} notes added")
for d in done: print(f"  {d.aircraft_id:<22} {d.image_file:<46} {d.old} -> {d.new}")
for d, why in skipped: print(f"  (skipped) {d.aircraft_id} {d.image_file}: {why}")
if not APPLY: print("dry run — re-run with --apply"); sys.exit(0)
if not done: print("nothing to write"); sys.exit(0)
bak = RECORD.parent / "_backups" / f"{RECORD.stem}.PRE_ACSTATE_{datetime.now():%Y%m%d_%H%M%S}.xlsx"; shutil.copy2(RECORD, bak)
wb = Workbook(); ws = wb.active; ws.title = "Review"
for r in [hdr] + body: ws.append(r)
wb.save(RECORD); print(f"written {RECORD}; backup {bak}")
subprocess.run([sys.executable, str(Path(__file__).with_name("split_wizard_all_export.py")), str(RECORD), "--apply"], check=True)
