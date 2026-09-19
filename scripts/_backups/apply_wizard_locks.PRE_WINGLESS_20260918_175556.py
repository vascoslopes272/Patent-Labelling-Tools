#!/usr/bin/env python3
"""apply_wizard_locks.py — write into the wizard record the value the wizard SHOWS wherever a lock hides a
different stored value (src/wizard_view.py replays the locks); then regenerate the five per-batch copies.
    python scripts/apply_wizard_locks.py            # dry run: list every change
    python scripts/apply_wizard_locks.py --apply    # back up, write the record, re-split the batch files
Notebook 04 flags WIZARD_VIEW_MISMATCH and fails its "record equals the wizard view" check until this has run.

Why the record held them (2026-09-18): the wizard drew a locked, disabled chip (Invariant on a fixed architecture's
figures, Whole-Body Pitch on a PTC/RC body) but exported whatever had been stored before the lock — e.g. a Hover picked
before TB and PTC joined the fixed set on 2026-08-30. Wizard v15.10 exports the locked value, so a new export cannot
bring them back; this script repairs the record already written. A lock whose screen shows no value (a TB body stored
as Fixed: the wizard clears it and asks) is listed, never guessed.
Cells keep their types (openpyxl) — the wizard's xlBool() reads the text "False" as true.
"""
import sys, shutil, subprocess, collections
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook, Workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.wizard_view import replay, strip_label

ROOT = Path("/mnt/storage_11tb/Drive_files_to_syncronize/3 - Images DataSets & Labelling Outputs/1639_LABELLED/0_labelling")
RECORD = ROOT / "inputs" / "record" / "reviewed_patents_Batch_ALL.xlsx"
APPLY = "--apply" in sys.argv[1:]
def s(v): return "" if v is None else str(v)

rows = [list(r) for r in load_workbook(RECORD)["Review"].iter_rows(values_only=True)]
hdr, body = rows[0], rows[1:]
ix = {c: i for i, c in enumerate(hdr)}
L = pd.DataFrame([[s(r[ix[c]]) for c in ("Patent_ID", "Section", "Sub_Dimension", "Field", "Value")] for r in body],
                 columns=["Patent_ID", "Section", "Sub_Dimension", "Field", "Value"])
D = replay(L)

# the "id — Label" composite the wizard writes for an id, taken from the record itself
composite = collections.defaultdict(collections.Counter)
for f, v in zip(L.Field, L.Value):
    if " — " in v: composite[(f, strip_label(v))][v] += 1
def label(field, vid): return composite[(field, vid)].most_common(1)[0][0] if composite[(field, vid)] else vid

where = {}                                   # (Patent_ID, Sub_Dimension, Field) -> body index
for i, r in enumerate(body): where[(s(r[ix["Patent_ID"]]), s(r[ix["Sub_Dimension"]]), s(r[ix["Field"]]))] = i
ids_with_arch = set(L.Patent_ID[L.Patent_ID.str.contains(r"_arch\d+$", regex=True)])

changed, inserted, human = [], [], []
insert_after = {}                            # body index -> rows to insert after it
for d in D.itertuples():
    if not d.wizard_shows:
        human.append(d); continue
    new = label(d.field, d.wizard_shows)
    if d.scope == "figure":
        key = (d.patent_id, d.figure, "acState")
    else:
        apid = f"{d.patent_id}_arch{d.arch}" if f"{d.patent_id}_arch{d.arch}" in ids_with_arch else d.patent_id
        key = (apid, next((k[1] for k in where if k[0] == apid and k[2] == d.field), d.field), d.field)
    i = where.get(key)
    if i is not None:
        changed.append((key, s(body[i][ix["Value"]]), new)); body[i][ix["Value"]] = new
    else:                                    # no row at all: add it after the last row of the same block
        anchor = max(j for k, j in where.items() if k[0] == key[0] and k[1] == key[1])
        row = [None] * len(hdr)
        for c, v in (("Patent_ID", key[0]), ("Section", "T2" if d.scope == "figure" else "M1"), ("Sub_Dimension", key[1]),
                     ("Field", key[2]), ("Value", new), ("Source", "human"), ("Image_Path", body[anchor][ix["Image_Path"]])):
            if c in ix: row[ix[c]] = v
        insert_after.setdefault(anchor, []).append(row); inserted.append((key, new))

print(f"record {RECORD.name}: {len(body)} rows; wizard-view differences {len(D)} -> {len(changed)} values to write, "
      f"{len(inserted)} rows to add, {len(human)} need a human pick")
for (pid, sub, f), old, new in changed: print(f"  {pid:<24} {f:<8} {old!r} -> {new!r}   {sub if f == 'acState' else ''}")
for (pid, sub, f), new in inserted: print(f"  + {pid:<22} {f:<8} {new!r}   {sub}")
for d in human: print(f"  ?? {d.patent_id} arch{d.arch} {d.field}={d.stored!r}: {d.rule}")
if not APPLY: print("dry run — re-run with --apply"); sys.exit(0)
if not (changed or inserted): print("nothing to write"); sys.exit(0)

out = [hdr]
for i, r in enumerate(body):
    out.append(r); out += insert_after.get(i, [])
ts = datetime.now().strftime("%Y%m%d_%H%M%S")
bdir = RECORD.parent / "_backups"; bdir.mkdir(exist_ok=True)
bak = bdir / f"{RECORD.stem}.PRE_LOCKS_{ts}.xlsx"; shutil.copy2(RECORD, bak)
wb = Workbook(); ws = wb.active; ws.title = "Review"
for r in out: ws.append(r)
wb.save(RECORD); print(f"written {RECORD} ({len(out) - 1} rows); backup {bak}")
subprocess.run([sys.executable, str(Path(__file__).with_name("split_wizard_all_export.py")), str(RECORD), "--apply"], check=True)
