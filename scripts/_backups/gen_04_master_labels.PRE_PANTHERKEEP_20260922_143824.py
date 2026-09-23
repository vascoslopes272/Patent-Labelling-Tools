"""Generate notebooks/04_master_labels.ipynb from cell sources (kept here so the
notebook can be regenerated deterministically instead of hand-patched JSON).

2026-09-17 rewrite (user rulings of 2026-09-16/17):
  * input = the ONE wizard record  0_labelling/inputs/record/reviewed_patents_Batch_ALL.xlsx
  * aircraft id = <patent>_ua<N> everywhere (a single-aircraft patent is still _ua1)
  * D1/D2 duplicates carry NOTHING from G1 to M3 — they POINT at the aircraft they repeat
    (`same_aircraft_as`); a D2 has no figures of its own either (T2 points as well)
  * names: the reviewed real name (+ vN) where one exists, otherwise the aircraft id;
    a D1/D2 carries the name of the aircraft it points at
  * outputs = 0_labelling/outputs/: ONE workbook 1639_LABELS.xlsx (README · Review = the record,
    wizard-loadable, names written in · ground_truth · patseer), images/<aircraft_id>/,
    tables/*.csv for the analysis stages, BUILD_LOG.md, ACTION_FLAGS.csv
"""
import nbformat as nbf
from pathlib import Path

cells = []
def md(s): cells.append(nbf.v4.new_markdown_cell(s.strip("\n")))
def code(s): cells.append(nbf.v4.new_code_cell(s.strip("\n")))

md(r"""
# 04 — Master labels (Stage 0 output: one workbook, one row per aircraft)

**Stage 04 = the join.** Reads the ONE wizard record (`0_labelling/inputs/record/reviewed_patents_Batch_ALL.xlsx`),
the ML feeds, the identity workbook (built by `scripts/build_identity_all.py` from the 03a sheets, the corrections and
`NAME_DECISIONS.csv`), the architecture ground truth and the PatSeer export; writes `0_labelling/outputs/` — disposable,
rebuilt on every run. Nothing here writes back to an input.

| what | rule |
|---|---|
| grain | one row per **aircraft** = `aircraft_id = <patent>_ua<N>` (`ua` = the wizard's aircraft number; a single-aircraft patent is `_ua1`); rejected patents are one row with empty labels |
| values | bare **ids** (the `"id — Label"` composite is stripped); booleans as True/False; counts as integers (the wizard stores 1 as `True`) |
| order | **one block per company, alphabetical** (ruling 2026-09-17): companies A→Z (`reference/family_map.csv` `company_canonical` → column `company`; the two pseudo-companies last); inside a company named aircraft A→Z, then unnamed by patent id, then rejected; a D1/D2 of the same company right under its original; a patent's aircraft together in `ua` order — Review sheet, ground_truth, patseer, both tables, MANIFEST and the record itself (`scripts/canonicalize_record.py --apply` copies the table order into the record + feed after every export / 04 run; 04 flags RECORD_ORDER_STALE / RECORD_NAMES_STALE) |
| duplicates | **D1/D2 carry nothing from G1 to M3**: their labels are blank and `same_aircraft_as` points at the aircraft they repeat (the pinned one from `review_decisions/duplicate_root_variant.csv`, else every aircraft of the original); a D2 has no figures of its own (`is_primary=False`). D3 keep their own labels and figures. Every duplicate is re-ordered directly under its original |
| names | reviewed real name with its `vN` where one exists (`NAME_DECISIONS.csv` → identity workbook); otherwise the aircraft id; a D1/D2 takes the name of the aircraft it points at. The typed wizard name stays in `aircraft_name_wizard` |
| ground truth | 7d: `arch_gt`, `arch_gt_provenance`, `arch_gt_visible` from `text_architecture/architecture_text_final_variants.csv` on primary approved aircraft |
| identity (03a) | 7b: `uav_final` / `is_electric_final` / `takeoff_final` + edge tags UAVSimilar / ElectricSimilar / STOLSimilar merged into `edgeTags` on primary rows |
| ML pre-labels | 7c: the feed's predictions with an `ml_` prefix (tables only, never in the workbook) |
| names | the wizard's names instead of four legacy ids (user 2026-09-18): gearArch `Wheeled Gear` (id FixedWheel), fusKin `Variable Incidence` (id LevelCabin), every `*_rmech` `Fixed` (id Exposed), figure qualityFlag `Partial Quality` (id generic); `dup_type` stays 1/2/3 (the analysis computes with it) with `dup_type_name` D1/D2/D3 beside it. TiltBody and `No Aircraft Image` are kept as they are |
| wizard view | 6b: every lock the wizard draws is replayed on the record (`src/wizard_view.py`); a stored value the screen does not show FAILS the build — fix with `scripts/apply_wizard_locks.py --apply`. `acState` is written as the wizard's display name **Invariant**, not its stored id `HoverCruise` |
| rules 2026-09-19 | 7e: `overrides`, `propulsor_units` (+ `_note`), `boom_thrust_state`, `wing_thrust_carrier`, `wing_borne_units` (+ `_note`), `wing_boom_candidate` — derived beside the labels, no label changed; a value an override hides is never read as zero / none / Fixed (blank + reason). Flags OVERRIDE_HIDES_VALUE, LEFTOVER_QUICK_COUNT, BOOM_TICK_AND_ROTOR_TILT, BOOM_ROTORS_WITHOUT_GROUP, STATE_NOTE_MENTIONS_TWO_STATES (R2-02(e)) |
| re-codes | two `T2.parts` free texts → `Other` (the old US2021284333A1 Oth→Trap recode was a wizard reload bug, fixed in the record 2026-09-18) |
| dropped | `codebook_version`, `timestamp`, `familyId`, `labelToken`, SBERT pre-labels, long bibliographic text, retired fields |
| outputs | `1639_LABELS.xlsx` (README · **Review** = the record with the names written in, loadable in the wizard · ground_truth · patseer) · `images/<aircraft_id>/<figure>.png` + MANIFEST.csv · `tables/aircraft_table.csv`, `figure_table.csv`, `data_dictionary.csv` (what stages 1–2 read) · `BUILD_LOG.md` · `ACTION_FLAGS.csv` |
""")

code(r"""
import os, re, sys, json, shutil, hashlib, subprocess, tempfile, collections
from pathlib import Path
from datetime import datetime
import numpy as np
import pandas as pd

repo_root = Path().resolve().parent
if str(repo_root) not in sys.path:
    sys.path.insert(0, str(repo_root))
from src.config_loader import load_config
cfg = load_config(); P = cfg["paths"]

# ── switches (env overrides so the notebook can run unattended) ─────────────
REASON_SOURCE = os.environ.get("NB04_REASON_SOURCE", "human")      # human | sbert | sbert_when_human_empty
SBERT_XLSX    = os.environ.get("NB04_SBERT_XLSX") or None           # the SBERT/LLM disapproval-reasons workbook (may not exist yet)
COPY_IMAGES   = os.environ.get("NB04_COPY_IMAGES", "1") == "1"      # real copies of every approved figure
BATCHES       = ["Batch_01", "Batch_02", "Batch_03", "Batch_04", "Batch_05"]

IN_ROOT     = Path(P["labelled"])                                   # 0_labelling/inputs
REC_DIR     = IN_ROOT / "record"                                    # THE wizard record (+ the ALL feed, the batch map)
RECORD      = REC_DIR / "reviewed_patents_Batch_ALL.xlsx"
LAB_DIR     = Path(P["corrected_wizard_exports"])                   # inputs/labels — the five generated copies + per-batch feeds
RD          = IN_ROOT / "review_decisions"
TA          = IN_ROOT / "text_architecture"
ID_XLSX     = IN_ROOT / "identity" / "aircraft_identity_ALL.xlsx"  # rebuilt by scripts/build_identity_all.py
PATSEER     = IN_ROOT / "reference" / "patseer_1639_08_06_26.xlsx"
HTML        = Path(P["html_template"])
OUT_DIR     = Path(os.environ.get("NB04_OUT_DIR") or P["labelled_outputs"])   # NB04_OUT_DIR: test builds elsewhere
TAB_DIR     = OUT_DIR / "tables"
WORKBOOK    = OUT_DIR / "1639_LABELS.xlsx"
TAGS_CSV    = Path("04_other_note_tags.csv")                        # beside this notebook; edit freely
SCHEMA_JS   = repo_root / "scripts" / "conformance" / "extract_html_schema.js"
assert REASON_SOURCE in ("human", "sbert", "sbert_when_human_empty"), REASON_SOURCE
RUN_TS = datetime.now().strftime("%Y-%m-%d %H:%M")
LOG: list[str] = [f"# 04_master_labels build log — {RUN_TS}", ""]
def log(s=""): print(s); LOG.append(str(s))
def sha(p, n=12): return hashlib.sha256(Path(p).read_bytes()).hexdigest()[:n]
def stamp(p): p = Path(p); return f"{p.name}  {datetime.fromtimestamp(p.stat().st_mtime):%Y-%m-%d %H:%M}  sha256:{sha(p)}"
INPUTS_USED = []
def used(p, what): INPUTS_USED.append((what, str(p), stamp(p))); log(f"  input  {what:<22} {stamp(p)}")
log(f"record : {RECORD}"); log(f"output : {OUT_DIR}"); log(f"wizard : {HTML.name}"); log(f"REASON_SOURCE={REASON_SOURCE}  SBERT_XLSX={SBERT_XLSX}  COPY_IMAGES={COPY_IMAGES}")
""")

md("## 1 — Load the record (long form); verify the five generated copies still equal it")
code(r"""
SEP = " — "
def strip_label(v):
    s = str(v); return s.split(SEP, 1)[0].strip() if SEP in s else s.strip()
def is_empty(v):
    return v is None or (isinstance(v, float) and np.isnan(v)) or str(v).strip() in ("", "nan", "None", "NaT")
def base_id(p): return re.sub(r"_arch\d+$", "", str(p))
def arch_idx(p):
    m = re.search(r"_arch(\d+)$", str(p)); return int(m.group(1)) if m else 1

used(RECORD, "wizard record")
L = pd.read_excel(RECORD, sheet_name="Review")
L["Field"] = L.Field.astype(str); L["Section"] = L.Section.astype(str)
L["base"] = L.Patent_ID.map(base_id); L["arch"] = L.Patent_ID.map(arch_idx)
bmap = pd.read_csv(REC_DIR / "patent_batch_map.csv", dtype=str).set_index("patent_id").batch.to_dict()
L["batch"] = L.base.map(bmap)
unmapped_b = sorted(L[L.batch.isna()].base.unique())
assert not unmapped_b, f"patents missing from patent_batch_map.csv: {unmapped_b[:8]}"
L["_pos"] = range(len(L))                            # record order (only checked against the corpus order below)
FAMILY_MAP = IN_ROOT / "reference" / "family_map.csv"; used(FAMILY_MAP, "company map")
_fam = pd.read_csv(FAMILY_MAP, dtype=str, keep_default_na=False)
comp_of = dict(zip(_fam.canonical_pub_number, _fam.company_canonical))
_nocomp = sorted(set(L.base) - set(comp_of)); assert not _nocomp, f"patents missing from family_map.csv: {_nocomp[:8]}"
L["company"] = L.base.map(comp_of)
log(f"  companies: {L.drop_duplicates('base').company.nunique()} normalised names (reference/family_map.csv company_canonical); row order = §8")
log(f"long table: {len(L)} rows, {L.base.nunique()} patents, {L.Patent_ID.nunique()} ids; per batch {L.drop_duplicates('base').batch.value_counts().sort_index().to_dict()}")

# The five files in inputs/labels are GENERATED from the record (scripts/split_wizard_all_export.py). Say whether they
# still equal it — 02a/03a and ~25 scripts read them, so a stale copy would make those stages disagree with this one.
def _rowset(df): return sorted(map(tuple, df[["Patent_ID", "Section", "Sub_Dimension", "Field", "Value", "Source", "Image_Path"]].fillna("").astype(str).values.tolist()))
SPLIT_STALE = []
for b in BATCHES:
    s = pd.read_excel(LAB_DIR / f"reviewed_patents_{b}.xlsx", sheet_name="Review")
    same = _rowset(s) == _rowset(L[L.batch == b])
    log(f"  {b}: generated copy equals the record -> {same}")
    if not same: SPLIT_STALE.append(b)
if SPLIT_STALE: log(f"  !! STALE generated copies {SPLIT_STALE}: run scripts/split_wizard_all_export.py <record> --apply")
""")

md("## 2 — Wizard schema (option ids + display names) for the data dictionary and value checks")
code(r"""
schema_out = Path(tempfile.mkstemp(suffix=".json")[1])
r = subprocess.run(["node", str(SCHEMA_JS), str(HTML), str(schema_out)], capture_output=True, text=True)
assert r.returncode == 0, r.stderr
SCHEMA = json.loads(schema_out.read_text())
LISTS  = {k: list(v) for k, v in SCHEMA["lists"].items()}
LABELS = SCHEMA.get("labels", {})
log(f"schema: {len(LISTS)} option lists, codebook_version {SCHEMA['codebook_version']!r} (not carried)")

# suffix -> option list (mirrors scripts/conformance/check_batches.py FIELD_LIST + FIELD_LIST_EXTRA)
CODED = {("G1","topType"):"TOP", ("G1","edgeTags"):"T1_EDGE_TAGS",
         ("M1","attach"):"BOOM_ATTACH", ("M1","long"):"BOOM_POS_LONG", ("M1","span"):"BOOM_POS_SPAN",
         ("M1","wingRel"):"BOOM_WING_REL", ("M1","wingIdx"):"BOOM_WING_IDX", ("M1","orient"):"BOOM_ORIENT",
         ("M1","fusShape"):"FUS_SHAPE", ("M1","fusKin"):"FUS_KIN", ("M1","gearArch"):"GEAR_ARCH",
         ("M1","dinoUnderstanding"):"DINO_UNDERSTANDING",
         ("M2","wingConf"):"WING_CONFIG", ("M2","posV"):"W_POS_V", ("M2","posL"):"W_POS_L", ("M2","plan"):"W_PLAN",
         ("M2","role"):"W_ROLE", ("M2","tilt"):"W_TILT", ("M2","empType"):"EMP_TYPE",
         ("M3","chord"):"CHORD", ("M3","bmech"):"BLADE_MECH", ("M3","rmech"):"RETRACT_MECH", ("M3","propKin"):"PROP_KIN",
         ("M3","orient"):"M3_ORIENT", ("M3","zone"):"M3_ZONE", ("M3","zoneChord"):"M3_ZONE_CHORD", ("M3","zoneSpan"):"M3_ZONE_SPAN",
         ("T2","per"):"T2_PER", ("T2","acSty"):"T2_AC_STY", ("T2","acCol"):"T2_AC_COL", ("T2","bgSty"):"T2_BG_STY",
         ("T2","bgCol"):"T2_BG_COL", ("T2","acState"):"AC_STATE", ("T2","qualityFlag"):"QUALITY_FLAGS", ("T2","parts"):"T2_PARTS_DEFAULT",
         ("META","duplicateType"):"DUP_TYPES", ("T1","t1DisapproveReason"):"T1_DISAPPROVE_REASONS"}
MULTI = {("M3","zone"), ("G1","edgeTags"), ("T2","parts")}
def suffix(field):
    f = re.sub(r"^\d+_", "", str(field))
    f = re.sub(r"^(boom|wing|hull_array|core_layout|emp|fuselage)\d*_", "", f)
    return re.sub(r"^t\d+_", "", f)
def list_for(section, field): return CODED.get((section, suffix(field)))
""")

md("""
## 3 — Column registry

What the tables carry, in order, and what they deliberately drop. Patent-level fields are renamed
to snake_case; G1–M3 keep the wizard's own Field names so the codebook maps 1:1.
""")
code(r"""
# ── patent level (T1 + META), in output order: (wizard Field, table column) ──
PATENT_COLS = [
    ("isApproved", "is_approved"), ("aircraftName", "aircraft_name"), ("assignee", "assignee"),
    ("title", "title"), ("app_year", "app_year"), ("pub_year", "pub_year"), ("pdf_link", "pdf_link"),
    ("archCount", "n_variants_declared"),
    ("isDuplicate", "is_duplicate"), ("duplicateType", "dup_type"), ("duplicateId", "dup_of"),
    ("t1DisapproveReason", "reason_human"), ("t1DisapproveReason_otherNote", "reason_note"),
    ("imgNotReflect", "img_not_reflect"),
    ("t1_humanUncertain", "t1_humanUncertain"), ("t1_uncertainNote", "t1_uncertainNote"),
    ("t1_quickOverride", "t1_quickOverride"), ("t1_quickNote", "t1_quickNote"),
]
LEGACY_MERGE = {"disapproveOther": "reason_note"}      # legacy name on 4 Batch_05 records
DROP_PATENT = {"abstract": "long bibliographic text", "description_of_drawings": "long bibliographic text",
               "scope": "unreviewed SBERT pre-label", "t1Field": "unreviewed SBERT pre-label", "t1Target": "unreviewed SBERT pre-label",
               "labelToken": "SBERT token", "timestamp": "save stamp", "codebook_version": "schema stamp (user decision: not carried)",
               "familyId": "hardcoded placeholder", "mainFigure": "legacy patent-level flag, superseded by per-figure isMain",
               "comments": "free text, not analysed", "patentImageComments": "free text, not analysed"}
DROP_VARIANT_RE = {r"^longSym$": "retired v13", r"^empTiltsNote$": "retired v15.5", r"_symLong$": "retired v15.6",
                   r"_symCirc$": "retired v15.6", r"^footAmbiguous$": "constant True", r"^boom\d+_cards$": "structural marker"}
G1_ORDER = ["topType", "notPureArch", "edgeTags", "edgeTags_wizard",
            "uav_final", "uav_uncertain", "is_electric_final", "electric_uncertain", "electric_similar_final",
            "takeoff_final", "takeoff_uncertain", "identity_from",
            "arch_gt", "arch_gt_provenance", "arch_gt_visible",
            "g1_humanUncertain", "g1_uncertainNote", "g1_quickOverride", "g1_quickNote"]
M1_FIXED = ["fusShape", "fusShapeOtherNote", "fusKin", "gearArch", "latSym", "dinoUnderstanding", "boomsPresent"]
M1_BOOM  = ["count", "attach", "attachOth", "wingRel", "wingRelOth", "wingIdx", "span", "long", "longOth",
            "orient", "orientOth", "sym", "circSym", "hasProps", "tilts", "retracts", "onFixedPart"]   # onFixedPart: wizard v15.16
M1_TAIL  = ["wingNotes", "boomNotes", "fusNotes", "empNotes", "gearNotes", "m1_humanUncertain", "m1_uncertainNote", "m1_quickOverride"]
M2_FIXED = ["wingConf", "wingConfOtherNote", "wCount", "empType", "empTypeOtherNote", "empTilts"]
M2_WING  = ["role", "role_otherNote", "tilt", "posV", "posL", "plan", "plan_otherNote", "tipJoin"]
M2_TAIL  = ["m2_humanUncertain", "m2_uncertainNote", "m2_quickOverride", "m2_quickNote"]
M3_CARDS = ["fuselage", "wing1", "wing2", "wing3", "emp", "boom", "core_layout", "hull_array"]
M3_CARD  = ["count", "sym", "ntypes", "zone", "zoneChord", "zoneSpan", "quickOverride", "quickCount", "notes",
            "chord", "orient", "bmech", "rmech", "propKin", "propKinOth", "ctrlOnly"]
M3_TYPE  = ["count", "chord", "orient", "bmech", "rmech", "propKin", "propKinOth", "ctrlOnly", "zone", "zoneChord", "zoneSpan"]
M3_TAIL  = ["m3_humanUncertain", "m3_uncertainNote"]
COUNT_FIELDS_RE = re.compile(r"(^|_)(count|ntypes|quickCount|wCount|archCount)$")
""")

md("## 4 — Pivot: patent-level table and aircraft-level table")
code(r"""
def pivot(section_rows, index):
    d = section_rows[["Patent_ID", "Field", "Value"]].copy()
    dup = d.duplicated(["Patent_ID", "Field"], keep=False)
    if dup.any():
        log(f"  !! {dup.sum()} duplicated (id, Field) cells — keeping the LAST (record order); ids: "
            f"{sorted(d[dup].Patent_ID.unique())[:8]}")
    return d.drop_duplicates(["Patent_ID", "Field"], keep="last").pivot(index="Patent_ID", columns="Field", values="Value")

PT = pivot(L[L.Section.isin(["T1", "META"])], "Patent_ID")          # keyed on the BASE id
VT = pivot(L[L.Section.isin(["G1", "M1", "M2", "M3"])], "Patent_ID") # keyed on the wizard's aircraft id (bare or _archN)
PT.index.name = VT.index.name = None
assert all(base_id(p) == p for p in PT.index), "T1/META rows on an _archN id"
batch_of = L.drop_duplicates("base").set_index("base").batch.to_dict()
pos_of   = L.drop_duplicates("base").set_index("base")._pos.to_dict()

bare = {p for p in VT.index if arch_idx(p) == 1 and not p.endswith("_arch1")}
archd = {base_id(p) for p in VT.index if re.search(r"_arch\d+$", p)}
both = sorted(bare & archd)
if both: log(f"  !! {len(both)} patents carry morphology on the bare id AND on _archN ids (bare ignored): {both[:6]}")
VT = VT.drop(index=[p for p in both])
log(f"patent table {PT.shape}  aircraft table {VT.shape}")
var_ids = collections.defaultdict(list)
for v in VT.index: var_ids[base_id(v)].append(v)
for k in var_ids: var_ids[k].sort(key=arch_idx)
""")

md("## 5 — Assemble one row per aircraft (`<patent>_ua<N>`); rejected patents get one row with empty labels")
code(r"""
patent_cols = [c for c, _ in PATENT_COLS]
rows = []
for pid in PT.index:
    variants = var_ids.get(pid, [])
    if not variants:
        rows.append(dict(patent_id=pid, variant=1, variant_id=pid, n_variants=0))
    for v in variants:
        rows.append(dict(patent_id=pid, variant=arch_idx(v), variant_id=v, n_variants=len(variants)))
M = pd.DataFrame(rows)
M["ua"] = M.variant.astype(int)
M["aircraft_id"] = M.patent_id + "_ua" + M.ua.astype(str)
M["batch"] = M.patent_id.map(batch_of)
for src, dst in PATENT_COLS:
    M[dst] = M.patent_id.map(PT[src]) if src in PT.columns else np.nan
for src, dst in LEGACY_MERGE.items():
    if src in PT.columns:
        legacy = M.patent_id.map(PT[src]); fill = M[dst].map(is_empty) & ~legacy.map(is_empty)
        M.loc[fill, dst] = legacy[fill]; log(f"  {src}: {int(fill.sum())} legacy notes merged into {dst}")
variant_cols = [c for c in VT.columns if not any(re.search(p, c) for p in DROP_VARIANT_RE)]
VTm = VT[variant_cols].reindex(M.variant_id)
VTm.index = M.index
M = pd.concat([M, VTm], axis=1)
dropped_v = sorted(set(VT.columns) - set(variant_cols))
log(f"skeleton: {len(M)} aircraft rows ({M.patent_id.nunique()} patents), {len(variant_cols)} label columns kept, dropped {dropped_v}")
""")

md("## 6 — Normalise values: ids only, booleans, integer counts, the re-codes, the `Other` tags")
code(r"""
sec_of = L.drop_duplicates("Field").set_index("Field").Section.to_dict()
def norm_cell(v, lst, multi):
    if is_empty(v): return np.nan
    parts = str(v).split("|") if multi else [str(v)]
    out = [strip_label(p) for p in parts if p.strip()]
    return "|".join(out)
coded_cols, bool_cols, int_cols = [], [], []
for c in variant_cols:
    sec = sec_of.get(c); lst = list_for(sec, c) if sec else None
    if lst:
        coded_cols.append(c)
        M[c] = M[c].map(lambda v, lst=lst, m=(sec, suffix(c)) in MULTI: norm_cell(v, lst, m))
for c in ["reason_human", "dup_type"]:
    M[c] = M[c].map(lambda v: np.nan if is_empty(v) else strip_label(v))
for c in list(M.columns):
    vals = {str(x) for x in M[c].dropna().unique()}
    if vals and vals <= {"True", "False"} and not COUNT_FIELDS_RE.search(c):   # a count of 1/0 stored as True/False stays a count
        M[c] = M[c].map(lambda v: np.nan if is_empty(v) else str(v) == "True").astype("boolean"); bool_cols.append(c)
for c in list(M.columns):
    if COUNT_FIELDS_RE.search(c) and c not in bool_cols and c != "ua":
        M[c] = M[c].map(lambda v: np.nan if is_empty(v) else (1 if str(v) == "True" else 0 if str(v) == "False" else int(float(v)))).astype("Int64"); int_cols.append(c)
for c in ["app_year", "pub_year"]:
    M[c] = pd.to_numeric(M[c], errors="coerce").astype("Int64")
log(f"normalised: {len(coded_cols)} coded columns stripped to ids, {len(bool_cols)} boolean, {len(int_cols)} integer")

bad = collections.Counter()
for c in coded_cols:
    lst = list_for(sec_of[c], c); ok = set(LISTS[lst]) | ({"Main"} if suffix(c) == "role" else set())
    for v in M[c].dropna():
        for p in str(v).split("|"):
            if p not in ok: bad[(c, p)] += 1
if bad: log(f"  !! values that are not current option ids: {dict(bad)}")
else:   log("  every coded value is a current option id")

# 2026-09-18: the US2021284333A1 wing1_plan Oth→Trap recode is gone — it was the wizard's reload bug (Trap missing from
# the ingest whitelist, fixed v15.11), not a typing slip, and apply_damage_decisions.py restored Trap in the record itself.
RECODES = []
for vid, col, old, new, why in RECODES:
    i = M.index[M.variant_id == vid]
    assert len(i) == 1 and M.loc[i[0], col] == old, (vid, col, M.loc[i, col].tolist())
    M.loc[i[0], col] = new; log(f"  recode {vid} {col}: {old} -> {new}  ({why})")

tags = pd.read_csv(TAGS_CSV)
tags["k"] = tags.note_text.map(lambda t: re.sub(r"\s+", " ", str(t).strip().lower()))
tag_of = {(r.note_field, r.k): r.tag for r in tags.itertuples()}
def note_field_key(col): return re.sub(r"^(boom|wing)\d+_", r"\1N_", col)
PAIRS = [("attach", "attachOth"), ("wingRel", "wingRelOth"), ("long", "longOth"), ("orient", "orientOth"),
         ("fusShape", "fusShapeOtherNote"), ("empType", "empTypeOtherNote"), ("plan", "plan_otherNote"),
         ("role", "role_otherNote"), ("wingConf", "wingConfOtherNote"), ("propKin", "propKinOth")]
tag_cols, unmapped = [], collections.Counter()
for note_col in [c for c in variant_cols if any(c.endswith(s) for _, s in PAIRS)]:
    coded_sfx = next(cs for cs, ns in PAIRS if note_col.endswith(ns))
    coded_col = note_col[: -len(next(ns for cs, ns in PAIRS if note_col.endswith(ns)))] + coded_sfx
    tcol = coded_col + "_otherTag"
    nf = note_field_key(note_col) if not note_col.endswith("propKinOth") else "propKinOth"
    def _t(v, nf=nf):
        if is_empty(v): return np.nan
        k = re.sub(r"\s+", " ", str(v).strip().lower()); t = tag_of.get((nf, k))
        if t is None: unmapped[(nf, k)] += 1
        return t or "Misc"
    M[tcol] = M[note_col].map(_t); tag_cols.append(tcol)
log(f"  {len(tag_cols)} _otherTag columns; unmapped notes -> Misc: {sum(unmapped.values())}")

rej_mask = ~M.is_approved.fillna(False).astype(bool)
had = M[rej_mask & M.topType.notna()][["patent_id", "topType", "reason_human"]]
REJECTED_WITH_MORPH = had.copy()
if len(had): log(f"  {len(had)} rejected patents carried a morphology, blanked: {had.to_dict('records')}")
M.loc[rej_mask, variant_cols + tag_cols] = np.nan
if unmapped: log("  " + "; ".join(f"{k[0]}: {k[1][:50]!r}" for k in list(unmapped)[:10]))
""")

md(r"""
## 6b — Wizard view: does the record store what the wizard SHOWS?

The wizard locks some fields once the architecture is known and draws the locked value as a disabled chip. Before
v15.10 it exported whatever was stored underneath, so the screen and the Excel could disagree (2026-09-18: 122
approved figures showed Invariant but stored Hover/Cruise/…; 13 PTC/RC bodies showed Whole-Body Pitch but stored
Fixed). `src/wizard_view.py` replays every lock on the record — Aircraft State (fixed architecture whose record moves
nothing → Invariant), Body Motion (PTC/RC → Whole-Body Pitch; TB never Fixed), card-level Propulsor Articulation
(TW/DS/SLC/SRW/MR/TB/PTC/RC → Fixed), Wing Tilt (TB/PTC → Fixed; TW wing 1 → Tilt; RC/MR wingless). Any difference
FAILS §12; the fix is `python scripts/apply_wizard_locks.py --apply` (same replay), never a hand edit here.

Quick Count Overrides are the other place where the screen hides a stored value: the wizard skips the overridden
page or card, but a value typed before the tick is still exported. §7e lists every such value (flag
OVERRIDE_HIDES_VALUE) and derives its columns without it (ruling 2026-09-19).
""")
code(r"""
from src.wizard_view import replay, unlocked_fixed, INVARIANT_ID, INVARIANT_NAME
WV = replay(L)
# patents the user is fixing in the wizard themselves (review_decisions/PENDING_WIZARD_REVIEW.csv): their differences are
# flagged WIZARD_REVIEW_PENDING, not a build failure — remove a row from that file once the patent is re-saved
PENDING_CSV = RD / "PENDING_WIZARD_REVIEW.csv"
PENDING = pd.read_csv(PENDING_CSV, dtype=str, keep_default_na=False) if PENDING_CSV.exists() else pd.DataFrame(columns=["patent_id", "what_to_check"])
if PENDING_CSV.exists(): used(PENDING_CSV, "pending wizard review")
WV_PENDING = WV[WV.patent_id.isin(set(PENDING.patent_id))]
WV = WV[~WV.patent_id.isin(set(PENDING.patent_id))]
log(f"  {len(WV_PENDING)} differences on {WV_PENDING.patent_id.nunique()} patents the user is reviewing in the wizard (PENDING_WIZARD_REVIEW.csv)")
log(f"wizard view: {len(WV)} places where the record stores something other than what the wizard shows")
for r in WV.head(40).itertuples():
    log(f"  !! {r.patent_id} arch{r.arch} {r.field} {r.figure}: stored {r.stored!r}, wizard shows {r.wizard_shows!r} — {r.rule}")
if len(WV): log("  fix: python scripts/apply_wizard_locks.py --apply   (then re-run this notebook)")
# a PTC is defined by nothing articulating, yet the wizard unlocks the figure state when the record says a part moves
PTC_MOVES = unlocked_fixed(L).query("topType == 'PTC'")
for r in PTC_MOVES.itertuples(): log(f"  PTC whose record says a part moves (figure state unlocked in the wizard): {r.patent_id} arch{r.arch}: {r.moves}")
""")

md(r"""
## 7 — Duplicates: D1/D2 point, D3 keep their own

Ruling 2026-09-17: a D1 or D2 is the SAME aircraft as its original, so it carries **nothing from G1 to M3** — no copy,
no inheritance. Its labels are blank and `same_aircraft_as` names the aircraft it repeats: the one pinned by the
reviewer in `review_decisions/duplicate_root_variant.csv` (a letter → `_ua<N>`), otherwise every aircraft of the
original. A D2 ("same figures") has no figures of its own either. Whatever the wizard record holds inside a D1/D2's
own G1–M3 block is reported (informational) and ignored. A D3 is a different aircraft: own labels, own figures.
""")
code(r"""
dup_of   = {p: strip_label(v) for p, v in PT["duplicateId"].dropna().items() if not is_empty(v)}
dup_type = {p: strip_label(v) for p, v in PT["duplicateType"].dropna().items() if not is_empty(v)}
known = set(PT.index)
def chain_root(p):
    seen = [p]
    while p in dup_of and dup_type.get(p) in ("1", "2"):
        p = dup_of[p]
        if p in seen or p not in known: return p, False
        seen.append(p)
    return p, True
def top_original(p):
    seen = [p]
    while p in dup_of:
        p = dup_of[p]
        if p in seen or p not in known: return p
        seen.append(p)
    return p

PIN_CSV = RD / "duplicate_root_variant.csv"; used(PIN_CSV, "duplicate pins")
_pin = pd.read_csv(PIN_CSV, keep_default_na=False, dtype=str)
PIN = {r.patent_id: ord(r.answer.strip().lower()) - 96 for r in _pin.itertuples() if re.fullmatch(r"[a-z]", r.answer.strip().lower())}
log(f"  duplicate pins: {len(PIN)} pinned to one aircraft, {int((_pin.answer.str.strip().str.lower() == 'all').sum())} 'all'")

M["dup_root"] = np.nan; M["same_aircraft_as"] = np.nan; M["labels_inherited_from"] = np.nan
M["is_primary"] = True; M["dup_of_missing"] = False; M["points_to_ua"] = pd.array([None] * len(M), dtype="Int64")
g1m3_cols = list(variant_cols) + tag_cols
root_uas = {pid: sorted(g.ua.tolist()) for pid, g in M.groupby("patent_id")}
missing_root, cross_batch, d1_diff = [], 0, []
out_rows = []
for pid, grp in M.groupby("patent_id", sort=False):
    t = dup_type.get(pid)
    if t in ("1", "2"):
        root, ok = chain_root(pid)
        grp = grp.copy(); grp["dup_root"] = root; grp["is_primary"] = False
        if not ok or root not in var_ids:
            missing_root.append((pid, root)); grp["dup_of_missing"] = True; out_rows.append(grp); continue
        if batch_of.get(root) != batch_of.get(pid): cross_batch += 1
        rus = root_uas[root]
        for i, r in grp.iterrows():
            for c in g1m3_cols:                       # informational: what the wizard held inside the duplicate's own block
                a = r[c]
                if not is_empty(a): d1_diff.append(dict(patent_id=pid, dup_type=t, root=root, column=c, own=a))
            if pid in PIN and PIN[pid] in rus: targets = [PIN[pid]]
            elif r.ua in rus and len(rus) > 1 and pid not in PIN and r.n_variants > 1: targets = [r.ua]
            else: targets = rus
            grp.at[i, "same_aircraft_as"] = "; ".join(f"{root}_ua{k}" for k in targets)
            grp.at[i, "points_to_ua"] = targets[0]
        grp["labels_inherited_from"] = root            # kept for older readers: "the labels are those of this patent's aircraft"
        grp[g1m3_cols] = np.nan
        out_rows.append(grp)
    elif t == "3":
        grp = grp.copy(); grp["dup_root"] = top_original(pid); out_rows.append(grp)
    else:
        out_rows.append(grp)
M = pd.concat(out_rows, ignore_index=True)
D2_SET = {p for p, v in dup_type.items() if v == "2"}
D1_SET = {p for p, v in dup_type.items() if v == "1"}
n_pt = M.same_aircraft_as.notna()
log(f"duplicates: {len(dup_of)} links — D1 {len(D1_SET)}, D2 {len(D2_SET)}, D3 {sum(1 for v in dup_type.values() if v=='3')}")
log(f"  D1/D2 rows that POINT (labels blank): {int(n_pt.sum())}  (cross-batch originals: {cross_batch}); pinned to one aircraft: {int(M.same_aircraft_as.fillna('').str.contains(';').eq(False).where(n_pt, False).sum())}")
log(f"  cells the wizard held inside D1/D2 blocks (ignored): {len(d1_diff)} on {len({d['patent_id'] for d in d1_diff})} patents")
if missing_root: log(f"  !! {len(missing_root)} D1/D2 whose original has no labels (dup_of_missing=True): {missing_root[:8]}")
log(f"aircraft rows: {len(M)}, {M.patent_id.nunique()} patents, primary rows {int(M.is_primary.sum())}")
""")

md(r"""
## 7b — Identity review (03a) → names, identity answers, edge tags

Rulings 2026-09-14: `uav_final == UAVSimilar` → **UAVSimilar**; `is_electric_final == No` (or the ElectricSimilar tag
given on the review page) → **ElectricSimilar**; `takeoff_final == STOL` → **STOLSimilar**. V/STOL is still a VTOL.
`*_uncertain` is carried as a flag. Per-aircraft answers (`*_variants`) apply per `ua`.

**Duplicate groups** = an original aircraft + every D1/D2 that points at it. An answer the reviewer gave on ANY member is
the group's answer when the reviewed members agree; reviewed members that disagree keep their own answers and are
flagged. Tags go on the PRIMARY rows only (a D1/D2 row carries no G1 block — it points). `edgeTags` = the wizard's tags ∪
the identity tags; the wizard's own value stays in `edgeTags_wizard`. Rejected rows carry no identity answers and no tags.

**Names (2026-09-17):** the reviewed real name with its `vN` where one exists, otherwise the aircraft id; a D1/D2 takes
the name of the aircraft it points at. A name is "real" when it differs from the generated company numbering
(`aircraft_group`) the identity build fills in for cleared proposals.
""")
code(r"""
used(ID_XLSX, "identity workbook")
IDT = pd.read_excel(ID_XLSX, sheet_name="Identity", keep_default_na=False, dtype=str).set_index("patent_id")
log(f"identity: {len(IDT)} patents (review_status done: {int((IDT.review_status == 'done').sum())})")
missing_id = sorted(set(M.patent_id) - set(IDT.index))
assert not missing_id, f"patents missing from the identity workbook: {missing_id[:10]}"
NAMES_CSV = RD / "NAME_DECISIONS.csv"; used(NAMES_CSV, "name decisions")

STOL_TAG = "STOLSimilar"
if STOL_TAG not in LISTS["T1_EDGE_TAGS"]:
    LISTS["T1_EDGE_TAGS"] = LISTS["T1_EDGE_TAGS"] + [STOL_TAG]
    LABELS.setdefault("T1_EDGE_TAGS", {})[STOL_TAG] = "STOL only (no vertical take-off), but similar enough"
ID_FIELDS = {"uav_final": ("uav_human", "uav_uncertain"), "is_electric_final": ("is_electric_human", "electric_uncertain"),
             "takeoff_final": ("takeoff_human", "takeoff_uncertain"), "electric_similar_final": ("electric_similar_human", None)}
ID_UNC = ["uav_uncertain", "electric_uncertain", "takeoff_uncertain"]

def id_value(pid, col, v):
    r = IDT.loc[pid]; vc = col + "_variants"
    if vc in IDT.columns and r[vc].replace(";", "").strip():
        parts = [p.strip() for p in r[vc].split(";")]
        if len(parts) >= v: return parts[v - 1]
    return r[col].strip()
def is_true(s): return str(s).strip().upper() == "TRUE"
def tagset(v): return set() if is_empty(v) else set(str(v).split("|"))

appr_mask = M.is_approved.fillna(False).astype(bool)
M["edgeTags_wizard"] = M["edgeTags"]
M["group_root"] = M.dup_root.where(M.same_aircraft_as.notna(), M.patent_id)
M["group_ua"] = M.points_to_ua.where(M.same_aircraft_as.notna(), M.ua).astype(int)

# ── names ──
def _norm(s): return re.sub(r"\s+", " ", str(s or "")).strip().lower()
# 2026-09-17: a real name can coincide with the generated group name (Odys Aviation = "ODYS AVIATION"); a known/text
# ruling in NAME_DECISIONS.csv then wins over the "name == generated" test that detects the clear rulings
_KNOWN = set()
_dec_f = RD / "NAME_DECISIONS.csv"
if _dec_f.exists():
    _dec = pd.read_csv(_dec_f, dtype=str, keep_default_na=False)
    _KNOWN = set(_dec.loc[_dec.decision.isin(["known", "text"]), "patent_id"])
def real_name_of(pid, ua):
    # the reviewed name of aircraft (pid, ua) if it is a REAL name; None when it is only the generated company numbering
    if pid not in IDT.index: return None
    n = id_value(pid, "aircraft_name_final", ua); g = id_value(pid, "aircraft_group", ua)
    # a known ruling wins only for a typed name: a generated sibling of a multi-aircraft patent is the exact group string
    if not n or (_norm(n) == _norm(g) and (pid not in _KNOWN or n == g)): return None
    # 2026-09-17: a multi-aircraft patent keeps per-aircraft groups ("LEONARDO 2a"), but its unreviewed name can still be the
    # patent-level generated group ("LEONARDO 2") — that is company numbering, not a real name
    if pid not in _KNOWN and _norm(n) == _norm(IDT.at[pid, "aircraft_group"]): return None
    # 2026-09-22: a group that is the raw assignee list ("A (US); B") yields its first assignee as the
    # name — that is still the generated group, not a real name (Georgia Tech, I-Hub patents)
    if pid not in _KNOWN and ";" in str(g) and _norm(n) in {_norm(x) for x in str(g).split(";")}: return None
    return n
def name_for(r):
    if r.same_aircraft_as is not np.nan and not is_empty(r.same_aircraft_as):
        out = []
        for tgt in str(r.same_aircraft_as).split("; "):
            tp, tu = tgt.rsplit("_ua", 1); out.append(real_name_of(tp, int(tu)) or tgt)
        return "; ".join(out)
    return real_name_of(r.patent_id, r.ua) or r.aircraft_id
M["aircraft_name_wizard"] = M["aircraft_name"]
M["aircraft_name"] = [name_for(r) for r in M.itertuples()]
M["name_is_real"] = [n != a and not re.search(r"_ua\d+$", str(n)) for n, a in zip(M.aircraft_name, M.aircraft_id)]
M.loc[~appr_mask, "name_is_real"] = False
_base_names = M.loc[M.name_is_real & appr_mask, "aircraft_name"].str.replace(r" v\d+$", "", regex=True).nunique()
log(f"  names: {int((M.name_is_real & appr_mask).sum())} approved rows carry a real name ({_base_names} base names); the rest are named by their id")

for c in list(ID_FIELDS) + ID_UNC + ["identity_from"]:
    M[c] = np.nan
M[list(ID_FIELDS) + ["identity_from"]] = M[list(ID_FIELDS) + ["identity_from"]].astype(object)
TAG_OF_FIELD = {"uav_final": "UAVSimilar", "is_electric_final": "ElectricSimilar", "electric_similar_final": "ElectricSimilar",
                "takeoff_final": STOL_TAG}
ID_CONFLICTS, lifted, src_of = [], collections.Counter(), collections.defaultdict(list)
CONTESTED = collections.defaultdict(set)
for (root, v), g in M[appr_mask].groupby(["group_root", "group_ua"]):
    members = list(dict.fromkeys([root] + g.patent_id.tolist()))
    for fcol, (hcol, ucol) in ID_FIELDS.items():
        reviewed = {}
        for p in members:
            if p not in IDT.index: continue
            val = id_value(p, fcol, v)
            if val and (id_value(p, hcol, v) or IDT.at[p, "review_status"] == "done"):
                reviewed[p] = val
        vals = set(reviewed.values())
        if TAG_OF_FIELD[fcol] in CONTESTED[(root, v)] and len(vals) == 1:
            vals = {None, *vals}
        for i, r in g.iterrows():
            if len(vals) == 1:   src = r.patent_id if r.patent_id in reviewed else next(iter(reviewed))
            elif len(vals) > 1:  src = r.patent_id if r.patent_id in reviewed else root
            else:                src = root
            if src != root: lifted[fcol] += 1
            val = id_value(src, fcol, v) if src in IDT.index else ""
            M.at[i, fcol] = val if val else np.nan
            if ucol and src in IDT.index and is_true(IDT.at[src, ucol]): M.at[i, ucol] = True
            if src not in src_of[i]: src_of[i].append(src)
        if len(vals) > 1 and None in vals: continue
        if len(vals) > 1:
            CONTESTED[(root, v)].add(TAG_OF_FIELD[fcol])
            ID_CONFLICTS.append(dict(root=root, variant=v, field=fcol, answers="; ".join(f"{p}={a}" for p, a in reviewed.items())))
for i, s in src_of.items(): M.at[i, "identity_from"] = "; ".join(s)
for c in ID_UNC: M[c] = M[c].astype("boolean")
log(f"  rows whose answer came from a reviewed group member other than the original: {dict(lifted)}")
if ID_CONFLICTS: log(f"  !! {len(ID_CONFLICTS)} duplicate groups whose reviewed members disagree (own answers kept): {ID_CONFLICTS}")

def id_tags(r):
    t = set()
    if r.uav_final == "UAVSimilar": t.add("UAVSimilar")
    if r.is_electric_final == "No" or r.electric_similar_final == "ElectricSimilar": t.add("ElectricSimilar")
    if r.takeoff_final == "STOL": t.add(STOL_TAG)
    return t
group_tags = collections.defaultdict(set)
for i, r in M[appr_mask].iterrows(): group_tags[(r.group_root, r.group_ua)] |= id_tags(r) - CONTESTED[(r.group_root, r.group_ua)]
ROW_TAGS = {i: group_tags[(r.group_root, r.group_ua)] | (id_tags(r) & CONTESTED[(r.group_root, r.group_ua)])
            for i, r in M[appr_mask].iterrows()}
TAG_ORDER = LISTS["T1_EDGE_TAGS"]
WIZARD_TAG_KEEP = {("US11124286B1", "UAVSimilar"): "user 2026-09-14: UAV but similar stays, although uav_final=No"}
def contradicted(i):
    out = {}
    if M.at[i, "uav_final"] == "No": out["UAVSimilar"] = "uav_final=No"
    if M.at[i, "is_electric_final"] in ("Yes", "Hybrid") and M.at[i, "electric_similar_final"] != "ElectricSimilar":
        out["ElectricSimilar"] = f"is_electric_final={M.at[i, 'is_electric_final']}"
    if M.at[i, "takeoff_final"] in ("VTOL", "V/STOL"): out[STOL_TAG] = f"takeoff_final={M.at[i, 'takeoff_final']}"
    return out
TAG_REMOVED, TAG_KEPT = [], []
def merged(i):
    if not M.at[i, "is_primary"]: return np.nan            # a D1/D2 points; its G1 block is blank
    t = tagset(M.at[i, "edgeTags_wizard"])
    if appr_mask[i]:
        pid = M.at[i, "patent_id"]; group = ROW_TAGS[i]
        for tag, why in contradicted(i).items():
            if tag in t and tag not in group:
                if (pid, tag) in WIZARD_TAG_KEEP: TAG_KEPT.append((pid, tag, why, WIZARD_TAG_KEEP[(pid, tag)]))
                else: t.discard(tag); TAG_REMOVED.append((pid, tag, why))
        t |= group
    return "|".join([x for x in TAG_ORDER if x in t] + sorted(t - set(TAG_ORDER))) if t else np.nan
M["edgeTags"] = [merged(i) for i in M.index]
TAG_REMOVED, TAG_KEPT = sorted(set(TAG_REMOVED)), sorted(set(TAG_KEPT))
unused_keep = set(WIZARD_TAG_KEEP) - {(p, t) for p, t, _, _ in TAG_KEPT}
if unused_keep: log(f"  !! WIZARD_TAG_KEEP entries that matched nothing: {sorted(unused_keep)}")
A_ = M[appr_mask & M.is_primary]
for t in TAG_ORDER:
    has = A_.edgeTags.map(lambda v, t=t: t in tagset(v)); had = A_.edgeTags_wizard.map(lambda v, t=t: t in tagset(v))
    log(f"  edge tag {t:<16}: {int(has.sum()):>4} primary rows / {A_[has].patent_id.nunique():>3} patents   (wizard alone: {int(had.sum())} rows)")
log(f"  takeoff_final (primary approved): {A_.takeoff_final.value_counts().to_dict()}   is_electric_final: {A_.is_electric_final.value_counts().to_dict()}")
log(f"  uncertainty flags carried: { {c: int(A_[c].fillna(False).sum()) for c in ID_UNC} }")
log(f"  wizard tags removed (identity answer contradicts): {TAG_REMOVED}")
log(f"  wizard tags kept by WIZARD_TAG_KEEP: {TAG_KEPT}")
M = M.drop(columns=["group_root", "group_ua"])
""")

md(r"""
## 7c — Wizard ML pre-labels (`ml_predict_labels_<batch>.xlsx`) beside the human labels — tables only

The machine side of the wizard record, carried with an `ml_` prefix in `tables/aircraft_table.csv` and never merged
into the human columns nor into the workbook. Every ML column is the patent's OWN prediction.
""")
code(r"""
mlf = []
for b in BATCHES:
    f = LAB_DIR / b / f"ml_predict_labels_{b}.xlsx"; used(f, f"ML feed {b}")
    d = pd.read_excel(f, sheet_name="Review", keep_default_na=False, dtype=str)
    d["batch"] = b; mlf.append(d)
ML = pd.concat(mlf, ignore_index=True); ML["Field"] = ML.Field.astype(str)
log(f"ML feed: {len(ML)} rows, {ML.Patent_ID.nunique()} patents")
missing_ml = sorted(set(PT.index) - set(ML.Patent_ID))
if missing_ml: log(f"  !! {len(missing_ml)} patents absent from the ML feed: {missing_ml[:8]}")

def ml_pivot(rows, what, index="Patent_ID"):
    keys = [index, "Field"] if isinstance(index, str) else list(index) + ["Field"]
    return rows.drop_duplicates(keys, keep="last").pivot(index=index, columns="Field", values=what)
def clean(v, coded=False):
    if is_empty(v): return np.nan
    return strip_label(v) if coded else str(v).strip()
def conf(v):
    try: return round(float(v), 3)
    except (TypeError, ValueError): return np.nan
def to_bool(col): return col.map(lambda v: np.nan if is_empty(v) else str(v) == "True").astype("boolean")

t1 = ML[ML.Section.isin(["T1", "META"])]
V, C, S = ml_pivot(t1, "Value"), ml_pivot(t1, "Confidence"), ml_pivot(t1, "Source")
ML_T1 = ["scope", "t1Field", "t1Target", "isApproved", "isDuplicate", "duplicateId", "duplicateType"]
ML_SOURCE_FOR = {"scope", "isApproved", "isDuplicate", "topType"}
ml_head = []
for f in ML_T1:
    if f not in V.columns: continue
    c = "ml_" + f
    M[c] = M.patent_id.map(V[f]).map(lambda v, k=(f == "duplicateType"): clean(v, k)); ml_head.append(c)
    if f in ("isApproved", "isDuplicate"): M[c] = to_bool(M[c])
    if f in C.columns and C[f].replace("", np.nan).notna().any():
        M[c + "_conf"] = M.patent_id.map(C[f]).map(conf); ml_head.append(c + "_conf")
    if f in ML_SOURCE_FOR and f in S.columns:
        M[c + "_source"] = M.patent_id.map(S[f]).map(clean); ml_head.append(c + "_source")
t2d = ML[(ML.Field == "dupOfPatent") & (ML.Value != "")]
M["ml_n_fig_dup_suggestions"] = M.patent_id.map(t2d.groupby("Patent_ID").size()).fillna(0).astype(int)
M["ml_fig_dup_patents"] = M.patent_id.map(t2d.groupby("Patent_ID").Value.agg(lambda s: "|".join(sorted(set(s)))))
ml_head += ["ml_n_fig_dup_suggestions", "ml_fig_dup_patents"]

morph = ML[ML.Section.isin(["G1", "M1", "M2", "M3"])]
ML_MORPH = ["topType", "fusShape", "fusKin", "gearArch", "latSym", "wingConf", "wCount", "empType"]
ml_morph = []
if len(morph):
    Vm, Cm, Sm = ml_pivot(morph, "Value"), ml_pivot(morph, "Confidence"), ml_pivot(morph, "Source")
    for f in ML_MORPH:
        if f not in Vm.columns: continue
        c = "ml_" + f
        M[c] = M.patent_id.map(Vm[f]).map(lambda v: clean(v, True)); ml_morph.append(c)
        if f in Cm.columns: M[c + "_conf"] = M.patent_id.map(Cm[f]).map(conf); ml_morph.append(c + "_conf")
        if f in ML_SOURCE_FOR and f in Sm.columns: M[c + "_source"] = M.patent_id.map(Sm[f]).map(clean); ml_morph.append(c + "_source")
    if "ml_latSym" in M.columns: M["ml_latSym"] = to_bool(M["ml_latSym"])
    if "ml_topType" in M.columns:
        n_tp = int((M.ml_topType == "TP").sum()); M["ml_topType"] = M.ml_topType.replace({"TP": "TR"})
        log(f"  ml_topType: {n_tp} rows recoded TP -> TR (retired code)")
    log(f"  G1–M3 predictions exist for {morph.Patent_ID.nunique()} patents (batches {sorted(morph.batch.unique())}); carried {ML_MORPH}")
else:
    log("  no G1–M3 predictions in any feed")

def agrees(a, b):
    both = a.notna() & b.notna()
    return pd.array(np.where(both, a.astype(str).values == b.astype(str).values, None), dtype="boolean")
M["ml_approved_agrees"] = agrees(M.ml_isApproved, M.is_approved)
M["ml_dup_agrees"] = agrees(M.ml_duplicateId, M.dup_of) if "ml_duplicateId" in M.columns else pd.array([None] * len(M), dtype="boolean")
ml_head += ["ml_approved_agrees", "ml_dup_agrees"]
if "ml_topType" in M.columns:
    M["ml_topType_agrees"] = agrees(M.ml_topType, M.topType); ml_morph.append("ml_topType_agrees")
    both = M.ml_topType.notna() & M.topType.notna()
    log(f"  ml_topType vs human topType: {int(M.ml_topType_agrees[both].sum())} / {int(both.sum())} agree")
log(f"  ML columns added: {len(ml_head) + len(ml_morph)}")
""")

md(r"""
## 7d — Architecture ground truth (03b) on the primary approved aircraft

`text_architecture/architecture_text_final_variants.csv` (one row per aircraft, `variant_id` in the wizard's `_archN`
form) → `arch_gt` (the type the whole document states), `arch_gt_provenance`, `arch_gt_visible` (is that type visible in
the figures). Joined on `aircraft_id`. A D1/D2 row points at its aircraft and carries no ground truth of its own.
`unsure` / `not_stated` are kept as written — the analysis decides what counts as evidence.
""")
code(r"""
GT_CSV = TA / "architecture_text_final_variants.csv"; used(GT_CSV, "architecture GT")
GT = pd.read_csv(GT_CSV, keep_default_na=False, dtype=str)
GT["aircraft_id"] = GT.variant_id.map(lambda v: re.sub(r"_arch(\d+)$", r"_ua\1", v) if re.search(r"_arch\d+$", v) else v + "_ua1")
dupgt = GT[GT.duplicated("aircraft_id", keep=False)]
if len(dupgt): log(f"  !! {len(dupgt)} GT rows share an aircraft_id (first kept): {sorted(dupgt.aircraft_id.unique())[:8]}")
GT = GT.drop_duplicates("aircraft_id").set_index("aircraft_id")
gt_mask = appr_mask & M.is_primary
for src, dst in [("arch_final", "arch_gt"), ("provenance", "arch_gt_provenance"), ("visible_in_figures", "arch_gt_visible")]:
    M[dst] = np.nan; M[dst] = M[dst].astype(object)
    M.loc[gt_mask, dst] = M.loc[gt_mask, "aircraft_id"].map(GT[src]).replace("", np.nan)
# no type + "not_stated_in_patent" (or the older "unsure") = the reviewer read the whole document and it does not state
# the type: a ruling, not a gap (user 2026-09-17: named "not stated in the patent")
NOT_STATED_PROV = {"not_stated_in_patent", "unsure"}
GT_UNSURE  = sorted(M[gt_mask & M.arch_gt.isna() & M.arch_gt_provenance.isin(NOT_STATED_PROV)].aircraft_id)
GT_MISSING = sorted(M[gt_mask & M.arch_gt.isna() & ~M.arch_gt_provenance.isin(NOT_STATED_PROV)].aircraft_id)
if GT_UNSURE: log(f"  ground truth NOT STATED IN THE PATENT (no type): {GT_UNSURE}")
GT_ORPHANS = sorted(set(GT.index) - set(M[gt_mask].aircraft_id))
log(f"ground truth: {len(GT)} aircraft in the file; {int((gt_mask & M.arch_gt.notna()).sum())} / {int(gt_mask.sum())} primary approved aircraft matched")
log(f"  provenance: {M.loc[gt_mask, 'arch_gt_provenance'].value_counts(dropna=False).to_dict()}")
if GT_MISSING: log(f"  !! {len(GT_MISSING)} primary approved aircraft without a ground truth: {GT_MISSING[:10]}")
if GT_ORPHANS: log(f"  !! {len(GT_ORPHANS)} GT rows that match no primary approved aircraft: {GT_ORPHANS[:10]}")
""")

md(r"""
## 7e — Overrides, boom thrust and wing-borne thrust (rulings 2026-09-19)

Columns derived per **primary approved** aircraft and placed beside the labels they read (a D1/D2 points and a rejected
row carries nothing, so theirs are blank). No label column is changed. A value a rule cannot determine is left blank
with its reason — never read as zero, none or Fixed — and the number left out is logged.

**1 · Overrides.** *An override keeps only what it records. After a G1 override the aircraft has no type; after an M1 or
M2 override the skipped fields are not determinable; at an M3 station only the propulsor count is kept (the wizard
stores `<station>_count = 0` and the real number in `<station>_quickCount`). A value that is not determinable is never
read as zero, none or Fixed: it is left out, and the number left out is reported.*
`overrides` = the overrides ticked (`G1|M2|M3:wing1`). `propulsor_units` = the non-control-only propulsors over every
M3 station, the quick count on a count-only station; blank for HB/PFV (the codebook excludes them), after a G1
override, and when an overridden station has no quick count (`propulsor_units_note` says why). A quick count on a
station whose override is off is ignored (LEFTOVER_QUICK_COUNT); a stored detail value behind an override is reported
(OVERRIDE_HIDES_VALUE), never used.

**2 · Tilting-boom rotors.** *The boom tick carries the tilt. A rotor on a boom group ticked Booms tilt is recorded Fixed,
as on a tilting wing. Tick Booms tilt only when the boom itself rotates; otherwise a rotor that changes angle on a boom
is recorded Tilt. A rotor that also pivots on a tilting boom is recorded Other, with a note.*
`boom_thrust_state` ∈ none / fixed / tilting / on_tilting_wing / mixed, over every (propeller-carrying boom group ×
non-control-only boom-card rotor type): a group rides a tilting wing when it is attached to Wings/Both, a wing is set to
Tilt (the one named by `boomN_wingIdx`, any when blank/Multi) and `boomN_onFixedPart` is not True; otherwise it tilts
when `boomN_tilts` is True or the rotor's propKin is Tilt; otherwise fixed. Groups with `hasProps` False are ignored;
blank when an M1 or boom-station override hides the groups or the rotors. Checks: BOOM_TICK_AND_ROTOR_TILT (Tilt rotor
on an aircraft whose every propeller-carrying boom group is ticked — 0 once migrated), BOOM_ROTORS_WITHOUT_GROUP.

**3 · Wing-borne thrust.** *A member on a wing that carries a propulsor is a boom when it carries propulsors both ahead of
and behind the wing, reaches beyond the leading or trailing edge by more than the local wing chord, or carries a tail
surface. Otherwise it is a nacelle or pylon, and the propulsor is recorded on the wing card at LE, TE, Above or Below.
Where the analysis asks whether the wing carries the thrust, both count as wing-borne.*
`wing_thrust_carrier` ∈ wing (wing cards only) / wing_boom (boom groups on the wing only) / both / none — winged aircraft
only. `wing_borne_units` = propulsors on the wing cards + on boom groups attached to Wings/Both; blank (with
`wing_borne_units_note`) when the boom-card rotors cannot be apportioned. `wing_boom_candidate` = the members to re-read
under rule 3: wing-attached boom groups that all lie only fore or only aft of the wing, vertical booms straddling the
wing, and wing-card propulsors at Above/Below.

**4 · Two flight states in one figure (R2-02(e))** — a figure check, §9: STATE_NOTE_MENTIONS_TWO_STATES.
""")
code(r"""
from src.wizard_view import PROPKIN_FIXED
def _s(v):
    if v is None or v is pd.NA or (isinstance(v, float) and np.isnan(v)): return ""
    s = str(v).strip(); return "" if s in ("nan", "None", "NaT", "<NA>") else s
def _t(v): return _s(v) == "True"
def _n(v):
    s = _s(v)
    if s in ("", "False"): return 0
    if s == "True": return 1
    try: return int(float(s))
    except ValueError: return 0
NO_UNITS_TYPES = {"HB", "PFV"}                       # codebook: hoverbikes and PFVs are not counted in propulsor units
WING_CARDS = ["wing1", "wing2", "wing3"]
WING_ATTACH = ("Wings", "Both")

def station_units(r, st):
    # (units, types, reason) at one M3 station: non-control-only propulsors; a count-only (override) station gives its
    # quick count and no type detail. units None = not determinable (override without a quick count)
    if _t(r.get(f"{st}_quickOverride")):
        q = _s(r.get(f"{st}_quickCount"))
        if not q: return None, [], f"M3:{st} override without a quick count"
        n = _n(q); return n, ([dict(count=n, propKin="", zoneChord="")] if n > 0 else []), ""
    nt = _n(r.get(f"{st}_ntypes")) or 1
    tys = []
    for pre in ([f"{st}_t{k}_" for k in range(1, nt + 1)] if nt > 1 else [f"{st}_"]):
        c = _n(r.get(pre + "count"))
        if c > 0 and not _t(r.get(pre + "ctrlOnly")):
            tys.append(dict(count=c, propKin=_s(r.get(pre + "propKin")), zoneChord=_s(r.get(pre + "zoneChord"))))
    return sum(t["count"] for t in tys), tys, ""

def boom_groups(r):
    # the boom groups that carry propulsors: recorded (count or attachment) and hasProps not False
    out = []
    for k in range(1, 7):
        att, cnt = _s(r.get(f"boom{k}_attach")), _n(r.get(f"boom{k}_count"))
        if (not att and cnt == 0) or _s(r.get(f"boom{k}_hasProps")) == "False": continue
        out.append(dict(k=k, attach=att, wingIdx=_s(r.get(f"boom{k}_wingIdx")), wingRel=_s(r.get(f"boom{k}_wingRel")),
                        orient=_s(r.get(f"boom{k}_orient")), tilts=_t(r.get(f"boom{k}_tilts")), on_fixed=_t(r.get(f"boom{k}_onFixedPart"))))
    return out
def rides_tilting_wing(g, tw):
    if g["attach"] not in WING_ATTACH or not tw or g["on_fixed"]: return False
    if g["wingIdx"] in ("", "Multi"): return True
    m = re.fullmatch(r"W(\d)", g["wingIdx"]); return bool(m) and int(m.group(1)) in tw

# what each override hides in the wizard (pageGate / pageM1 / pageM2 return early; m3Card shows only the quick count
# and the notes). Values the wizard writes by itself are not answers: defaults (False, count 0, one type, latSym True,
# wing role Main; the card `sym` tick defaults per card) and the propKin lock (Fixed on PROPKIN_FIXED types).
HIDDEN_BY = {"G1": ["topType", "notPureArch"],
             "M1": M1_FIXED + [f"boom{k}_{f}" for k in range(1, 7) for f in M1_BOOM],
             "M2": M2_FIXED + [f"wing{k}_{f}" for k in range(1, 4) for f in M2_WING]}
def hidden_fields(o):
    if o in HIDDEN_BY: return HIDDEN_BY[o]
    st = o.split(":", 1)[1]
    return ([f"{st}_{f}" for f in M3_CARD if f not in ("quickOverride", "quickCount", "notes", "sym")]
            + [f"{st}_t{k}_{f}" for k in range(1, 8) for f in M3_TYPE])   # wizard M3_MAX_TYPES = 7 (2026-09-19)
def is_default(field, v, tt):
    s, f = _s(v), suffix(field)
    return (s in ("", "False") or (f in ("count", "wCount") and s == "0") or (f == "ntypes" and s in ("0", "1"))
            or (f == "latSym" and s == "True") or (f == "role" and s == "Main") or (f == "propKin" and s == "Fixed" and tt in PROPKIN_FIXED))

RULE_COLS = ["overrides", "propulsor_units", "propulsor_units_note", "boom_thrust_state",
             "wing_thrust_carrier", "wing_borne_units", "wing_borne_units_note", "wing_boom_candidate"]
RV = {c: [None] * len(M) for c in RULE_COLS}
OVR_HIDDEN, LEFTOVER_QC, BOOM_TICK_TILT, BOOM_NO_GROUP = [], [], [], []
_pa = M.is_approved.fillna(False).astype(bool) & M.is_primary.astype(bool)
for pos, (i, row) in enumerate(M.iterrows()):
    if not _pa[i]: continue
    r = row.to_dict(); tt = _s(r.get("topType")); aid, pid = r["aircraft_id"], r["patent_id"]
    # ── 1 · overrides ──
    ov = [s for s in ("G1", "M1", "M2") if _t(r.get(f"{s.lower()}_quickOverride"))]
    ov += [f"M3:{st}" for st in M3_CARDS if _t(r.get(f"{st}_quickOverride"))]
    RV["overrides"][pos] = "|".join(ov) if ov else None
    for o in ov:
        hid = [f"{f}={_s(r.get(f))}" for f in hidden_fields(o) if not is_default(f, r.get(f), tt)]
        if hid: OVR_HIDDEN.append((pid, aid, o, hid))
    for st in M3_CARDS:
        q = _s(r.get(f"{st}_quickCount"))
        if q and not _t(r.get(f"{st}_quickOverride")): LEFTOVER_QC.append((pid, aid, st, q))
    su = {st: station_units(r, st) for st in M3_CARDS}
    why = ([f"{tt}: the codebook excludes hoverbikes and PFVs from propulsor counts"] if tt in NO_UNITS_TYPES else []) \
        + (["G1 override: the aircraft has no type"] if "G1" in ov else []) + [su[st][2] for st in M3_CARDS if su[st][0] is None]
    if why: RV["propulsor_units_note"][pos] = "; ".join(why)
    else:   RV["propulsor_units"][pos] = int(sum(su[st][0] for st in M3_CARDS))
    # ── 2 · boom thrust ──
    gs = boom_groups(r); tw = {k for k in range(1, 5) if _s(r.get(f"wing{k}_tilt")) == "Tilt"}
    bu, btys, _ = su["boom"]; m1_ovr, boom_ovr = "M1" in ov, "M3:boom" in ov
    boom_card_n = (bu or 0) if boom_ovr else _n(r.get("boom_count"))
    if boom_card_n > 0 and not gs and not m1_ovr: BOOM_NO_GROUP.append((pid, aid, boom_card_n))
    if gs and not m1_ovr and not boom_ovr and all(g["tilts"] for g in gs) and any(t["propKin"] == "Tilt" for t in btys):
        BOOM_TICK_TILT.append((pid, aid, tt, len(gs)))
    if m1_ovr or boom_ovr: state = None
    elif not btys: state = "none"
    else:
        cats = set()
        for t in btys:
            for g in (gs or [None]):
                if g is not None and rides_tilting_wing(g, tw): cats.add("on_tilting_wing")
                elif (g is not None and g["tilts"]) or t["propKin"] == "Tilt": cats.add("tilting")
                else: cats.add("fixed")
        state = cats.pop() if len(cats) == 1 else "mixed"
    RV["boom_thrust_state"][pos] = state
    # ── 3 · wing-borne thrust ──
    if "M2" in ov: wwhy = ["M2 override: the wings are not determinable"]
    elif _n(r.get("wCount")) <= 0: wwhy = ["wingless"]
    else: wwhy = (["M1 override: the boom groups are not determinable"] if m1_ovr else []) + [su[st][2] for st in WING_CARDS + ["boom"] if su[st][0] is None]
    if wwhy:
        RV["wing_borne_units_note"][pos] = "; ".join(wwhy); continue
    wty = [t for st in WING_CARDS for t in su[st][1]]
    wgs = [g for g in gs if g["attach"] in WING_ATTACH]; ogs = [g for g in gs if g["attach"] not in WING_ATTACH]
    on_wing_boom = bool(btys) and bool(wgs)
    RV["wing_thrust_carrier"][pos] = "both" if (wty and on_wing_boom) else "wing_boom" if on_wing_boom else "wing" if wty else "none"
    rel = {g["wingRel"] for g in wgs}
    RV["wing_boom_candidate"][pos] = bool((on_wing_boom and (rel in ({"Fore"}, {"Aft"}) or any(g["orient"] == "Vert" and g["wingRel"] == "Straddle" for g in wgs)))
                                          or any(t["zoneChord"] in ("Above", "Below") for t in wty))
    uwhy = ([f"{tt}: the codebook excludes hoverbikes and PFVs from propulsor counts"] if tt in NO_UNITS_TYPES else []) \
         + (["G1 override: the aircraft has no type"] if "G1" in ov else [])
    if bu > 0 and not gs: uwhy.append("the boom card's rotors have no propeller-carrying boom group")
    elif bu > 0 and wgs and ogs: uwhy.append("the boom card's rotors cannot be apportioned: boom groups mix wing (Wings/Both) and other attachment")
    if uwhy: RV["wing_borne_units_note"][pos] = "; ".join(uwhy)
    else:    RV["wing_borne_units"][pos] = int(sum(su[st][0] for st in WING_CARDS) + (bu if wgs else 0))
for c in RULE_COLS:
    M[c] = pd.Series(RV[c], index=M.index, dtype=object)
for c in ("propulsor_units", "wing_borne_units"): M[c] = pd.array(M[c].tolist(), dtype="Int64")
M["wing_boom_candidate"] = pd.array(M["wing_boom_candidate"].tolist(), dtype="boolean")

# column dictionary entries (tables/data_dictionary.csv, section "derived" — not a label slot)
RULE_DOC = {
    "overrides": ("", "Quick Count Overrides ticked on this aircraft: G1 | M1 | M2 | M3:<station> (empty = none). An override keeps only what it records: after G1 the aircraft has no type, after M1/M2 the skipped fields are not determinable, at an M3 station only the propulsor count is kept (ruling 2026-09-19)"),
    "propulsor_units": ("", "non-control-only propulsors over every M3 station; a count-only (override) station counts its quick count. Blank for HB/PFV (the codebook excludes them), after a G1 override, and when an overridden station has no quick count — never read as 0"),
    "propulsor_units_note": ("", "why propulsor_units is blank"),
    "boom_thrust_state": ("none=no non-control-only rotor on the boom card | fixed=every boom rotor fixed | tilting=boom group ticked Booms tilt or rotor propKin Tilt | on_tilting_wing=group rides a tilting wing | mixed=more than one of these",
                          "state of the thrust carried on booms (ruling 2026-09-19: the boom tick carries the tilt; a rotor on a boom group ticked Booms tilt is recorded Fixed, as on a tilting wing; tick Booms tilt only when the boom itself rotates, otherwise a rotor that changes angle on a boom is Tilt; a rotor that also pivots on a tilting boom is Other, with a note). Groups with hasProps False ignored; blank when an M1 or boom-station override hides the groups or rotors"),
    "wing_thrust_carrier": ("wing=propulsors on the wing cards only | wing_boom=on boom groups attached to Wings/Both only | both=both | none=neither",
                            "winged aircraft only (blank when wingless or after an M2/M1 override). Ruling 2026-09-19: a member on a wing carrying a propulsor is a boom when it carries propulsors ahead of and behind the wing, reaches beyond the LE/TE by more than the local chord, or carries a tail surface; otherwise a nacelle/pylon recorded on the wing card (LE/TE/Above/Below). Both count as wing-borne"),
    "wing_borne_units": ("", "propulsors on the wing cards + on boom groups attached to Wings/Both (non-control-only; quick counts on count-only stations). Blank with wing_borne_units_note when the boom-card rotors cannot be apportioned, and for HB/PFV / G1 override as propulsor_units"),
    "wing_borne_units_note": ("", "why wing_borne_units is blank (wingless, override, boom rotors that cannot be apportioned)"),
    "wing_boom_candidate": ("", "True = a member to re-read under the 2026-09-19 boom/nacelle rule: the wing-attached boom groups all lie only fore or only aft of the wing, a vertical boom straddles the wing, or a wing-card propulsor sits Above/Below the wing"),
    "empTilts": (None, "True when the empennage tilts, with the wing or on its own"),
}

_ppa = M[_pa]; _gate = _ppa.edgeTags.isna()
log("rules 2026-09-19 (7e):")
log("  1 overrides — an override keeps only what it records; a value not determinable is left out, never read as zero, none or Fixed")
log(f"    overrides on primary approved aircraft: {int(_ppa.overrides.notna().sum())}  {_ppa.overrides.value_counts().to_dict()}")
log(f"    propulsor_units: {int(_ppa.propulsor_units.notna().sum())} / {len(_ppa)} primary approved non-blank "
    f"({int(_ppa[_gate].propulsor_units.notna().sum())} / {int(_gate.sum())} without a Similar tag); left out: "
    f"{_ppa.propulsor_units_note.dropna().str.replace(r'^(HB|PFV):.*', 'HB/PFV', regex=True).value_counts().to_dict()}")
log(f"    values stored behind an override: {len(OVR_HIDDEN)} (OVERRIDE_HIDES_VALUE)   quick counts on a station whose override is off: {len(LEFTOVER_QC)} (ignored, LEFTOVER_QUICK_COUNT)")
log("  2 booms — the boom tick carries the tilt: a rotor on a ticked boom group is Fixed, as on a tilting wing; Other if it also pivots")
log(f"    boom_thrust_state: {_ppa.boom_thrust_state.value_counts(dropna=False).to_dict()}")
log(f"    boom tick + boom-card rotor Tilt: {len(BOOM_TICK_TILT)} (BOOM_TICK_AND_ROTOR_TILT; 0 once migrated)   boom rotors without a boom group: {len(BOOM_NO_GROUP)}")
log("  3 wing-borne thrust — a wing member is a boom only when it carries propulsors fore and aft, reaches beyond LE/TE by more than the chord, or carries a tail; both count as wing-borne")
_w = _ppa.wing_thrust_carrier.notna()
log(f"    wing_thrust_carrier (winged, determinable: {int(_w.sum())}): {_ppa.wing_thrust_carrier.value_counts().to_dict()}")
log(f"    wing_borne_units non-blank: {int(_ppa.wing_borne_units.notna().sum())}; blank because: {_ppa[_w].wing_borne_units_note.dropna().str.replace(r'^(HB|PFV):.*', 'HB/PFV', regex=True).value_counts().to_dict()}")
log(f"    wing_boom_candidate: {int(_ppa.wing_boom_candidate.fillna(False).sum())} to re-read under rule 3")
""")

md("## 8 — Order: company A→Z → named aircraft A→Z, then unnamed by patent id, then rejected → a D1/D2 of the same company right under its original")
code(r"""
# Ruling 2026-09-17 (user): ONE block per company, inside a company per aircraft, alphabetical.
#   company  : `company` (reference/family_map.csv company_canonical), A→Z case-insensitive; the two pseudo-companies
#              "Individual Inventor" and "Unknown / Independent" go last
#   patent   : placed by its best aircraft — tier 0 approved with a real name (A→Z by name), tier 1 approved id-named
#              (by patent id), tier 2 rejected (by patent id); a patent's aircraft stay together in ua order
#   D1 / D2  : directly under the original they point at when that original is in the SAME company; otherwise they are
#              filed in their own company as a standalone entry (their name is the original's name, so they still sort
#              with the aircraft they repeat by name). D3 are different aircraft: they sort on their own.
M["company"] = M.patent_id.map(comp_of)
PSEUDO_COMPANIES = {"Individual Inventor", "Unknown / Independent"}
def block_of(pid):
    if dup_type.get(pid) in ("1", "2"):
        root, ok = chain_root(pid)
        if ok and root in var_ids and comp_of.get(root) == comp_of.get(pid): return root
    return pid
M["_block"] = M.patent_id.map(block_of)
_appr = M.is_approved.fillna(False).astype(bool)
M["_tier"] = np.where(_appr & M.name_is_real.fillna(False).astype(bool), 0, np.where(_appr, 1, 2))
M["_nkey"] = np.where(M._tier == 0, M.aircraft_name.astype(str).str.lower(), M.patent_id.astype(str))
_pk = {pid: min(zip(g._tier, g._nkey)) for pid, g in M.groupby("patent_id")}      # a patent is placed by its best aircraft
M["_ptier"] = M._block.map(lambda b: _pk[b][0]); M["_pname"] = M._block.map(lambda b: _pk[b][1])
M["_c0"] = M.company.isin(PSEUDO_COMPANIES).astype(int); M["_c1"] = M.company.astype(str).str.lower()
M["_self"] = (M.patent_id != M._block).astype(int)
M["_dt"]   = M.dup_type.fillna("0")
M = M.sort_values(["_c0", "_c1", "_ptier", "_pname", "_block", "_self", "_dt", "patent_id", "ua"]).reset_index(drop=True)
M = M.drop(columns=["_block", "_tier", "_nkey", "_ptier", "_pname", "_c0", "_c1", "_self", "_dt"])
_runs = int((M.company != M.company.shift()).sum())
log(f"order: {M.company.nunique()} companies in {_runs} runs (must be equal); first: {M.company.iloc[0]} / {M.aircraft_name.iloc[0]}")
assert _runs == M.company.nunique(), "a company is split into more than one run"
""")

md("## 9 — Figures: one row per figure file of a non-D2 record; counts and the main figure per aircraft")
code(r"""
T2 = L[L.Section == "T2"].copy()
T2["block"] = T2.Sub_Dimension.astype(str)
placeholder = T2.block.str.contains(r"\(fig", regex=True)
n_placeholder = T2[placeholder & (T2.Field == "status")].groupby("base").size()
T2 = T2[~placeholder]
F = T2.drop_duplicates(["base", "block", "Field"], keep="last").pivot(index=["base", "block"], columns="Field", values="Value").reset_index()
F.columns.name = None
ip = T2.dropna(subset=["Image_Path"]).drop_duplicates(["base", "block"]).set_index(["base", "block"]).Image_Path
F["image_path"] = [ip.get((b, k), np.nan) for b, k in zip(F.base, F.block)]
F = F.rename(columns={"base": "patent_id", "figKey": "fig_key", "isMain": "is_main"})
F["batch"] = F.patent_id.map(batch_of); F["company"] = F.patent_id.map(comp_of)
F["image_file"] = F.image_path.map(lambda p: Path(str(p)).name if not is_empty(p) else np.nan)
F["file_exists"] = F.image_path.map(lambda p: (not is_empty(p)) and Path(str(p)).exists())
for c in ["per", "acSty", "acCol", "bgSty", "bgCol", "acState", "qualityFlag", "parts"]:
    if c in F.columns: F[c] = F[c].map(lambda v, m=(c == "parts"): norm_cell(v, None, m))
if "parts" in F.columns:
    okp = set(LISTS["T2_PARTS_DEFAULT"])
    for i in F.index[F.parts.notna()]:
        ps = str(F.at[i, "parts"]).split("|"); badp = [p for p in ps if p not in okp]
        if badp:
            F.at[i, "parts"] = "|".join([p for p in ps if p in okp] + ["Other"])
            F.at[i, "comment"] = ("" if is_empty(F.at[i, "comment"]) else str(F.at[i, "comment"]) + " | ") + "parts: " + "; ".join(badp)
            log(f"  parts recode {F.at[i,'patent_id']} {F.at[i,'fig_key']}: {badp} -> Other (text on comment)")
F["arch"] = pd.to_numeric(F.get("arch"), errors="coerce").astype("Int64")
F["ua"] = F.arch.fillna(1).astype(int)
F["aircraft_id"] = F.patent_id + "_ua" + F.ua.astype(str)
for c in ["is_main", "hasLegends"]:
    if c in F.columns: F[c] = F[c].map(lambda v: np.nan if is_empty(v) else str(v) == "True").astype("boolean")
F["rotation_deg"] = F.rotation_deg.map(lambda v: 0 if str(v) == "False" else v).pipe(pd.to_numeric, errors="coerce").astype("Int64") if "rotation_deg" in F.columns else np.nan
F = F.rename(columns={"edgeTags": "fig_tags"})

# ── per-figure SigLIP pre-labels from the ML feed, keyed on (patent, image file name) ──
t2m = ML[ML.Section == "T2"].copy()
t2m["image_file"] = [Path(str(p)).name if not is_empty(p) else re.sub(r"^Image:\s*", "", str(s)).strip()
                     for p, s in zip(t2m.Image_Path, t2m.Sub_Dimension)]
t2m = t2m[t2m.image_file != ""]
Vf, Cf = ml_pivot(t2m, "Value", ["Patent_ID", "image_file"]), ml_pivot(t2m, "Confidence", ["Patent_ID", "image_file"])
ML_FIG = ["per", "acSty", "acCol", "bgSty", "bgCol", "acState", "parts", "tiltedInView", "match_status", "dupOfPatent", "dupOfFig"]
ML_FIG_CODED = {"per", "acSty", "acCol", "bgSty", "bgCol", "acState", "parts"}
fkeys = list(zip(F.patent_id, F.image_file))
ml_fig_cols = []
for f in ML_FIG:
    if f not in Vf.columns: continue
    F["ml_" + f] = [clean(Vf[f].get(k, np.nan), f in ML_FIG_CODED) for k in fkeys]; ml_fig_cols.append("ml_" + f)
    if f in ML_FIG_CODED and f in Cf.columns:
        F["ml_" + f + "_conf"] = [conf(Cf[f].get(k, np.nan)) for k in fkeys]; ml_fig_cols.append("ml_" + f + "_conf")
# the tables carry what the wizard SHOWS: its "Invariant" option is stored under the legacy id 'HoverCruise' (kept so
# saved records load). The Review sheet keeps the id — the wizard reads that sheet back.
for c in ("acState", "ml_acState"):
    if c in F.columns: F[c] = F[c].replace({INVARIANT_ID: INVARIANT_NAME})
for f in [c for c in ("per", "acSty", "acCol", "bgSty", "bgCol", "acState") if c in F.columns and "ml_" + c in F.columns]:
    both = F[f].notna() & F["ml_" + f].notna()
    F["ml_" + f + "_agrees"] = pd.array(np.where(both, F[f].astype(str).values == F["ml_" + f].astype(str).values, None), dtype="boolean")
    ml_fig_cols.append("ml_" + f + "_agrees")
FIG_COLS = ["aircraft_id", "patent_id", "ua", "batch", "company", "fig_key", "block", "image_file", "image_path", "file_exists", "status", "arch", "is_main",
            "per", "acSty", "acCol", "bgSty", "bgCol", "parts", "qualityFlag", "acState", "stateNote", "hasLegends",
            "dupOf", "comment", "fig_tags", "rotation_deg"] + ml_fig_cols
F = F[[c for c in FIG_COLS if c in F.columns]]
F["is_primary"] = ~F.patent_id.isin(D1_SET | D2_SET)
F["approved_copy_path"] = pd.Series([None] * len(F), index=F.index, dtype=object)

# counts per patent BEFORE the D2 figures are set aside (the record does hold their own crops)
# a REJECTED patent keeps no figures — the wizard exports only its verdict and reason (2026-09-18: CN113086184A and
# WO2023272353A1 still carried 14 figure rows and 2 copied images)
_rej_p = set(M[~M.is_approved.fillna(False).astype(bool)].patent_id)
_n_rej_fig = int(F.patent_id.isin(_rej_p).sum())
F = F[~F.patent_id.isin(_rej_p)].reset_index(drop=True)
log(f"figures of rejected patents dropped: {_n_rej_fig}")
F_all = F.copy()
appr_all = F_all[F_all.status == "approved"]
M["n_figures"] = M.patent_id.map(F_all.groupby("patent_id").size()).fillna(0).astype(int)
M["n_approved"] = M.patent_id.map(appr_all.groupby("patent_id").size()).fillna(0).astype(int)
M["n_fig_placeholders"] = M.patent_id.map(n_placeholder).fillna(0).astype(int)
n_d2_fig = int(F.patent_id.isin(D2_SET).sum()); n_d2_appr = int(appr_all.patent_id.isin(D2_SET).sum())
F = F[~F.patent_id.isin(D2_SET)].reset_index(drop=True)      # a D2 has the same figures as its original: it points
log(f"figures: {len(F_all)} with a file reference ({int(F_all.file_exists.sum())} files present, {len(appr_all)} approved); "
    f"{int(n_placeholder.sum())} '(fig N)' placeholders dropped; {n_d2_fig} figures of D2 records set aside ({n_d2_appr} approved) — they point at the original's")

# rule 2026-09-19 (codebook R2-02(e)), two flight states in one figure: the state drawn in solid lines is recorded, the
# dashed one ignored; both drawn with the same weight = Both (state added 2026-09-19); nothing readable = Other; lift units beside cruise units are not two states;
# Invariant never for a figure showing two positions. acState passes through unchanged — a stateNote that names two
# states / both modes is a WARNING for the reviewer (STATE_NOTE_MENTIONS_TWO_STATES, §13).
TWO_STATE_RE = re.compile(r"\bboth\b[^.;]*\b(?:modes?|states?|configurations?|positions?)\b|\b(?:two|2)\s+(?:flight\s+)?(?:modes|states|positions|configurations)\b"
                          r"|\b(?:hover|vtol|take-?off)\b[^.;]*\band\b[^.;]*\bcruise\b|\bcruise\b[^.;]*\band\b[^.;]*\b(?:hover|vtol)\b|\b(?:dashed|dotted|phantom)\b", re.I)
TWO_STATE_FIGS = F[F.stateNote.map(lambda v: not is_empty(v) and bool(TWO_STATE_RE.search(str(v))))] if "stateNote" in F.columns else F.iloc[0:0]
log(f"  figure stateNotes that name two states (R2-02(e) warning): {len(TWO_STATE_FIGS)}  acState {TWO_STATE_FIGS.acState.value_counts(dropna=False).to_dict()}")

appr = F[F.status == "approved"]
byv = appr.groupby(["patent_id", "ua"]).size()
mainfig = appr[appr.is_main == True].drop_duplicates(["patent_id", "ua"]).set_index(["patent_id", "ua"]).image_path
M["n_approved_this_variant"] = [int(byv.get((p, v), 0)) for p, v in zip(M.patent_id, M.ua)]
M["main_figure"] = [mainfig.get((p, v), np.nan) for p, v in zip(M.patent_id, M.ua)]
# a pointing row takes the figure count and main figure of the aircraft it points at (first target)
for i in M.index[M.same_aircraft_as.notna()]:
    tp, tu = str(M.at[i, "same_aircraft_as"]).split("; ")[0].rsplit("_ua", 1); tu = int(tu)
    M.at[i, "n_approved_this_variant"] = int(byv.get((tp, tu), 0))
    if is_empty(M.at[i, "main_figure"]): M.at[i, "main_figure"] = mainfig.get((tp, tu), np.nan)
""")

md("## 10 — Disapproval reason: the human/SBERT precedence switch and the disagreement report")
code(r"""
M["reason_sbert"] = np.nan; M["sub_reason"] = np.nan
sbert = None
if SBERT_XLSX and Path(SBERT_XLSX).exists():
    sbert = pd.read_excel(SBERT_XLSX)
    sbert.columns = [str(c).strip().lower() for c in sbert.columns]
    need = {"patent_id", "reason"}
    assert need <= set(sbert.columns), f"SBERT workbook needs columns {need} (+ optional sub_reason, confidence); has {list(sbert.columns)}"
    sbert = sbert.drop_duplicates("patent_id").set_index("patent_id")
    M["reason_sbert"] = M.patent_id.map(sbert["reason"])
    if "sub_reason" in sbert.columns: M["sub_reason"] = M.patent_id.map(sbert["sub_reason"])
    dis = M[~M.is_approved.fillna(False).astype(bool)].drop_duplicates("patent_id")[["batch", "patent_id", "aircraft_name", "reason_human", "reason_note", "reason_sbert", "sub_reason"]]
    dis["agree"] = dis.reason_human.fillna("").str.lower() == dis.reason_sbert.fillna("").str.lower()
    TAB_DIR.mkdir(parents=True, exist_ok=True)
    dis.to_csv(TAB_DIR / "reason_disagreement_report.csv", index=False)
    log(f"SBERT: {len(sbert)} rows; rejected patents compared {len(dis)}, agree {int(dis.agree.sum())}, disagree {int((~dis.agree).sum())} -> tables/reason_disagreement_report.csv")
elif REASON_SOURCE != "human":
    log(f"  !! REASON_SOURCE={REASON_SOURCE} but no SBERT workbook at {SBERT_XLSX!r} — falling back to human"); REASON_SOURCE = "human"

if REASON_SOURCE == "human":
    M["reason"] = M.reason_human; M["reason_source"] = np.where(M.reason_human.notna(), "human", None)
elif REASON_SOURCE == "sbert":
    M["reason"] = M.reason_sbert; M["reason_source"] = np.where(M.reason_sbert.notna(), "sbert", None)
else:
    use_h = M.reason_human.notna()
    M["reason"] = np.where(use_h, M.reason_human, M.reason_sbert)
    M["reason_source"] = np.where(use_h, "human", np.where(M.reason_sbert.notna(), "sbert", None))
log(f"reason column filled from: {M.reason_source.value_counts(dropna=False).to_dict()}")
""")

md(r"""
## 11 — Write the outputs

`0_labelling/outputs/` is wiped of stale files and rebuilt: `tables/` (what the analysis stages read),
`images/<aircraft_id>/` (approved figures of every non-D2 record, manifest with sha256, stale copies pruned),
`1639_LABELS.xlsx` (README · Review · ground_truth · patseer), then BUILD_LOG.md and ACTION_FLAGS.csv in sections 12–13.
""")
code(r"""
def expand(prefix_re, order, cols):
    groups = sorted({int(m.group(1)) for c in cols for m in [re.match(prefix_re, c)] if m})
    out = []
    for g in groups:
        pre = re.match(prefix_re, next(c for c in cols if re.match(prefix_re, c) and int(re.match(prefix_re, c).group(1)) == g)).group(0)
        out += [pre + f for f in order if pre + f in cols]
    return out
def m3_block(cols):
    out = []
    for card in M3_CARDS:
        out += [f"{card}_{f}" for f in M3_CARD if f"{card}_{f}" in cols]
        types = sorted({int(m.group(1)) for c in cols for m in [re.match(rf"^{card}_t(\d+)_", c)] if m})
        for t in types: out += [f"{card}_t{t}_{f}" for f in M3_TYPE if f"{card}_t{t}_{f}" in cols]
    return out
cols = set(M.columns)
HEAD = ["aircraft_id", "patent_id", "ua", "n_variants", "batch", "company", "aircraft_name", "name_is_real", "aircraft_name_wizard",
        "variant", "variant_id",
        "assignee", "title", "app_year", "pub_year", "pdf_link",
        "is_approved", "reason", "reason_source", "sub_reason", "reason_human", "reason_sbert", "reason_note",
        "is_duplicate", "dup_type", "dup_of", "dup_root", "is_primary", "same_aircraft_as", "points_to_ua", "labels_inherited_from", "dup_of_missing",
        "n_figures", "n_approved", "n_approved_this_variant", "n_fig_placeholders", "main_figure", "img_not_reflect",
        "t1_humanUncertain", "t1_uncertainNote", "t1_quickOverride", "t1_quickNote"]
g1 = [c for c in G1_ORDER if c in cols]
m1 = [c for c in M1_FIXED if c in cols] + (["fusShape_otherTag"] if "fusShape_otherTag" in cols else [])
m1 += expand(r"^boom(\d+)_", M1_BOOM + [f"{x}_otherTag" for x in ("attach", "wingRel", "long", "orient")], cols)
m1 += [c for c in M1_TAIL if c in cols]
m2 = [c for c in M2_FIXED if c in cols] + [c for c in ("wingConf_otherTag", "empType_otherTag") if c in cols]
m2 += expand(r"^wing(\d+)_", M2_WING + ["role_otherTag", "plan_otherTag"], cols) + [c for c in M2_TAIL if c in cols]
m3 = m3_block(cols) + sorted(c for c in cols if c.endswith("propKin_otherTag")) + [c for c in M3_TAIL if c in cols]   # sorted: a set's order changes with the hash seed
ORDER = HEAD + g1 + m1 + m2 + m3 + [c for c in ml_head if c in cols]
def insert_after(order, anchor, items):
    items = [i for i in items if i in cols and i not in order]
    if anchor in order: i = order.index(anchor) + 1; order[i:i] = items
    else: order += items
for f in ML_MORPH:
    insert_after(ORDER, f, [f"ml_{f}", f"ml_{f}_conf", f"ml_{f}_source", f"ml_{f}_agrees"])
def insert_before(order, anchor, items):
    items = [i for i in items if i in cols and i not in order]
    if anchor in order: i = order.index(anchor); order[i:i] = items
    else: order += items
# 7e rule columns beside what they read: overrides after the G1 block, the unit totals ahead of the M3 cards, the
# wing-thrust columns ahead of the wing-1 card, the boom state ahead of the boom card
insert_after(ORDER, "g1_quickNote", ["overrides"])
if m3: insert_before(ORDER, m3[0], ["propulsor_units", "propulsor_units_note"])
insert_before(ORDER, "wing1_count", ["wing_thrust_carrier", "wing_borne_units", "wing_borne_units_note", "wing_boom_candidate"])
insert_before(ORDER, "boom_count", ["boom_thrust_state"])
leftover =[c for c in M.columns if c not in ORDER and c != "n_variants_declared"]
if leftover: log(f"  columns not in the template (appended at the end): {leftover}")
M = M[ORDER + leftover]
log(f"aircraft table: {M.shape[0]} rows x {M.shape[1]} columns")

# ── data dictionary (tables only) ──
dd = []
for c in M.columns:
    sec = sec_of.get(c); lst = list_for(sec, c) if sec else None
    kind = ("tag" if c.endswith("_otherTag") else "id" if lst else "bool" if str(M[c].dtype) == "boolean" else
            "int" if str(M[c].dtype) == "Int64" else "text")
    opts = " | ".join(f"{i}={LABELS[lst].get(i, i)}" if LABELS.get(lst) else i for i in LISTS[lst]) if lst else ""
    doc_opts, desc = RULE_DOC.get(c, (None, ""))
    dd.append(dict(column=c, section=sec or ("derived" if c in HEAD or c in RULE_COLS else ""), kind=kind,
                   options=opts if doc_opts is None else doc_opts, non_empty=int(M[c].notna().sum()), description=desc))
DD = pd.DataFrame(dd)

# ── outputs folder: only what this run writes ──
OUT_DIR.mkdir(parents=True, exist_ok=True); TAB_DIR.mkdir(exist_ok=True)
KEEP_TOP = {"1639_LABELS.xlsx", "images", "tables", "BUILD_LOG.md", "ACTION_FLAGS.csv", "README.md"}
for p in OUT_DIR.iterdir():
    if p.name not in KEEP_TOP:
        (shutil.rmtree if p.is_dir() else os.remove)(p); log(f"  removed stale output {p.name}")

# ── approved images: one folder per aircraft, sha256 manifest, stale copies pruned ──
if COPY_IMAGES:
    root = OUT_DIR / "images"; man = []; missing = 0; copied = 0
    for i, r in F[(F.status == "approved")].iterrows():
        if not r.file_exists: missing += 1; continue
        dst = root / r.aircraft_id / r.image_file
        dst.parent.mkdir(parents=True, exist_ok=True)
        src = Path(r.image_path)
        if not dst.exists() or dst.stat().st_size != src.stat().st_size:
            shutil.copy2(src, dst); copied += 1
        F.at[i, "approved_copy_path"] = str(dst)
        man.append(dict(aircraft_id=r.aircraft_id, patent_id=r.patent_id, company=r.company, ua=r.ua, fig_key=r.fig_key, figure_file=r.image_file,
                        src=str(src), dst=str(dst), bytes=dst.stat().st_size, sha256=hashlib.sha256(dst.read_bytes()).hexdigest()))
    MAN = pd.DataFrame(man); MAN.to_csv(root / "MANIFEST.csv", index=False)
    keep = set(MAN.dst) | {str(root / "MANIFEST.csv")}; pruned = 0
    for p in list(root.rglob("*")):
        if p.is_file() and str(p) not in keep: p.unlink(); pruned += 1
    for dd_ in sorted([x for x in root.rglob("*") if x.is_dir()], key=lambda x: -len(str(x))):
        if not any(dd_.iterdir()): dd_.rmdir()
    log(f"images: {len(man)} approved figures in {len(MAN.aircraft_id.unique())} aircraft folders ({copied} newly copied, {pruned} stale files pruned, {missing} approved figures whose file is missing on disk)")

# ── tables ──
# ── names the wizard shows instead of four legacy ids (user ruling 2026-09-18: ONLY these; TiltBody and
# "No Aircraft Image" stay as they are). dup_type stays the code 1/2/3 — notebooks 10/20/23 compute with it —
# and dup_type_name carries D1/D2/D3 beside it.
DISPLAY = {"gearArch": {"FixedWheel": "Wheeled Gear"}, "fusKin": {"LevelCabin": "Variable Incidence"},
           "qualityFlag": {"generic": "Partial Quality"}}
_rmech = [c for c in M.columns if re.search(r"(^|_)rmech$", c)]
for c in list(M.columns):
    base_c = c[3:] if c.startswith("ml_") else c
    m = DISPLAY.get(base_c) or ({"Exposed": "Fixed"} if base_c in _rmech or re.search(r"(^|_)rmech$", base_c) else None)
    if m and M[c].dtype == object: M[c] = M[c].replace(m)
for c in ("qualityFlag", "ml_qualityFlag"):
    if c in F.columns: F[c] = F[c].replace(DISPLAY["qualityFlag"])
M.insert(M.columns.get_loc("dup_type") + 1, "dup_type_name", M.dup_type.map(lambda v: f"D{v}" if str(v) in ("1", "2", "3") else np.nan))
log("display names: FixedWheel→Wheeled Gear, LevelCabin→Variable Incidence, Exposed→Fixed (rmech), generic→Partial Quality; dup_type_name D1/D2/D3")
M.to_csv(TAB_DIR / "aircraft_table.csv", index=False)
F.to_csv(TAB_DIR / "figure_table.csv", index=False)
DD.to_csv(TAB_DIR / "data_dictionary.csv", index=False)
log(f"tables: aircraft_table.csv {M.shape}, figure_table.csv {F.shape}, data_dictionary.csv -> {TAB_DIR}")

# ── the workbook ──
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
used(PATSEER, "PatSeer export")
PS_COLS = [("Record Number", "patent_id"), ("Title", "title"), ("Record Type", "record_type"), ("Assignee", "assignee"), ("Assignee Country", "assignee_country"),
           ("Current Assignee", "current_assignee"), ("Applicant", "applicant"), ("Inventors", "inventors"), ("Application No.", "application_no"),
           ("Filing/Application Date", "filing_date"), ("Publication/Issue Date", "publication_date"), ("Priority Date (Record)", "priority_date"),
           ("Priority Country Code", "priority_country"), ("Publication Country", "publication_country"), ("CPC Main Group (CPCG)", "cpc_main_group"),
           ("Simple Family ID", "simple_family_id"), ("No. of Simple Family Members", "simple_family_members"), ("Extended Family ID", "extended_family_id"),
           ("Backward Citation Count", "backward_citations"), ("No. of Forward Citations (Individual)", "forward_citations"), ("Legal Status (Dead/Alive)", "legal_status")]
PS = pd.read_excel(PATSEER, usecols=[a for a, _ in PS_COLS], dtype=str).rename(columns=dict(PS_COLS))
PS["patent_id"] = PS.patent_id.str.strip(); PS = PS.drop_duplicates("patent_id").set_index("patent_id")
order = M.drop_duplicates("patent_id").patent_id.tolist()
PS = PS.reindex(order).reset_index(); PS.insert(1, "batch", PS.patent_id.map(batch_of)); PS.insert(2, "company", PS.patent_id.map(comp_of))
log(f"patseer sheet: {int(PS.title.notna().sum())} / {len(PS)} patents found in the export")

_gsel = M[M.is_approved.fillna(False).astype(bool) & M.is_primary]
GTS = _gsel[["aircraft_id", "patent_id", "company", "aircraft_name", "topType", "arch_gt", "arch_gt_visible"]].rename(columns={"topType": "image_label"})
# flag column (user 2026-09-17): what a reader must know before using the row as ground truth
def _gt_flag(r):
    out = []
    if is_empty(r.arch_gt) and str(r.arch_gt_provenance) in NOT_STATED_PROV: out.append("GT not stated in the patent")
    elif is_empty(r.arch_gt): out.append("GT missing: not read yet")
    elif str(r.arch_gt_provenance) == "unsure": out.append("GT unsure")
    if str(r.g1_humanUncertain) == "True": out.append("image label marked uncertain in the wizard")
    return "; ".join(out)
GTS["flag"] = [_gt_flag(r) for r in _gsel.itertuples()]

# names per patent, in ua order, as ONE string for the wizard's aircraftName field ("a; b" for a multi-aircraft patent)
approved_pids = set(M[M.is_approved.fillna(False).astype(bool)].patent_id)
NAME_STR = {pid: "; ".join(g.sort_values("ua").aircraft_name.tolist()) for pid, g in M[M.patent_id.isin(approved_pids)].groupby("patent_id")}
has_name_row = set(L[(L.Section == "T1") & (L.Field == "aircraftName")].Patent_ID)

wb = Workbook(); ws0 = wb.active; ws0.title = "README"
readme = [("workbook", WORKBOOK.name), ("built", RUN_TS), ("built by", "Patent-Labelling-Tools/notebooks/04_master_labels.ipynb (generator scripts/gen_04_master_labels.py)"),
          ("wizard", HTML.name), ("", "")]
readme += [("sheet Review", "the wizard record exactly as exported (Patent_ID, Section, Sub_Dimension, Field, Value, Source, Image_Path), with the aircraft names written into T1/aircraftName. Load it in the wizard with Resume (it picks the sheet named Review)."),
           ("sheet ground_truth", "one row per primary approved aircraft: the wizard type (image_label) next to the whole-document ground truth (arch_gt) and whether that type is visible in the figures; column flag = GT not stated in the patent / GT missing / GT unsure / image label marked uncertain in the wizard"),
           ("sheet patseer", "bibliographic facts from the PatSeer export, one row per patent, joined on patent_id"),
           ("aircraft id", "<patent>_ua<N>: N = the aircraft number in the wizard (a single-aircraft patent is _ua1). Names: real name (+ vN) where reviewed, else the id."),
           ("order", "every sheet, table and the record itself: ONE block per company, companies A→Z (column company = reference/family_map.csv company_canonical, e.g. Bell Helicopter Textron / Bell Textron / Textron Innovations = Bell / Textron; Individual Inventor and Unknown / Independent last); inside a company the named aircraft A→Z, then the unnamed by patent id, then the rejected records; a D1/D2 of the same company directly under its original; a patent's aircraft together in ua order"),
           ("flight state", "acState per figure. The wizard SHOWS 'Invariant' for the option it stores under the old id HoverCruise: tables/figure_table.csv writes Invariant; the Review sheet keeps 'HoverCruise — Invariant' because the wizard reads that sheet back. Every lock the wizard draws is checked against the record on each build (notebook 04 §6b). Two flight states in one figure (codebook R2-02(e), ruling 2026-09-19): when a moving part is drawn in two positions, the state drawn in solid lines is recorded and the dashed one is ignored; when both are drawn with the same weight, the figure is Both (a state added the same day; Other is only 'no configuration can be read'). Lift units drawn beside cruise units do not make a two-state figure. Invariant is never used for a figure showing two positions. acState passes through unchanged; a stateNote that names two states is flagged STATE_NOTE_MENTIONS_TWO_STATES (warning)."),
           ("overrides", "ruling 2026-09-19: an override keeps only what it records. After a G1 override the aircraft has no type; after an M1 or M2 override the skipped fields are not determinable; at an M3 station only the propulsor count is kept (the wizard stores <station>_count=0 and the real number in <station>_quickCount). A value that is not determinable is never read as zero, none or Fixed: it is left out, and the number left out is reported. aircraft_table: overrides (G1|M1|M2|M3:<station>), propulsor_units (non-control-only propulsors over every M3 station, quick count on a count-only station; blank for HB/PFV, after a G1 override, or when an overridden station has no quick count) and propulsor_units_note (why blank). Flags: OVERRIDE_HIDES_VALUE (a value stored behind an override), LEFTOVER_QUICK_COUNT (a quick count on a station whose override is off — ignored)."),
           ("tilting booms", "ruling 2026-09-19: the boom tick carries the tilt. A rotor on a boom group ticked Booms tilt is recorded Fixed, as on a tilting wing. Tick Booms tilt only when the boom itself rotates; otherwise a rotor that changes angle on a boom is recorded Tilt. A rotor that also pivots on a tilting boom is recorded Other, with a note. aircraft_table: boom_thrust_state = none / fixed / tilting / on_tilting_wing / mixed (a group rides a tilting wing when attached to Wings/Both, the wing it names — any when blank/Multi — is set to Tilt and onFixedPart is not True; it tilts when Booms tilt is ticked or its rotors are propKin Tilt; groups with hasProps False ignored). Flags: BOOM_TICK_AND_ROTOR_TILT, BOOM_ROTORS_WITHOUT_GROUP."),
           ("wing-borne thrust", "ruling 2026-09-19: a member on a wing that carries a propulsor is a boom when it carries propulsors both ahead of and behind the wing, reaches beyond the leading or trailing edge by more than the local wing chord, or carries a tail surface. Otherwise it is a nacelle or pylon, and the propulsor is recorded on the wing card at LE, TE, Above or Below. Where the analysis asks whether the wing carries the thrust, both count as wing-borne. aircraft_table (winged aircraft): wing_thrust_carrier = wing / wing_boom / both / none, wing_borne_units (+ wing_borne_units_note when the boom rotors cannot be apportioned), wing_boom_candidate (members to re-read under this rule)."),
           ("tilting empennage", "empTilts = True: the empennage tilts, with the wing or on its own."),
           ("duplicates", "D1/D2 carry nothing from G1 to M3 and point at the aircraft they repeat (tables/aircraft_table.csv, column same_aircraft_as); a D2 has no figures of its own"),
           ("images", "images/<aircraft_id>/<figure>.png — every approved figure of a non-D2 record, MANIFEST.csv with sha256"),
           ("tables", "tables/aircraft_table.csv (one row per aircraft, every label + identity + ground truth + the 2026-09-19 rule columns + ml_ pre-labels), figure_table.csv (one row per figure), data_dictionary.csv (every column: section, kind, options; description for the derived rule columns) — what stages 1 and 2 read"),
           ("", "")]
readme += [("input " + what, stmp) for what, _, stmp in INPUTS_USED]
readme += [("", ""), ("rows Review", len(L) + sum(1 for p in approved_pids if p not in has_name_row)), ("rows ground_truth", len(GTS)), ("rows patseer", len(PS)),
           ("rows aircraft_table", len(M)), ("rows figure_table", len(F))]
for k, v in readme: ws0.append([k, v])
ws0.column_dimensions["A"].width = 28; ws0.column_dimensions["B"].width = 140

wsR = wb.create_sheet("Review")
src_ws = load_workbook(RECORD, read_only=True)["Review"]
it = src_ws.iter_rows(values_only=True); hdr = list(next(it)); wsR.append(hdr)
NAME_WRITTEN, NAME_ROWS_ADDED = 0, 0
for row in it:
    row = list(row); pid = row[0]
    if row[1] == "T1" and row[3] == "aircraftName" and pid in NAME_STR:
        if str(row[4] or "") != NAME_STR[pid]: NAME_WRITTEN += 1
        row[4] = NAME_STR[pid]
    wsR.append(row)
    if row[1] == "T1" and row[3] == "isApproved" and pid in NAME_STR and pid not in has_name_row:
        wsR.append([pid, "T1", "Aircraft / Prototype Name", "aircraftName", NAME_STR[pid], None, None]); NAME_ROWS_ADDED += 1
log(f"Review sheet: {wsR.max_row - 1} rows; names written into aircraftName on {NAME_WRITTEN} patents, {NAME_ROWS_ADDED} name rows added")

def df_to_sheet(ws, df):
    ws.append(list(df.columns))
    for rec in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) or v is pd.NA else v for v in rec])
    for i, c in enumerate(df.columns, 1): ws.column_dimensions[get_column_letter(i)].width = min(60, max(12, int(df[c].astype(str).str.len().quantile(0.9)) + 2))
df_to_sheet(wb.create_sheet("ground_truth"), GTS)
df_to_sheet(wb.create_sheet("patseer"), PS)
wb.save(WORKBOOK)
log(f"workbook: {WORKBOOK.name} — sheets {wb.sheetnames}")
""")

md("## 12 — Checks and summary")
code(r"""
checks = []
def chk(name, ok, detail=""): checks.append((name, bool(ok), detail)); log(f"  [{'OK' if ok else 'FAIL'}] {name}  {detail}")
chk("every patent present", M.patent_id.nunique() == PT.shape[0], f"{M.patent_id.nunique()} / {PT.shape[0]}")
chk("the record stores what the wizard shows (every lock replayed, §6b)", WV.empty,
    f"{len(WV)} differences" + ("; fix: python scripts/apply_wizard_locks.py --apply" if len(WV) else ""))
chk("figure_table acState uses the wizard's display name Invariant, never the id HoverCruise", not (F.acState == INVARIANT_ID).any())
chk("no label separator left in a coded column", not any(M[c].astype(str).str.contains(SEP, regex=False).any() for c in coded_cols))
prim_appr = M[M.is_primary & M.is_approved.fillna(False).astype(bool)]
unclass = prim_appr[prim_appr.g1_quickOverride.fillna(False).astype(bool)]
gap = prim_appr[prim_appr.topType.isna() & ~prim_appr.g1_quickOverride.fillna(False).astype(bool)]
chk("approved primary aircraft carry a topType (g1_quickOverride exempt)", gap.empty, f"{len(gap)} without; {len(unclass)} unclassifiable by design: {unclass.patent_id.tolist()}")
rej = M[~M.is_approved.fillna(False).astype(bool)]
chk("rejected rows carry no morphology", rej.topType.isna().all(), f"{len(rej)} rejected rows")
noreason = rej[rej.reason.isna()].patent_id.tolist()
log(f"  [{'OK' if not noreason else 'FINDING'}] rejected rows without a reason: {noreason}")
pt = M[M.same_aircraft_as.notna()]
prim_ids = set(prim_appr.aircraft_id) | set(M[M.is_primary].aircraft_id)
chk("D1/D2 rows carry no G1–M3 and point at existing aircraft", pt[g1m3_cols].isna().all().all() and
    all(t in prim_ids for v in pt.same_aircraft_as for t in str(v).split("; ")), f"{len(pt)} pointing rows")
chk("every approved aircraft has a name", M[M.is_approved.fillna(False).astype(bool)].aircraft_name.map(lambda v: not is_empty(v)).all())
chk("aircraft_id unique", not M.aircraft_id.duplicated().any())
chk("(patent_id, ua) unique", not M.duplicated(["patent_id", "ua"]).any())
_pos = pd.Series(range(len(M)), index=M.index); _first = M.groupby("patent_id").apply(lambda g: _pos[g.index].min())
def _contiguous(groups):
    return {k: bool(_pos[g.index].max() - _pos[g.index].min() + 1 == len(g)) for k, g in groups if len(g)}
_blk = M.patent_id.map(block_of)
_under = M[_blk != M.patent_id].drop_duplicates("patent_id")
_before = [r.patent_id for r in _under.itertuples() if _first[block_of(r.patent_id)] > _first[r.patent_id]]
_bad_blocks = [k for k, ok in {**_contiguous(M.groupby("company")), **_contiguous(M.groupby("patent_id")), **_contiguous(M.groupby(_blk))}.items() if not ok]
chk("one block per company; a patent's aircraft together; a same-company D1/D2 directly under its original", not _before and not _bad_blocks,
    f"{len(_before)} before their original; non-contiguous: {_bad_blocks[:6]}")
chk("ML feed covers every patent", not missing_ml, f"{len(missing_ml)} missing")
chk("no ML value landed in a human column", all(not c.startswith("ml_") for c in variant_cols + [d for _, d in PATENT_COLS]))
chk("every edge tag is a known tag", all(t in LISTS["T1_EDGE_TAGS"] for v in M.edgeTags.dropna() for t in str(v).split("|")))
chk("rejected rows carry no identity answers or tags", rej[["uav_final", "is_electric_final", "takeoff_final", "edgeTags"]].isna().all().all())
_removed = {(p, t) for p, t, _ in TAG_REMOVED}
chk("wizard tags removed only where the identity answer contradicts them",
    all((p, t) in _removed for p, w, e, pr in zip(M.patent_id, M.edgeTags_wizard, M.edgeTags, M.is_primary) if pr for t in tagset(w) - tagset(e)), f"{len(TAG_REMOVED)} removed")
chk("V/STOL rows carry no STOLSimilar tag", not M[M.takeoff_final == "V/STOL"].edgeTags.map(lambda v: STOL_TAG in tagset(v)).any())
chk("no D2 figure in the figure table / images", not F.patent_id.isin(D2_SET).any())
chk("figure_table aircraft_ids exist", F.aircraft_id.isin(set(M.aircraft_id)).all(), f"{int((~F.aircraft_id.isin(set(M.aircraft_id))).sum())} orphan figure rows")
chk("workbook Review sheet holds every record row", wsR.max_row - 1 == len(L) + NAME_ROWS_ADDED)
log(f"  [{'OK' if not GT_MISSING else 'FINDING'}] primary approved aircraft without a ground truth: {len(GT_MISSING)}")
# rules 2026-09-19 (§7e / §9)
_pa12 = M.is_primary & M.is_approved.fillna(False).astype(bool)
chk("rule columns (7e) only on primary approved aircraft", M.loc[~_pa12, RULE_COLS].isna().all().all())
chk("propulsor_units / wing_borne_units blank exactly where a note gives the reason",
    ((M.propulsor_units.isna() == M.propulsor_units_note.notna())[_pa12]).all()
    and ((M.wing_borne_units.isna() == M.wing_borne_units_note.notna())[_pa12]).all())
for _what, _items in [("values stored behind an override (OVERRIDE_HIDES_VALUE)", OVR_HIDDEN),
                      ("quick counts on a station whose override is off, ignored (LEFTOVER_QUICK_COUNT)", LEFTOVER_QC),
                      ("boom-card rotor Tilt on a fully ticked tilting boom (BOOM_TICK_AND_ROTOR_TILT; 0 once migrated)", BOOM_TICK_TILT),
                      ("boom-card rotors without a propeller-carrying boom group (BOOM_ROTORS_WITHOUT_GROUP)", BOOM_NO_GROUP)]:
    log(f"  [{'OK' if not _items else 'FINDING'}] {_what}: {len(_items)}")
log(f"  [{'OK' if TWO_STATE_FIGS.empty else 'WARNING'}] figure stateNotes that name two states (STATE_NOTE_MENTIONS_TWO_STATES, R2-02(e)): {len(TWO_STATE_FIGS)}")
log("")
log(f"rows per batch: {M.groupby('batch').size().to_dict()}")
log(f"primary approved aircraft: {len(prim_appr)}   pointing D1/D2 rows: {len(pt)}   rejected: {len(rej)}")
log(f"topType (primary approved): {prim_appr.topType.value_counts().to_dict()}")
(OUT_DIR / "BUILD_LOG.md").write_text("\n".join(LOG) + "\n")
assert all(ok for _, ok, _ in checks), [n for n, ok, _ in checks if not ok]
print("\nBUILD_LOG.md written; all checks passed.")
""")

md("""
## 13 — ACTION FLAGS — what the inputs still need from you

Everything below is a finding in the record or the review files that this notebook reports but must not fix. Written to
`outputs/ACTION_FLAGS.csv` on every run; a flag disappears only when the input is corrected.
""")
code(r"""
flags = []
def flag(kind, pid, detail, action): flags.append(dict(flag=kind, patent_id=pid, detail=detail, action=action))
for pid in rej[rej.reason.isna()].patent_id.unique():
    flag("REJECTED_NO_REASON", pid, "isApproved=False and no t1DisapproveReason", "open in the wizard, pick a disapproval reason, re-export")
for r in REJECTED_WITH_MORPH.itertuples():
    flag("REJECTED_CARRIES_MORPHOLOGY", r.patent_id, f"topType={r.topType}, reason={r.reason_human}",
         "record is rejected but still labelled — blanked here; decide whether the rejection or the labels are right")
for r in M[M.dup_of_missing].drop_duplicates("patent_id").itertuples():
    flag("DUP_ROOT_MISSING", r.patent_id, f"dup_type={r.dup_type}, dup_of={r.dup_of}", "the original carries no labels; fix the link or label the record")
for r in F[(F.status == "approved") & ~F.file_exists].itertuples():
    flag("APPROVED_FIGURE_FILE_MISSING", r.patent_id, f"fig {r.fig_key}: {r.image_path}", "restore the crop or disapprove the figure")
for (nf, k), n in unmapped.items():
    flag("OTHER_NOTE_UNTAGGED", "", f"{nf}: {k[:80]!r} ({n}x)", "add a row to notebooks/04_other_note_tags.csv")
for a in GT_UNSURE:
    flag("GROUND_TRUTH_NOT_STATED", a, "not stated in the patent: the whole-document reading does not settle the type (no arch_final)", "informational — flagged in the ground_truth sheet; tick G1 Uncertain in the wizard if the figure label is uncertain too")
for b in SPLIT_STALE:
    flag("SPLIT_COPY_STALE", b, "labels/reviewed_patents_<batch>.xlsx differs from the record", "python scripts/split_wizard_all_export.py <record> --apply")
_rec_order = L.drop_duplicates("base").base.tolist(); _canon = M.drop_duplicates("patent_id").patent_id.tolist()
if _rec_order != _canon:
    flag("RECORD_ORDER_STALE", "", f"the record is not in the corpus order ({sum(a != b for a, b in zip(_rec_order, _canon))} patents out of place; first expected {_canon[0]}, found {_rec_order[0]})", "python scripts/canonicalize_record.py --apply  (reorders record + feed, writes the names, re-splits the batch copies)")
if NAME_WRITTEN or NAME_ROWS_ADDED:
    flag("RECORD_NAMES_STALE", "", f"{NAME_WRITTEN} names differ / {NAME_ROWS_ADDED} name rows missing in the record (the Review sheet has them)", "python scripts/canonicalize_record.py --apply")
d1 = pd.DataFrame(d1_diff)
if len(d1):
    for pid, g in d1.groupby("patent_id"):
        flag("DUP_BLOCK_IGNORED", pid, f"{len(g)} G1–M3 cells inside this {'D' + str(g.dup_type.iloc[0])} record (points at {g.root.iloc[0]}): {', '.join(g.column.tolist()[:6])}",
             "informational — a D1/D2 points, its own block is ignored (ruling 2026-09-17)")
for c in ID_CONFLICTS:
    flag("IDENTITY_DUP_CONFLICT", c["root"], f"{c['field']} (aircraft {c['variant']}): {c['answers']}",
         "reviewed duplicates of one aircraft answered differently — each keeps its own answer; settle it in the identity review")
for pid, tag, why in TAG_REMOVED:
    flag("WIZARD_TAG_REMOVED", pid, f"wizard edgeTags had {tag} but {why}", "removed here (identity review wins); drop it in the wizard too, or add it to WIZARD_TAG_KEEP")
for pid, tag, why, note in TAG_KEPT:
    flag("WIZARD_TAG_KEPT", pid, f"{tag} kept although {why}", f"informational — {note}")
for pid in unclass.patent_id:
    flag("UNCLASSIFIABLE_BY_DESIGN", pid, "g1_quickOverride=True, no topType", "informational — excluded from topType counts")
for r in PENDING.itertuples():
    n = int((WV_PENDING.patent_id == r.patent_id).sum())
    flag("WIZARD_REVIEW_PENDING", r.patent_id, r.what_to_check + (f" ({n} screen-vs-record differences until re-saved)" if n else ""),
         "open it in the wizard, fix and save; then remove its row from review_decisions/PENDING_WIZARD_REVIEW.csv")
for r in PTC_MOVES.itertuples():
    flag("PTC_RECORD_MOVES", r.patent_id, f"arch{r.arch} is PTC but the record says a part moves ({r.moves}), so the wizard unlocks its figure state",
         "PTC means nothing articulates: fix the type or the moving part in the wizard; the figures then lock to Invariant")
for a in GT_MISSING:
    flag("GROUND_TRUTH_MISSING", a, "primary approved aircraft with no row in architecture_text_final_variants.csv", "decide it on the 03b page and re-apply")
for a in GT_ORPHANS:
    flag("GROUND_TRUTH_ORPHAN", a, "ground-truth row that matches no primary approved aircraft", "informational — the aircraft was merged, renumbered or became a duplicate")
# rules 2026-09-19 (§7e / §9)
for pid, aid, o, hid in OVR_HIDDEN:
    flag("OVERRIDE_HIDES_VALUE", pid, f"{aid}: {o} override, but the record still stores {', '.join(hid)} — the wizard does not show it and the tables do not use it",
         "an override keeps only what it records: if these are real answers untick the override in the wizard, otherwise clear them; re-export")
for pid, aid, st, q in LEFTOVER_QC:
    flag("LEFTOVER_QUICK_COUNT", pid, f"{aid}: {st}_quickCount={q} but the {st} override is off — ignored, the card's own count is used",
         "clear the quick count in the wizard (or tick the override if the station really is count-only); re-export")
for pid, aid, tt, n in BOOM_TICK_TILT:
    flag("BOOM_TICK_AND_ROTOR_TILT", pid, f"{aid} ({tt}): every propeller-carrying boom group ({n}) is ticked Booms tilt and a boom-card rotor is also propKin Tilt",
         "the boom tick carries the tilt: record the rotor Fixed (Other with a note if it also pivots on the boom), or untick Booms tilt if the boom itself does not rotate")
for pid, aid, n in BOOM_NO_GROUP:
    flag("BOOM_ROTORS_WITHOUT_GROUP", pid, f"{aid}: {n} propulsors on the boom card but no boom group carries propulsors (none recorded, or every group hasProps=False)",
         "add the boom group in M1 (or set its hasProps), or move the propulsors to the card that carries them")
for r in TWO_STATE_FIGS.itertuples():
    flag("STATE_NOTE_MENTIONS_TWO_STATES", r.patent_id, f"{r.aircraft_id} fig {r.fig_key} ({r.status}): acState={r.acState}, stateNote={str(r.stateNote)!r}",
         "WARNING — R2-02(e): record the state drawn in solid lines (ignore the dashed one); both drawn with the same weight = Both; nothing readable = Other; never Invariant; lift units beside cruise units are not two states")
AF =pd.DataFrame(flags, columns=["flag", "patent_id", "detail", "action"])
AF.to_csv(OUT_DIR / "ACTION_FLAGS.csv", index=False)
INFO = ["DUP_BLOCK_IGNORED", "UNCLASSIFIABLE_BY_DESIGN", "WIZARD_TAG_KEPT", "GROUND_TRUTH_ORPHAN"]
pd.set_option("display.width", 220); pd.set_option("display.max_colwidth", 90); pd.set_option("display.max_rows", 200)
print(f"{len(AF)} flags -> {OUT_DIR / 'ACTION_FLAGS.csv'}\n")
print(AF.groupby("flag").size().to_string(), "\n")
print(AF[~AF.flag.isin(INFO)].to_string(index=False))
print("\ninformational:")
print(AF[AF.flag.isin(INFO)][["flag", "patent_id", "detail"]].to_string(index=False))
""")

nb = nbf.v4.new_notebook(cells=cells)
nb.metadata["kernelspec"] = {"display_name": "nb_01_review", "language": "python", "name": "nb_01_review"}
nb.metadata["language_info"] = {"name": "python"}
out = Path("/home/vasco/Vasco Workspace/Tese_Vasco_Lnx/Patent-Labelling-Tools/notebooks/04_master_labels.ipynb")
nbf.write(nb, out)
print("wrote", out, len(cells), "cells")
