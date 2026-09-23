"""Write the 2026-09-22 domain tags into the wizard record (TRL pass rulings).

User 2026-09-22: "if they have turbine they are not electric, but similar" and "change on the xlsx".
  ElectricSimilar — named turbine aircraft: AW609 (US2022073201A1), AW609 v1 (US2024167561A1, D3) and
                    its two D2s (EP4467450A1, US2023234701A1), S-97 Raider (US2023257103A1)
  UAVSimilar      — Panther (US2012091257A1 aircraft 2 only: "may be in the form of a UAV, for example a
                    tactical UAV"), KARI Quad Tilt-Prop UAV (KR20230075012A), Aergility ATLIS
                    (US2022043465A1, uncrewed cargo aircraft — same ruling as Bell APT 2026-09-21)

Writes a G1 / Edge-Case Tags / edgeTags row (after notPureArch; at the end of the record block when the
record has no G1 block, as for D1/D2 duplicates) or extends an existing edgeTags cell. Idempotent.
Also appends wizard_edge_tags rows to CORRECTIONS_identity.csv for the single-aircraft patents (a
patent-wide row would tag Panther's aircraft 1 too, so Panther gets the record tag only).

    python scripts/apply_domain_tags_0922.py            # dry run
    python scripts/apply_domain_tags_0922.py --apply    # backup + write
Then: split_wizard_all_export.py <record> --apply → build_identity_all.py → notebook 04 →
canonicalize_record.py --apply → notebook 04 → notebook 10.
"""
import csv
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

INPUTS = Path("/mnt/storage_11tb/Drive_files_to_syncronize/3 - Images DataSets & Labelling Outputs/1639_LABELLED/0_labelling/inputs")
RECORD = INPUTS / "record" / "reviewed_patents_Batch_ALL.xlsx"
CORR = INPUTS / "review_decisions" / "CORRECTIONS_identity.csv"
TS = datetime.now().strftime("%Y%m%d_%H%M%S")

TARGETS = [  # record id, tag, correction row (patent-wide) or None, reason
    ("US2022073201A1", "ElectricSimilar", True, "AW609 — turbine aircraft by fact (Leonardo AW609 tiltrotor)"),
    ("US2024167561A1", "ElectricSimilar", True, "AW609 v1 (D3 of US2022073201A1) — turbine aircraft by fact"),
    ("EP4467450A1", "ElectricSimilar", True, "D2 of US2024167561A1 (AW609 v1), inherits ElectricSimilar"),
    ("US2023257103A1", "ElectricSimilar", True, "S-97 Raider — turbine compound helicopter by fact"),
    ("US2023234701A1", "ElectricSimilar", True, "D2 of US2024167561A1 (AW609 v1), inherits ElectricSimilar"),
    ("US2012091257A1_arch2", "UAVSimilar", False, "IAI Panther tactical UAV; patent: 'may be in the form of a UAV, for example a tactical UAV'"),
    ("KR20230075012A", "UAVSimilar", True, "KARI Quad Tilt-Prop UAV — uncrewed demonstrator by fact"),
    ("US2022043465A1", "UAVSimilar", True, "Aergility ATLIS — uncrewed cargo aircraft by fact (same ruling as Bell APT 2026-09-21)"),
]


def main(apply: bool) -> None:
    wb = openpyxl.load_workbook(RECORD)
    ws = wb["Review"]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    plan = []
    for rid, tag, _, _ in TARGETS:
        rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, col["Patent_ID"]).value == rid]
        if not rows:
            raise SystemExit(f"{rid}: no rows in the record")
        tag_rows = [r for r in rows if ws.cell(r, col["Field"]).value == "edgeTags"]
        if tag_rows:
            r = tag_rows[0]
            cur = [t for t in str(ws.cell(r, col["Value"]).value or "").split("|") if t]
            if tag in cur:
                plan.append((rid, "already tagged", None))
            else:
                plan.append((rid, f"extend {'|'.join(cur)} -> {'|'.join(cur + [tag])}", ("set", r, "|".join(cur + [tag]))))
            continue
        after = [r for r in rows if ws.cell(r, col["Field"]).value == "notPureArch"]
        at = (after[0] if after else max(rows)) + 1
        plan.append((rid, f"insert edgeTags={tag} at row {at}" + ("" if after else " (end of block, no G1)"), ("insert", at, tag)))
    for rid, what, _ in plan:
        print(f"  {rid:24s} {what}")
    if not apply:
        print("dry run — pass --apply to write")
        return
    bak = RECORD.parent / "_backups" / f"reviewed_patents_Batch_ALL.PRE_DOMAINTAGS0922_{TS}.xlsx"
    shutil.copy2(RECORD, bak)
    # inserts from the bottom so earlier row numbers stay valid
    for rid, _, act in sorted(plan, key=lambda p: -(p[2][1] if p[2] else 0)):
        if act is None:
            continue
        kind, r, val = act
        if kind == "set":
            ws.cell(r, col["Value"]).value = val
        else:
            ws.insert_rows(r)
            for h, v in (("Patent_ID", rid), ("Section", "G1"), ("Sub_Dimension", "Edge-Case Tags"),
                         ("Field", "edgeTags"), ("Value", val)):
                ws.cell(r, col[h]).value = v
    wb.save(RECORD)
    # downstream mirror for the identity finals (patent-wide, single-aircraft patents only)
    with open(CORR, newline="", encoding="utf-8") as f:
        existing = {(r["patent_id"], r["field"], r["new_value"]) for r in csv.DictReader(f)}
    shutil.copy2(CORR, CORR.parent / "_backups" / f"CORRECTIONS_identity.PRE_DOMAINTAGS0922_{TS}.csv")
    with open(CORR, "a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for rid, tag, corr, why in TARGETS:
            if corr and (rid, "wizard_edge_tags", tag) not in existing:
                w.writerow([rid, "wizard_edge_tags", tag, f"user 2026-09-22 (TRL pass): {why}", "2026-09-22"])
    print(f"written; backup {bak.name}")


if __name__ == "__main__":
    main("--apply" in sys.argv)
