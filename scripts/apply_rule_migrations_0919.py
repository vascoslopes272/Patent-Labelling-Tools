#!/usr/bin/env python3
"""apply_rule_migrations_0919.py — record changes that follow from the rulings of 2026-09-19 (plan
0_labelling/audit_2026-09-18/PLAN_labelling_rules_2026-09-19.md, accepted by the user "yes to all").
    python scripts/apply_rule_migrations_0919.py            # dry run: list every change
    python scripts/apply_rule_migrations_0919.py --apply    # back up, write the record, re-split the batch files
(a) Tilting-boom rotors (rule 3.3, "the boom tick carries the tilt"): on an aircraft whose EVERY propeller-carrying boom
    group is ticked "Booms tilt", a boom-card rotor type recorded Tilt becomes Fixed. Partly ticked aircraft are NOT
    touched (the boom card cannot say which rotor type sits on which group): they are listed for the wizard.
(b) US2024132207A1 kept TR (user 2026-09-19 "the last one is TR"): the tail is recorded like the wing, as a tilting
    rotor — empTilts False, emp_propKin Tilt (otherwise the wizard blocks TR: rotor tilt + tilting surface = CVT).
(c) Quick counts left over after a station's override was switched off (rule 3.1): blanked.
Only the named cells change; values keep their types (openpyxl) — the wizard's xlBool() reads the text "False" as true.
After --apply: notebook 04 → canonicalize_record.py --apply → 04, and the user reloads the record in the wizard.
"""
import sys, re, shutil, subprocess, collections
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook, Workbook

ROOT = Path("/mnt/storage_11tb/Drive_files_to_syncronize/3 - Images DataSets & Labelling Outputs/1639_LABELLED/0_labelling")
RECORD = ROOT / "inputs" / "record" / "reviewed_patents_Batch_ALL.xlsx"
APPLY = "--apply" in sys.argv[1:]
FIXED = "Fixed — Fixed (No Articulation)"
TILT = "Tilt — Tilt (Vectoring Unit)"
def s(v): return "" if v is None else str(v)
def sid(v): return s(v).split(" — ")[0].strip()
def true(v): return v is True or s(v) in ("True", "1")

rows = [list(r) for r in load_workbook(RECORD)["Review"].iter_rows(values_only=True)]
hdr, body = rows[0], rows[1:]
ix = {c: i for i, c in enumerate(hdr)}
by = collections.defaultdict(dict)                     # Patent_ID (with _archN) -> field -> body index
for i, r in enumerate(body): by[s(r[ix["Patent_ID"]])].setdefault(s(r[ix["Field"]]), i)
def val(pid, f): i = by[pid].get(f); return body[i][ix["Value"]] if i is not None else None
changes, manual = [], []
def setv(pid, f, new, why):
    i = by[pid].get(f)
    if i is None: manual.append((pid, f"{f}: no row to change ({why})")); return
    old = body[i][ix["Value"]]
    if s(old) == s(new): return
    changes.append((pid, f, s(old), s(new), why)); body[i][ix["Value"]] = new

# (a) tilting-boom rotors
for pid, fields in by.items():
    groups = [int(m.group(1)) for f in fields for m in [re.fullmatch(r"boom(\d+)_(count|attach)", f)] if m]
    carrying = sorted({g for g in groups if not (s(val(pid, f"boom{g}_hasProps")) in ("False", "0") or val(pid, f"boom{g}_hasProps") is False)
                       and (s(val(pid, f"boom{g}_count")) not in ("", "0", "False") or s(val(pid, f"boom{g}_attach")))})
    if not carrying or not any(true(val(pid, f"boom{g}_tilts")) for g in carrying): continue
    tilt_types = [f for f in fields if re.fullmatch(r"boom(_t\d+)?_propKin", f) and sid(val(pid, f)) == "Tilt"]
    tilt_types = [f for f in tilt_types if s(val(pid, f.replace("_propKin", "_count"))) not in ("", "0", "False")]
    if not tilt_types: continue
    if all(true(val(pid, f"boom{g}_tilts")) for g in carrying):
        for f in tilt_types: setv(pid, f, FIXED, "rotor on a boom ticked 'Booms tilt' → Fixed (the boom carries the tilt)")
    else:
        manual.append((pid, f"boom groups partly ticked {[g for g in carrying if true(val(pid, f'boom{g}_tilts'))]} of {carrying}; "
                            f"rotor types marked Tilt: {', '.join(tilt_types)} — decide in the wizard which rotors ride a tilting boom"))
# (b) US2024132207A1 kept TR
setv("US2024132207A1", "empTilts", False, "kept TR (user): tail recorded like the wing")
setv("US2024132207A1", "emp_propKin", TILT, "kept TR (user): tail rotor tilts")
# (c) left-over quick counts
for pid, fields in by.items():
    for f in fields:
        m = re.fullmatch(r"(.+)_quickCount", f)
        if m and s(val(pid, f)) not in ("", "0", "False") and not true(val(pid, m.group(1) + "_quickOverride")):
            if s(val(pid, m.group(1) + "_count")) in ("", "0", "False"):   # the quick count is the ONLY record of these units
                manual.append((pid, f"{m.group(1)}: card counts 0 but a quick count {s(val(pid, f))} is left over — switch the "
                                    f"station to count only (override on) or enter its detail; never blanked here"))
            else:
                setv(pid, f, None, f"quick count left over — {m.group(1)} override is off, the card records its units")

print(f"record {RECORD.name}: {len(body)} rows → {len(changes)} cells to change on {len({c[0].split('_arch')[0] for c in changes})} patents; {len(manual)} for the wizard")
for pid, f, old, new, why in changes: print(f"  {pid:<24} {f:<18} {old[:34]!r:<38} -> {new[:34]!r}   {why}")
for pid, why in manual: print(f"  (wizard) {pid}: {why}")
if not APPLY: print("dry run — re-run with --apply"); sys.exit(0)
if not changes: print("nothing to write"); sys.exit(0)
ts = datetime.now().strftime("%Y%m%d_%H%M%S")
bdir = RECORD.parent / "_backups"; bdir.mkdir(exist_ok=True)
bak = bdir / f"{RECORD.stem}.PRE_RULES0919_{ts}.xlsx"; shutil.copy2(RECORD, bak)
wb = Workbook(); ws = wb.active; ws.title = "Review"
for r in [hdr] + body: ws.append(r)
wb.save(RECORD); print(f"written {RECORD}; backup {bak}")
subprocess.run([sys.executable, str(Path(__file__).with_name("split_wizard_all_export.py")), str(RECORD), "--apply"], check=True)
