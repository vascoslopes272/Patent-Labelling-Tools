#!/usr/bin/env python3
"""apply_mr_fields.py — write the multirotor fields picked on mr_fields_review.html into the wizard record.
    python scripts/apply_mr_fields.py [MR_FIELDS_DECISIONS.csv]            # dry run: list every change
    python scripts/apply_mr_fields.py [MR_FIELDS_DECISIONS.csv] --apply    # back up, write the record, re-split the batch files
Default decisions file: 0_labelling/inputs/review_decisions/MR_FIELDS_DECISIONS.csv (the page's export).
Only blanks are written: a field that already holds a value in the record is listed and left alone (never overwritten).
Values go in as the wizard writes them — orient as the bare id, chord / rmech / propKin as "id — Label" — into the
field's own M3 block ("Propulsion: <card>" or "Propulsion: <card> — Type N"), in the wizard's field order.
Rows marked in_wizard=yes are added to review_decisions/PENDING_WIZARD_REVIEW.csv instead.
Run notebook 04 → canonicalize_record.py --apply → 04 afterwards, as after any record change.
"""
import sys, re, json, shutil, subprocess, tempfile
from datetime import date, datetime
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook, Workbook

REPO = Path(__file__).resolve().parents[1]
ROOT = Path("/mnt/storage_11tb/Drive_files_to_syncronize/3 - Images DataSets & Labelling Outputs/1639_LABELLED/0_labelling")
RECORD = ROOT / "inputs" / "record" / "reviewed_patents_Batch_ALL.xlsx"
PENDING = ROOT / "inputs" / "review_decisions" / "PENDING_WIZARD_REVIEW.csv"
args = [a for a in sys.argv[1:] if not a.startswith("--")]
DEC = Path(args[0]) if args else ROOT / "inputs" / "review_decisions" / "MR_FIELDS_DECISIONS.csv"
APPLY = "--apply" in sys.argv[1:]
ORDER = ["count", "chord", "orient", "bmech", "rmech", "propKin", "ctrlOnly", "zone"]   # the wizard's field order in a type block
LISTS = {"chord": "CHORD", "rmech": "RETRACT_MECH", "propKin": "PROP_KIN"}             # orient is stored as the bare id
def s(v): return "" if v is None else str(v)

tmp = Path(tempfile.mkstemp(suffix=".json")[1])
subprocess.run(["node", str(REPO / "scripts" / "conformance" / "extract_html_schema.js"),
                str(REPO / "notebooks" / "UI_for_taxonomy_caracterization_15_4.html"), str(tmp)], check=True, capture_output=True)
LAB = json.loads(tmp.read_text())["labels"]
def value(field, vid):
    if field == "orient": return vid
    return f"{vid} — {LAB[LISTS[field]][vid]}"

D = pd.read_csv(DEC, dtype=str, keep_default_na=False)
wiz = D[D.in_wizard == "yes"]; D = D[(D.in_wizard != "yes") & (D.value != "")]
bad = D[~D.apply(lambda x: x.value in (LAB[LISTS[x.field]] if x.field in LISTS else {"Vertical", "Horizontal", "Mixed"}), axis=1)]
assert bad.empty, f"unknown option ids:\n{bad}"

rows = [list(r) for r in load_workbook(RECORD)["Review"].iter_rows(values_only=True)]
hdr, body = rows[0], rows[1:]
ix = {c: i for i, c in enumerate(hdr)}
pids = {s(r[ix["Patent_ID"]]) for r in body}
where = {(s(r[ix["Patent_ID"]]), s(r[ix["Field"]])): i for i, r in enumerate(body)}

written, inserted, kept = [], [], []
insert_after = {}
for d in D.itertuples():
    rid = f"{d.patent_id}_arch{d.ua}" if f"{d.patent_id}_arch{d.ua}" in pids else d.patent_id
    assert rid in pids, f"{d.aircraft_id}: no record rows under {rid}"
    field, new = f"{d.group}_{d.field}", value(d.field, d.value)
    i = where.get((rid, field))
    if i is not None:
        old = s(body[i][ix["Value"]])
        if old.strip(): kept.append((rid, field, old, new)); continue
        body[i][ix["Value"]] = new; written.append((rid, field, new)); continue
    m = re.match(r"^(.+?)(?:_t(\d+))?$", d.group)
    sub = f"Propulsion: {m.group(1)}" + (f" — Type {m.group(2)}" if m.group(2) else "")
    block = sorted(j for (p, f), j in where.items() if p == rid and s(body[j][ix["Sub_Dimension"]]) == sub)
    assert block, f"{rid}: no {sub!r} block in the record"
    before = ORDER[:ORDER.index(d.field)]
    prev = [j for j in block if s(body[j][ix["Field"]])[len(d.group) + 1:] in before]
    anchor = max(prev) if prev else max(block)
    row = [None] * len(hdr)
    for c, v in (("Patent_ID", rid), ("Section", "M3"), ("Sub_Dimension", sub), ("Field", field), ("Value", new),
                 ("Source", "human"), ("Image_Path", body[anchor][ix["Image_Path"]])):
        if c in ix: row[ix[c]] = v
    insert_after.setdefault(anchor, []).append(row); inserted.append((rid, field, new))
for lst in insert_after.values():                     # several fields after one anchor: keep the wizard's order
    lst.sort(key=lambda r: ORDER.index(s(r[ix["Field"]]).rsplit("_", 1)[-1]))

print(f"{DEC.name}: {len(D)} picks, {len(wiz)} aircraft left for the wizard -> {len(written)} blank cells filled, "
      f"{len(inserted)} rows added, {len(kept)} already hold a value (left alone)")
for rid, f, new in written + inserted: print(f"  {rid:<22} {f:<22} {new}")
for rid, f, old, new in kept: print(f"  (kept) {rid} {f} = {old!r}, page said {new!r}")
for w in wiz.itertuples(): print(f"  -> wizard: {w.aircraft_id} {w.note}")
if not APPLY: print("dry run — re-run with --apply"); sys.exit(0)

if len(wiz):
    P = pd.read_csv(PENDING, dtype=str, keep_default_na=False)
    new = pd.DataFrame([dict(patent_id=w.patent_id, what_to_check="multirotor propulsion fields never recorded (arrangement / thrust direction / retraction / articulation)"
                              + (f"; note: {w.note}" if w.note else ""), why="user chose the wizard on mr_fields_review.html", added=str(date.today()))
                        for w in wiz.itertuples() if w.patent_id not in set(P.patent_id)])
    pd.concat([P, new]).to_csv(PENDING, index=False); print(f"{len(new)} patents added to {PENDING.name}")
if not (written or inserted): print("nothing to write"); sys.exit(0)
out = [hdr]
for i, r in enumerate(body):
    out.append(r); out += insert_after.get(i, [])
ts = datetime.now().strftime("%Y%m%d_%H%M%S")
bdir = RECORD.parent / "_backups"; bdir.mkdir(exist_ok=True)
bak = bdir / f"{RECORD.stem}.PRE_MRFIELDS_{ts}.xlsx"; shutil.copy2(RECORD, bak)
wb = Workbook(); ws = wb.active; ws.title = "Review"
for r in out: ws.append(r)
wb.save(RECORD); print(f"written {RECORD} ({len(out) - 1} rows); backup {bak}")
subprocess.run([sys.executable, str(Path(__file__).with_name("split_wizard_all_export.py")), str(RECORD), "--apply"], check=True)
