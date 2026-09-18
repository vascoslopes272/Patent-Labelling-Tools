#!/usr/bin/env python3
"""build_damage_review_page.py — the review page for labels that past wizard saves damaged (audit 2026-09-18, audit_7).
    python scripts/build_damage_review_page.py
Reads  0_labelling/audit_2026-09-18/audit_7_history/damage.csv (rows still bad in the current record), the whole
current record, the snapshot each good value comes from, notebook 04's aircraft table (names, PDF links, identity
answers) and the wizard's option lists (scripts/conformance/extract_html_schema.js).
Writes notebooks/post-process/damage_review.html (template scripts/damage_review_template.html) — self-contained,
figures by file:// path.

v2 (user 2026-09-18: "not sufficient info to choose — I only see a few images, I don't know if it is the main"):
every card shows ALL figures of the patent grouped by aircraft with their labels (status, MAIN, perspective, flight
state, style), a now-vs-after table, warnings (two mains, a missing aircraft, …) and lets the user set the main figure,
status, perspective and flight state of the figures involved. The flight state follows the wizard's lock
(src/wizard_view.py): a locked aircraft shows Invariant and cannot be changed.
Export: DAMAGE_RESTORE_DECISIONS.csv into 1639_LABELLED/review_decisions/ (decision + edits + main per aircraft).
Nothing is written to the record here.
"""
import json, re, subprocess, sys, tempfile
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from src.wizard_view import architectures, moves, CONVERTIBLE

L0 = Path("/mnt/storage_11tb/Drive_files_to_syncronize/3 - Images DataSets & Labelling Outputs/1639_LABELLED/0_labelling")
DAMAGE = L0 / "audit_2026-09-18" / "audit_7_history" / "damage.csv"
RECORD = L0 / "inputs" / "record" / "reviewed_patents_Batch_ALL.xlsx"
TABLES = L0 / "outputs" / "tables"
IMG_ROOT = L0 / "inputs" / "images"
HTML = REPO / "notebooks" / "UI_for_taxonomy_caracterization_15_4.html"
OUT = REPO / "notebooks" / "post-process" / "damage_review.html"
FIG_FIELDS = ["status", "isMain", "arch", "per", "acState", "acSty", "acCol", "bgSty", "bgCol", "parts", "qualityFlag", "rotation_deg", "comment", "dupOf"]

def sid(v):
    v = "" if v is None else str(v)
    return v.split(" — ", 1)[0].strip()
def ua_of(a): a = str(a).strip(); return int(a) if a.isdigit() and int(a) > 0 else 1
def base(p): return re.sub(r"_arch\d+$", "", str(p))

def read_rows(path, pids):
    ws = load_workbook(path, read_only=True)
    ws = ws["Review"] if "Review" in ws.sheetnames else ws.worksheets[0]
    it = ws.iter_rows(values_only=True); hdr = [str(h) for h in next(it)]; ix = {c: i for i, c in enumerate(hdr)}
    out = [{c: ("" if r[ix[c]] is None else str(r[ix[c]])) for c in hdr}
           for r in it if pids is None or base(r[ix["Patent_ID"]] or "") in pids]
    df = pd.DataFrame(out); df["base"] = df.Patent_ID.map(base); return df

_file_index = None
def resolve(path, fname, pid):
    """An image path from any snapshot → a file that exists today, looked up ONLY in the patent's own folder
    (generic names such as 'Pasted image.png' or 'FT_8_crop_0_Fu.png' repeat across patents)."""
    global _file_index
    p = (path or "").replace("file://", "")
    if p and pid in p and Path(p).exists(): return p
    if _file_index is None:
        _file_index = {}
        for q in IMG_ROOT.rglob("*.png"): _file_index.setdefault((q.parent.name.split("_")[0], q.name), str(q))
    return _file_index.get((pid, fname), "")

def rot(v):
    try: return int(float(v))
    except Exception: return 0

def block(df, pid, sub):
    b = df[(df.base == pid) & (df.Sub_Dimension == sub)]
    vals = {r.Field: sid(r.Value) for r in b.itertuples()}
    ip = next((p for p in b.Image_Path if p), "") if "Image_Path" in b else ""
    return vals, ip

def figdict(pid, fname, vals, ip, **extra):
    d = {k: vals.get(k, "") for k in FIG_FIELDS}
    src = resolve(ip, fname, pid)
    d.update(file=fname, src=("file://" + src) if src else "", ua=ua_of(vals.get("arch") or "1"),
             isMain=(vals.get("isMain") == "True"), rot=rot(vals.get("rotation_deg")))
    d.update(extra); return d

# ── inputs ──
tmp = Path(tempfile.mkstemp(suffix=".json")[1])
subprocess.run(["node", str(REPO / "scripts" / "conformance" / "extract_html_schema.js"), str(HTML), str(tmp)], check=True, capture_output=True)
SCHEMA = json.loads(tmp.read_text())
PER = SCHEMA["lists"]["T2_PER"]
STATE = [["Hover", "Hover"], ["Transition", "Transition"], ["Cruise", "Cruise"], ["HoverCruise", "Invariant"], ["Other", "Other"]]
TOPNAME = SCHEMA["labels"]["TOP"]

D = pd.read_csv(DAMAGE, dtype=str, keep_default_na=False)
D = D[D.still_bad_in_current_record == "yes"].reset_index(drop=True)
REC = read_rows(RECORD, None)
AT = pd.read_csv(TABLES / "aircraft_table.csv", dtype=str, keep_default_na=False, low_memory=False)
ARCH = architectures(REC)                                   # (pid, arch) -> fields, for every architecture with a topType
T2 = REC[(REC.Section == "T2") & REC.Sub_Dimension.str.startswith("Image: ") & ~REC.Sub_Dimension.str.contains(r"\(fig", regex=True)]
BLOCKS = {}
for (pid, sub), g in T2.groupby(["base", "Sub_Dimension"], sort=False):
    BLOCKS.setdefault(pid, []).append((sub, {r.Field: sid(r.Value) for r in g.itertuples()}, next((p for p in g.Image_Path if p), "")))
SNAP = {}
def snap(path):
    if path not in SNAP: SNAP[path] = read_rows(path, set(D.patent_id))
    return SNAP[path]

PATENTS = {}
def patent(pid):
    """Everything a card shows about one patent: its aircraft (type, lock, wings) and all its figures."""
    if pid in PATENTS: return PATENTS[pid]
    n_ua = int(max([a for (p, a) in ARCH if p == pid] + [1]))
    ac = {}
    for a in range(1, n_ua + 1):
        d = ARCH.get((pid, a), {})
        tt = d.get("topType", "")
        wing = {k: d[k] for k in ("wingConf", "wCount", "wing1_role", "wing1_tilt", "wing1_plan", "wing1_posV") if d.get(k)}
        ac[a] = dict(type=tt, typeName=TOPNAME.get(tt, ""), locked=bool(tt) and tt not in CONVERTIBLE and not moves(d),
                     moves="; ".join(moves(d)) if d else "", wing=wing, tags=d.get("edgeTags", ""))
    figs = [figdict(pid, sub[7:], v, ip) for sub, v, ip in BLOCKS.get(pid, [])]
    a = AT[AT.patent_id == pid]
    r0 = a.iloc[0] if len(a) else None
    get = lambda c: (str(r0[c]) if r0 is not None and c in r0 else "")
    PATENTS[pid] = dict(n_ua=n_ua, aircraft=ac, figs=figs, title=get("title"), company=get("company"), pdf=get("pdf_link"),
                        approved=get("is_approved"), dup=get("dup_type"), dupOf=get("dup_of"),
                        names={int(x.ua): x.aircraft_name for x in a.itertuples()},
                        identity={int(x.ua): dict(uav=x.uav_final, electric=x.is_electric_final, takeoff=x.takeoff_final) for x in a.itertuples()})
    return PATENTS[pid]

def show_state(v): return "Invariant" if v == "HoverCruise" else v
def short(f): return ", ".join(f"{k}={show_state(f[k])}" for k in ("status", "per", "acState") if f.get(k)) + (", MAIN" if f.get("isMain") else "")

# ── cards ──
items, w7, HANDLED = [], {}, set()
for r in D.itertuples():
    ua = ua_of(r.arch); P = patent(r.patent_id)
    head = dict(pid=r.patent_id, ua=ua, bug=r.bug, conf=r.confidence, snapshot=r.last_good_snapshot,
                good=r.good_value, bad=r.bad_value, field=r.field_or_figure, note=r.note.split(" | current")[0],
                warn=[], changes=[], involved={}, editable=[], restoreFig=None)
    if r.bug == "W2":
        fname = re.search(r"Image: (.+?\.png)", r.field_or_figure).group(1)
        gv, gip = block(snap(r.last_good_snapshot), r.patent_id, "Image: " + fname)
        m = re.search(r"same-FIG in current: (.+?\.png)", r.note); repl = m.group(1) if m else ""
        cur = {f["file"]: f for f in P["figs"]}
        cv = cur.get(repl, {}); still = cur.get(fname)
        lost_src = resolve(gip, fname, r.patent_id)
        stem = re.sub(r"_F[^_]*\.png$", "", fname)
        renamed = [f["file"] for f in P["figs"] if f["file"] != fname and re.sub(r"_F[^_]*\.png$", "", f["file"]) == stem]
        repl_blank = bool(repl) and not cv.get("status")
        if repl_blank: HANDLED.add((r.patent_id, repl))
        if still and not still.get("status"): HANDLED.add((r.patent_id, fname))
        good = figdict(r.patent_id, fname, gv, gip)
        g_ua = good["ua"]
        if renamed:
            rn = cur[renamed[0]]
            head["involved"] = {rn["file"]: "renamed crop — holds your labels"}
            if repl: head["involved"][repl] = "unreviewed leftover" if repl_blank else "reviewed by you"
            same = all(rn.get(k) == good.get(k) for k in ("status", "per", "acState"))
            q = (f"Nothing was lost: the crop you reviewed as {fname} ({short(good)}) was re-cropped/renamed to {rn['file']}"
                 + (", which carries the same labels." if same else f", which you reviewed again later ({short(rn)})."))
            q += (f" {repl} has the same figure number, was never reviewed and has no status. Remove it?" if repl_blank
                  else f" {repl} has the same figure number and was reviewed by you. Leave everything as it is?" if repl else " Leave it as it is?")
            opts = ([["restore", "Remove the leftover"], ["keep", "Keep the leftover"], ["unsure", "Not sure"]] if repl_blank
                    else [["keep", "OK, leave it"], ["unsure", "Not sure"]])
            kind, rec = "renamed", ("restore" if repl_blank else "keep")
            if repl_blank: head["changes"] = [["figure " + repl, "in the record, never reviewed, no status", "removed"]]
        elif not lost_src:
            q = (f"Your reviewed figure {fname} ({short(good)}) was dropped on a save and its image file no longer exists anywhere "
                 "on disk, so it cannot be restored. Leave it as it is?")
            if repl: head["involved"] = {repl: "holds the figure number now"}
            opts = [["keep", "OK, leave it"], ["unsure", "Not sure"]]
            kind, rec = "gone", "keep"
        else:
            good.update(src="file://" + lost_src, restore=True)
            head["restoreFig"] = good
            head["involved"] = {fname: "your figure — to restore"}
            if repl: head["involved"][repl] = "unreviewed replacement" if repl_blank else "reviewed by you — kept"
            head["editable"] = [fname]
            q = ("A save dropped the figure you reviewed — shown below in its aircraft with a green RESTORE frame and its old labels"
                 + (", and put an unreviewed crop of the same figure number in its place (red frame)" if repl_blank else
                    f"; the same figure number is now held by {repl}, which you reviewed ({short(cv) or 'no status'}) — maybe your own later decision" if repl else "")
                 + ". Restore it? Check or correct its labels and the main figure below first.")
            head["changes"] = [["figure " + fname, ("in the record, no status / default labels" if still else "not in the record"), "restored: " + short(good)]]
            if repl_blank: head["changes"].append(["figure " + repl, "in the record, never reviewed", "removed"])
            if g_ua > P["n_ua"]:
                head["warn"].append(f"This figure belonged to aircraft {g_ua}, but the patent now has only {P['n_ua']} aircraft — you probably "
                                    f"removed aircraft {g_ua} on purpose. If you restore it, choose its aircraft on the figure.")
            if good["isMain"]:
                tgt = min(g_ua, P["n_ua"])
                cur_main = [f["file"] for f in P["figs"] if f["isMain"] and f["ua"] == tgt and f["status"] == "approved"]
                if cur_main:
                    head["warn"].append(f"It was the MAIN figure; aircraft {tgt} now has {cur_main[0]} as main. An aircraft keeps ONE main — "
                                        "the page keeps the current one unless you click “make MAIN” on the restored figure.")
                    good["isMain"] = False; good["wasMain"] = True
            opts = [["restore", "Restore my figure"], ["keep", "Keep what is there now"], ["unsure", "Not sure"]]
            kind, rec = "lost", ("restore" if r.confidence == "high" and (repl_blank or not repl) else "")
        items.append(dict(head, key=f"W2|{r.patent_id}|{fname}", title=f"Figure {fname}", q=q, opts=opts, rec=rec, kind=kind,
                          aux=json.dumps(dict(kind=kind, file=fname, replacement=repl, replacement_blank=repl_blank, renamed=renamed))))
    elif r.bug == "W7":
        fname = re.search(r"Image: (.+?\.png)", r.field_or_figure).group(1)
        fld = r.field_or_figure.rsplit("/", 1)[-1].strip()
        k = (r.patent_id, fname)
        if k not in w7:
            src = re.search(r"source figure (\S+?\.png)", r.note)
            w7[k] = dict(head, key=f"W7|{r.patent_id}|{fname}", title=f"Figure {fname}", kind="copied", fields={},
                         involved={fname: "labels overwritten", **({src.group(1): "the figure it duplicates"} if src else {})},
                         editable=[fname], opts=[["restore", "Restore its own labels"], ["keep", "Keep the copied labels"], ["unsure", "Not sure"]],
                         rec="restore" if r.confidence in ("high", "medium") else "")
            items.append(w7[k])
        w7[k]["fields"][fld] = [sid(r.good_value), sid(r.bad_value)]
        w7[k]["changes"] = [[f, show_state(b), show_state(g)] for f, (g, b) in w7[k]["fields"].items()]
        w7[k]["q"] = ("On reload the wizard copied the labels of the figure this one duplicates over its own. Restore its own labels "
                      "(table)? You can also correct them on the figure below.")
        w7[k]["aux"] = json.dumps(dict(file=fname, fields=w7[k]["fields"]))
    elif r.bug == "W3":
        ac = P["aircraft"].get(ua, {}); tt = ac.get("type", "")
        wing = ", ".join(f"{k}={v}" for k, v in ac.get("wing", {}).items()) or "wing fields"
        q = (f"Aircraft {ua} is a {tt} ({ac.get('typeName','')}) and is stored with a wing the wizard never shows for this type ({wing}). "
             + ("Before a save it had no wing. " if r.confidence == "high" else "It has been like this since its first save — no backup has a good value. ")
             + "Look at its figures below: is it wingless?")
        head["changes"] = [["wings of aircraft " + str(ua), wing, "no wing (count 0, wing fields removed)"]]
        items.append(dict(head, key=f"W3|{r.patent_id}|{ua}", title=f"Wings of aircraft {ua}", q=q, kind="wing",
                          opts=[["restore", "Wingless (count 0)"], ["keep", "It has a real wing — keep (then its type is probably wrong)"], ["unsure", "Not sure"]],
                          rec="restore", aux=""))
    else:
        fld = re.search(r"/(\w+)", r.field_or_figure).group(1) if "/" in r.field_or_figure else r.field_or_figure
        g, b = sid(r.good_value), sid(r.bad_value)
        idn = P["identity"].get(ua, {})
        txt = {"W1": (f"A save blanked this aircraft's tag “{g}”. Identity review for this aircraft: UAV = {idn.get('uav') or '—'}, "
                      f"electric = {idn.get('electric') or '—'}, take-off = {idn.get('takeoff') or '—'}. Restore the tag?", "Restore the tag", "Keep it without the tag"),
               "W4": ("A reload turned your wing planform “Trapezoidal” into “Other” with the note “Trap”. Restore Trapezoidal?", "Restore Trapezoidal", "Keep Other"),
               "W6": (f"A save dropped the quick count of {fld.replace('_quickCount','')} (was {g}). Restore it?", "Restore the count", "Keep it empty"),
               "W8": (f"A reload removed the thrust direction “Mixed” from {fld}. Restore Mixed?", "Restore Mixed", "Keep it empty")}[r.bug]
        head["changes"] = [[fld, b or "(empty)", g]]
        items.append(dict(head, key=f"{r.bug}|{r.patent_id}|{ua}|{fld}", title=fld, q=txt[0], kind="value",
                          opts=[["restore", txt[1]], ["keep", txt[2]], ["unsure", "Not sure"]],
                          rec="restore" if r.confidence == "high" else "", aux=""))

# unreviewed crops with no status on an APPROVED patent, not already part of a W2 card
appr = set(REC[(REC.Field == "isApproved") & ~REC.Patent_ID.str.contains("_arch") & (REC.Value == "True")].base)
for pid, blocks in BLOCKS.items():
    if pid not in appr: continue
    for sub, v, ip in blocks:
        fname = sub[7:]
        if v.get("status") or (pid, fname) in HANDLED: continue
        patent(pid)
        items.append(dict(pid=pid, ua=ua_of(v.get("arch") or "1"), bug="W2x", conf="high", snapshot="", good="(no block)", bad="block with no status",
                          field=sub, note="Same feed merge as the dropped figures, but nothing of yours was dropped.", warn=[],
                          changes=[["figure " + fname, "in the record, never reviewed, no status", "removed"]],
                          involved={fname: "unreviewed — no status"}, editable=[], restoreFig=None, kind="leftover",
                          key=f"W2x|{pid}|{fname}", title=f"Figure {fname}",
                          q="An unreviewed crop (red frame) slipped into this reviewed patent on a save — no status, default labels. You never reviewed it. Remove it?",
                          opts=[["restore", "Remove it"], ["keep", "Keep it (I will review it in the wizard)"], ["unsure", "Not sure"]],
                          rec="restore", aux=json.dumps(dict(kind="leftover", file=fname))))

ORDER = {"W2": 0, "W2x": 1, "W1": 2, "W7": 3, "W4": 4, "W8": 5, "W6": 6, "W3": 7}
items.sort(key=lambda x: (ORDER[x["bug"]], x["conf"] != "high", x["pid"], x["ua"]))
BUGNAME = {"W1": "Tag lost", "W2": "Reviewed figure dropped", "W2x": "Unreviewed crop slipped in", "W3": "Phantom wing",
           "W4": "Trap → Other", "W6": "Quick count lost", "W7": "Figure labels copied", "W8": "“Mixed” removed"}
used = {it["pid"] for it in items}
data = json.dumps(dict(items=items, patents={p: PATENTS[p] for p in used}, bugname=BUGNAME, per=PER, state=STATE), ensure_ascii=False,
                  default=lambda o: o.item() if hasattr(o, "item") else str(o))   # numpy ints from the arch groupby
n_img = sum(len(PATENTS[p]["figs"]) for p in used); n_missing = sum(1 for p in used for f in PATENTS[p]["figs"] if not f["src"])
PAGE = (REPO / "scripts" / "damage_review_template.html").read_text(encoding="utf-8")
OUT.write_text(PAGE.replace("__DATA__", data), encoding="utf-8")
print(f"wrote {OUT}: {len(items)} cards {pd.Series([i['bug'] for i in items]).value_counts().to_dict()}; "
      f"{len(used)} patents, {n_img} figures ({n_missing} without an image file)")
