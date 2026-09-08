#!/usr/bin/env python3
"""
build_labelled_folder.py — assemble `1639_LABELLED/` from `1639_DS/`.

The labelling work tree, cut down to what is actually needed to open the wizard,
review, and join the batches from here on. `1639_DS/` is left untouched and
becomes the archive (raw downloads, triage JSON, the frozen 03_HUMAN / 03d
originals, the July 02a/02b outputs).

    1639_LABELLED/
      README.md  MANIFEST.csv  MANIFEST.sha256
      batches.xlsx  DATA_DICTIONARY.md  family_map.csv  descriptions.csv  dataset_summary.csv
      labels/
        reviewed_patents_Batch_NN.xlsx      human record  (= 03c at build time), FLAT, like 03c
        CORRECTIONS_LOG.csv
        Batch_NN/
          ml_predict_labels_Batch_NN.xlsx   the wizard feed        (like data_matched/)
          aircraft_identity_Batch_NN.xlsx   Stage 03a
          crops_mapping_Batch_NN.csv        the crop ledger
      images/                                mirror of 00b2_figure_crops, only files an xlsx points at
      wizard/                                read-only copy of the HTML the labels were made with
      joined/                                empty; the join notebook writes here

Why the two shapes under labels/: the reviewed files stay FLAT so that
`paths.corrected_wizard_exports` can point at `labels/` and 02a / 03a read them
unchanged; the per-batch subfolders mirror `paths.data_matched`, so 03a keeps
writing `aircraft_identity_<batch>.xlsx` next to the feed. Zero code changes,
one config edit.

The wizard renders images through `file://` + the absolute `Image_Path` cell,
so every Image_Path in the copied feed and reviewed workbooks is rewritten from
`1639_DS/matched/...` or `1639_DS/00b2_figure_crops/...` to
`1639_LABELLED/images/...`. Only the Image_Path column is touched.

Usage:
    python scripts/build_labelled_folder.py            # dry run: what would be copied / rewritten
    python scripts/build_labelled_folder.py --apply    # do it, then verify
    python scripts/build_labelled_folder.py --verify   # re-check an existing build
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import shutil
import sys
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path("/mnt/storage_11tb/Drive_files_to_syncronize/3 - Images DataSets & Labelling Outputs")
SRC = ROOT / "1639_DS"
DST = ROOT / "1639_LABELLED"
REPO = Path(__file__).resolve().parents[1]

BATCHES = [f"Batch_{i:02d}" for i in range(1, 6)]
OLD_IMAGE_ROOTS = [SRC / "matched", SRC / "00b2_figure_crops"]     # matched -> 00b2_figure_crops (symlink)
NEW_IMAGE_ROOT = DST / "images"

FEED_DIR = SRC / "data" / "00b2_crops_and_01a_MACHINE_feed"
HUMAN_DIR = SRC / "data" / "03c_CORRECTED_wizard_exports"
STATS_DIR = SRC / "data" / "Global Statistics"
WIZARD_HTML = REPO / "notebooks" / "UI_for_taxonomy_caracterization_15_4.html"

STATS_FILES = ["batches.xlsx", "DATA_DICTIONARY.md", "family_map.csv",
               "descriptions.csv", "dataset_summary.csv"]


# ─── helpers ─────────────────────────────────────────────────────────────────

def rebase(path: str) -> tuple[str, bool]:
    """Old absolute Image_Path -> new one. (unchanged, False) if not under an old root."""
    p = str(path).strip()
    for root in OLD_IMAGE_ROOTS:
        prefix = str(root) + "/"
        if p.startswith(prefix):
            return str(NEW_IMAGE_ROOT / p[len(prefix):]), True
    return p, False


def image_paths(xlsx: Path) -> list[str]:
    ws = openpyxl.load_workbook(xlsx, read_only=True)["Review"]
    rows = ws.iter_rows(values_only=True)
    hdr = next(rows)
    i = hdr.index("Image_Path")
    out = []
    for r in rows:
        v = r[i]
        if v is not None and str(v).strip() not in ("", "nan", "None", "null"):
            out.append(str(v).strip())
    return out


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def rewrite_xlsx(src: Path, dst: Path) -> tuple[int, int]:
    """Copy src -> dst rewriting the Image_Path column. Returns (rewritten, left)."""
    wb = openpyxl.load_workbook(src)
    ws = wb["Review"]
    hdr = [c.value for c in ws[1]]
    col = hdr.index("Image_Path") + 1
    rewritten = left = 0
    for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        cell = row[0]
        if cell.value is None or str(cell.value).strip() in ("", "nan", "None", "null"):
            continue
        new, changed = rebase(cell.value)
        if changed:
            cell.value = new
            rewritten += 1
        else:
            left += 1
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return rewritten, left


# ─── plan ────────────────────────────────────────────────────────────────────

def build_plan() -> dict:
    plan = {"copies": [], "xlsx": [], "images": {}, "missing_src": [], "unrebased": set(),
            "warnings": []}

    for f in STATS_FILES:
        s = STATS_DIR / f
        if s.exists():
            plan["copies"].append((s, DST / f))
        else:
            plan["warnings"].append(f"missing {s}")

    for b in BATCHES:
        feed = FEED_DIR / b / f"ml_predict_labels_{b}.xlsx"
        human = HUMAN_DIR / f"reviewed_patents_{b}.xlsx"
        ident = FEED_DIR / b / f"aircraft_identity_{b}.xlsx"
        crops = FEED_DIR / b / f"crops_mapping_{b}.csv"
        for s, d, kind in [(feed, DST / "labels" / b / feed.name, "feed"),
                           (human, DST / "labels" / human.name, "human")]:
            if s.exists():
                plan["xlsx"].append((s, d, kind))
                for p in image_paths(s):
                    new, changed = rebase(p)
                    if not changed:
                        plan["unrebased"].add(p)
                        continue
                    plan["images"][p] = new
            else:
                plan["warnings"].append(f"missing {s}")
        for s, d in [(ident, DST / "labels" / b / ident.name), (crops, DST / "labels" / b / crops.name)]:
            if s.exists():
                plan["copies"].append((s, d))
            else:
                plan["warnings"].append(f"missing {s}")

    log = HUMAN_DIR / "CORRECTIONS_LOG.csv"
    if log.exists():
        plan["copies"].append((log, DST / "labels" / log.name))

    if WIZARD_HTML.exists():
        plan["copies"].append((WIZARD_HTML, DST / "wizard" /
                               f"{WIZARD_HTML.stem}_{date.today():%Y%m%d}.html"))
    else:
        plan["warnings"].append(f"missing {WIZARD_HTML}")

    plan["missing_src"] = sorted(p for p in plan["images"] if not Path(p).exists())
    return plan


def print_plan(plan: dict) -> None:
    n_img = len(plan["images"])
    size = sum(Path(p).stat().st_size for p in plan["images"] if Path(p).exists())
    print(f"Destination : {DST}")
    print(f"Plain copies: {len(plan['copies'])}")
    for s, d in plan["copies"]:
        print(f"   {s.relative_to(SRC) if s.is_relative_to(SRC) else s}  ->  {d.relative_to(DST)}")
    print(f"Workbooks with Image_Path rewrite: {len(plan['xlsx'])}")
    for s, d, k in plan["xlsx"]:
        print(f"   [{k}] {s.relative_to(SRC)}  ->  {d.relative_to(DST)}")
    print(f"Images to copy: {n_img} files, {size/1e6:.0f} MB "
          f"({len(plan['missing_src'])} referenced but missing in 1639_DS — will stay missing)")
    for p in plan["missing_src"][:8]:
        print(f"   missing: {p}")
    if plan["unrebased"]:
        print(f"Image_Path values NOT under an old image root (left as they are): {len(plan['unrebased'])}")
        for p in sorted(plan["unrebased"])[:5]:
            print(f"   {p}")
    for w in plan["warnings"]:
        print(f"⚠  {w}")
    sample = list(plan["images"].items())[:2]
    for old, new in sample:
        print(f"rewrite example:\n   {old}\n   -> {new}")


# ─── apply ───────────────────────────────────────────────────────────────────

def apply(plan: dict) -> None:
    if DST.exists() and any(DST.iterdir()):
        print(f"⚠  {DST} already exists and is not empty — files are overwritten in place, "
              f"nothing is deleted.")
    for d in ("labels", "images", "wizard", "joined"):
        (DST / d).mkdir(parents=True, exist_ok=True)

    for s, d in plan["copies"]:
        d.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(s, d)
    for s, d, kind in plan["xlsx"]:
        rw, left = rewrite_xlsx(s, d)
        print(f"  {d.relative_to(DST)}: {rw} Image_Path cells rewritten, {left} left unchanged")

    copied = skipped = 0
    for old, new in plan["images"].items():
        src, dst = Path(old), Path(new)
        if not src.exists():
            continue
        if dst.exists() and dst.stat().st_size == src.stat().st_size:
            skipped += 1
            continue
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
        copied += 1
        if copied % 2000 == 0:
            print(f"  images: {copied} copied…")
    print(f"  images: {copied} copied, {skipped} already present")

    for html in (DST / "wizard").glob("*.html"):
        os.chmod(html, 0o444)
    (DST / "joined" / "README.md").write_text(
        "Outputs of the batch-join stage go here (corpus-wide xlsx). Empty at build time.\n")

    write_manifest(plan)
    write_readme(plan)


def write_manifest(plan: dict) -> None:
    sources = {d: s for s, d in plan["copies"]}
    sources.update({d: s for s, d, _ in plan["xlsx"]})
    sources.update({Path(new): Path(old) for old, new in plan["images"].items()})
    rows = []
    for p in sorted(x for x in DST.rglob("*") if x.is_file() and x.name not in ("MANIFEST.csv", "MANIFEST.sha256")):
        rows.append({"dest": str(p.relative_to(DST)), "bytes": p.stat().st_size,
                     "sha256": sha256(p), "source": str(sources.get(p, ""))})
    with open(DST / "MANIFEST.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["dest", "bytes", "sha256", "source"])
        w.writeheader(); w.writerows(rows)
    with open(DST / "MANIFEST.sha256", "w") as f:
        for r in rows:
            f.write(f"{r['sha256']}  {r['dest']}\n")
    print(f"  manifest: {len(rows)} files hashed")


def write_readme(plan: dict) -> None:
    n_img = len(plan["images"])
    (DST / "README.md").write_text(f"""# 1639_LABELLED — the labelling work tree

Built {date.today().isoformat()} from `1639_DS/` by `Patent-Labelling-Tools/scripts/build_labelled_folder.py`.
`1639_DS/` is the ARCHIVE from that date on: raw downloads, triage JSON, the frozen pass-1 / pass-2
wizard originals (`03_HUMAN_wizard_exports`, `03d_HUMAN_PASS2_wizard_exports`, with their hash
manifests), the July 02a/02b outputs, the legacy conformance folder. Nothing here is needed to
re-derive those; nothing there is needed to work from here.

## What is where

| Path | What | Source in 1639_DS |
|---|---|---|
| `batches.xlsx` | which patent is in which batch (+ company_canonical, prototype_label) | `data/Global Statistics/` |
| `DATA_DICTIONARY.md`, `family_map.csv`, `descriptions.csv`, `dataset_summary.csv` | corpus reference tables | `data/Global Statistics/` |
| `labels/reviewed_patents_Batch_NN.xlsx` | **the human record** — the wizard export after every correction (= `03c` at build time). Load it in the wizard with *Resume*. | `data/03c_CORRECTED_wizard_exports/` |
| `labels/CORRECTIONS_LOG.csv` | audit trail of every post-export correction | same |
| `labels/Batch_NN/ml_predict_labels_Batch_NN.xlsx` | the machine feed you *Load Batch* in the wizard | `data/00b2_crops_and_01a_MACHINE_feed/Batch_NN/` |
| `labels/Batch_NN/aircraft_identity_Batch_NN.xlsx` | Stage 03a: real aircraft name / electric / VTOL review sheet | same |
| `labels/Batch_NN/crops_mapping_Batch_NN.csv` | crop ledger: page image -> crop -> OCR label -> description line | same |
| `images/` | every cropped figure an xlsx points at ({n_img} files), same relative layout as `00b2_figure_crops/` | `00b2_figure_crops/` (alias `matched/`) |
| `wizard/` | read-only copy of the HTML version these labels were made with. **Edit the one in the repo, not this.** | `Patent-Labelling-Tools/notebooks/` |
| `joined/` | where the batch-join stage writes the corpus-wide tables | — |
| `MANIFEST.csv` / `MANIFEST.sha256` | hash + source path of every file at build time (`sha256sum -c MANIFEST.sha256`) | — |

## Rules

1. The reviewed files stay FLAT under `labels/` (like 03c) so `paths.corrected_wizard_exports`
   can point at `labels/`; the feed and 03a files sit in `labels/Batch_NN/` (like `data_matched`).
   Filenames are unchanged on purpose (48 code sites use them).
2. Every `Image_Path` cell in the feed and reviewed workbooks was rewritten to `{NEW_IMAGE_ROOT}/…`.
   The wizard renders images through `file://` + that absolute path, so if this folder ever moves,
   rewrite the cells again (same script, edit `DST`).
3. {len(plan['missing_src'])} referenced images were already missing in `1639_DS` at build time and are
   still missing here (listed in the build log).
4. Human exports are the record of what the annotator clicked. Corrections go into the file in
   `labels/` AND a row in `CORRECTIONS_LOG.csv`; the frozen originals in `1639_DS` are never touched.
5. Re-freeze: when a batch is final, hash it (`sha256sum labels/reviewed_patents_Batch_NN.xlsx`)
   and note the hash in `CORRECTIONS_LOG.csv`. `1639_DS/data/03e_CORRECTED_FROZEN_20260908` predates
   the 17:30 edits of that day; the files here are newer.

Config: `Patent-Labelling-Tools/config.yaml` → `paths.labelled` points here; `matched`,
`data_matched`, `corrected_wizard_exports` and `batches_xlsx` derive from it. The archive keys
(`raw_images`, `triage`, `html_review_exports`, `processed`, …) still point at `1639_DS`.
""")


# ─── verify ──────────────────────────────────────────────────────────────────

def verify() -> bool:
    ok = True
    print(f"Verifying {DST}")
    for xlsx in sorted(DST.glob("labels/reviewed_patents_Batch_*.xlsx")) + sorted(DST.glob("labels/Batch_*/ml_predict_labels_*.xlsx")):
        paths = image_paths(xlsx)
        stale = [p for p in paths if "/1639_DS/" in p]
        absent = {p for p in paths if p.startswith(str(NEW_IMAGE_ROOT)) and not Path(p).exists()}
        # A file that was already absent under every old root is expected to be
        # absent here too — the wizard showed "no image on disk" for it before.
        def _at_source(p):
            rel = p[len(str(NEW_IMAGE_ROOT)) + 1:]
            return any((r / rel).exists() for r in OLD_IMAGE_ROOTS)
        lost = sorted(p for p in absent if _at_source(p))          # real copy failures
        expected = sorted(p for p in absent if not _at_source(p))
        other = [p for p in paths if not p.startswith(str(NEW_IMAGE_ROOT)) and "/1639_DS/" not in p]
        flag = "OK " if not stale and not lost else "BAD"
        if flag == "BAD":
            ok = False
        print(f"  {flag} {xlsx.relative_to(DST)}: {len(paths)} Image_Path cells, "
              f"{len(stale)} still point at 1639_DS, {len(lost)} lost in copy, "
              f"{len(expected)} distinct files already missing at source, "
              f"{len(other)} outside images/ (placeholder)")
        for p in lost[:3]:
            print(f"        LOST: {p}")
    n_img = sum(1 for _ in NEW_IMAGE_ROOT.rglob("*") if _.is_file())
    print(f"  images/ holds {n_img} files")
    for must in ["batches.xlsx", "README.md", "MANIFEST.sha256", "labels/CORRECTIONS_LOG.csv"]:
        if not (DST / must).exists():
            print(f"  BAD missing {must}"); ok = False
    print("VERIFY:", "PASS" if ok else "FAIL")
    return ok


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--apply", action="store_true", help="copy + rewrite + manifest, then verify")
    ap.add_argument("--verify", action="store_true", help="only re-check an existing build")
    a = ap.parse_args()
    if a.verify:
        return 0 if verify() else 1
    plan = build_plan()
    print_plan(plan)
    if not a.apply:
        print("\nDry run only. Re-run with --apply to build.")
        return 0
    print("\nApplying…")
    apply(plan)
    return 0 if verify() else 1


if __name__ == "__main__":
    sys.exit(main())
