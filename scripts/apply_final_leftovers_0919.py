#!/usr/bin/env python3
"""apply_final_leftovers_0919.py — the last two record fixes the user chose on 2026-09-19 (21:20) instead of another
wizard pass ("no more wizard work"):
  (a) US2023056709A1 aircraft 2: boom-card Retraction Kinematics was never recorded → Fixed, as on aircraft 1.
  (b) values stored behind an M2 override the user KEPT (GB202302720D0, US2019135426A1): cleared. The wizard does not
      show them and notebook 04 ignores them (flag OVERRIDE_HIDES_VALUE); rule "an override keeps only what it records".
    python scripts/apply_final_leftovers_0919.py            # dry run
    python scripts/apply_final_leftovers_0919.py --apply    # back up, write the record, re-split the batch files
Values keep their types (openpyxl). After --apply: notebook 04 → canonicalize_record.py --apply → 04.
"""
import sys, shutil, subprocess
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook, Workbook

ROOT = Path("/mnt/storage_11tb/Drive_files_to_syncronize/3 - Images DataSets & Labelling Outputs/1639_LABELLED/0_labelling")
RECORD = ROOT / "inputs" / "record" / "reviewed_patents_Batch_ALL.xlsx"
APPLY = "--apply" in sys.argv[1:]
RMECH_FIXED = "Exposed — Fixed (same configuration in every phase of flight)"
CLEAR = {"GB202302720D0": ["wingConf", "wingConfOtherNote", "empType", "empTypeOtherNote", "wing1_tilt", "wing1_role"],
         "US2019135426A1": ["wingConf", "wingConfOtherNote", "wing1_role"]}          # M2 fields behind the kept override
def s(v): return "" if v is None else str(v)

rows = [list(r) for r in load_workbook(RECORD)["Review"].iter_rows(values_only=True)]
hdr, body = rows[0], rows[1:]
ix = {c: i for i, c in enumerate(hdr)}
where = {(s(r[ix["Patent_ID"]]), s(r[ix["Field"]])): i for i, r in enumerate(body)}
changes, inserts = [], {}
# (b) clear hidden M2 values — only while the M2 override is still on
for pid, fields in CLEAR.items():
    assert s(body[where[(pid, "m2_quickOverride")]][ix["Value"]]) == "True", f"{pid}: M2 override is no longer on"
    for f in fields:
        i = where.get((pid, f))
        if i is not None and s(body[i][ix["Value"]]).strip():
            changes.append((pid, f, s(body[i][ix["Value"]]), "")); body[i][ix["Value"]] = None
# (a) US2023056709A1 aircraft 2 boom rmech
pid = "US2023056709A1_arch2"
i = where.get((pid, "boom_rmech"))
if i is None:
    anchor = where[(pid, "boom_bmech")]                       # the wizard's order: count, chord, orient, bmech, rmech, propKin
    row = list(body[anchor]); row[ix["Field"]] = "boom_rmech"; row[ix["Value"]] = RMECH_FIXED
    inserts[anchor] = row; changes.append((pid, "boom_rmech", "(no row)", RMECH_FIXED))
elif not s(body[i][ix["Value"]]).strip():
    changes.append((pid, "boom_rmech", "", RMECH_FIXED)); body[i][ix["Value"]] = RMECH_FIXED

print(f"{len(changes)} changes:"); [print(f"  {p:<22} {f:<18} {o[:40]!r} -> {n[:40]!r}") for p, f, o, n in changes]
if not APPLY: print("dry run — re-run with --apply"); sys.exit(0)
if not changes: print("nothing to write"); sys.exit(0)
bak = RECORD.parent / "_backups" / f"{RECORD.stem}.PRE_LEFTOVERS_{datetime.now():%Y%m%d_%H%M%S}.xlsx"; shutil.copy2(RECORD, bak)
wb = Workbook(); ws = wb.active; ws.title = "Review"; ws.append(hdr)
for k, r in enumerate(body):
    ws.append(r)
    if k in inserts: ws.append(inserts[k])
wb.save(RECORD); print(f"written {RECORD}; backup {bak}")
subprocess.run([sys.executable, str(Path(__file__).with_name("split_wizard_all_export.py")), str(RECORD), "--apply"], check=True)
