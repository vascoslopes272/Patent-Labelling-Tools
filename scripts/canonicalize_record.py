#!/usr/bin/env python3
"""canonicalize_record.py — put the wizard record (and the wizard feed) in the order notebook 04 writes and write the
aircraft names derived by notebook 04 into it; then regenerate the five per-batch copies.
    python scripts/canonicalize_record.py            # dry run: report what would change
    python scripts/canonicalize_record.py --apply    # back up, write record + feed, re-split the batch files
Run it after every notebook 04 run that flags RECORD_ORDER_STALE / RECORD_NAMES_STALE, then run 04 again (0 names
written, no RECORD_ flag).

ORDER — the single source of truth is outputs/tables/aircraft_table.csv, which notebook 04 (§8) sorts by the ruling of
2026-09-17: ONE block per company (A→Z, the two pseudo-companies last), inside a company the named aircraft A→Z, then
the unnamed by patent id, then the rejected records; a D1/D2 of the same company directly under its original; a
patent's aircraft together in ua order. This script only copies that patent order into the record and the feed
(rows of one patent — its bare id and its _archN ids — stay together, in their existing relative order).
The wizard keys everything by Patent_ID, so any row order loads; but a saved patent's rows move to the END of the
wizard's export (REVIEWED_ROWS filter+concat), so every export needs this script again before it is the record.
NAMES (same rule as gen_04 §11): one string per approved patent, "a; b" for a multi-aircraft patent; a T1/aircraftName
row is inserted after isApproved when none exists. D1/D2 carry the name of the aircraft they point at.
Cells keep their types (openpyxl) — the wizard's xlBool() reads the text "False" as true.
"""
import re, sys, shutil, subprocess
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook, Workbook
ROOT = Path("/mnt/storage_11tb/Drive_files_to_syncronize/3 - Images DataSets & Labelling Outputs/1639_LABELLED/0_labelling")
RECORD = ROOT / "inputs" / "record" / "reviewed_patents_Batch_ALL.xlsx"
FEED = ROOT / "inputs" / "record" / "ml_predict_labels_Batch_ALL.xlsx"
TABLE = ROOT / "outputs" / "tables" / "aircraft_table.csv"
APPLY = "--apply" in sys.argv[1:]
base = lambda s: re.sub(r"_arch\d+$", "", str(s))
def s(v): return "" if v is None else str(v)
M = pd.read_csv(TABLE, keep_default_na=False, dtype=str, low_memory=False); M["ua"] = M.ua.astype(int)
ORDER = M.drop_duplicates("patent_id").patent_id.tolist()      # the order notebook 04 wrote (§8)
COMP = dict(zip(M.patent_id, M.company))
POS = {p: i for i, p in enumerate(ORDER)}

# ── the record ──
src = load_workbook(RECORD)["Review"]
rows = [list(r) for r in src.iter_rows(values_only=True)]
hdr, body = rows[0], rows[1:]
ix = {c: i for i, c in enumerate(hdr)}
pids = list(dict.fromkeys(base(r[ix["Patent_ID"]]) for r in body))
missing = sorted(set(pids) - set(POS)); extra = sorted(set(POS) - set(pids))
assert not missing and not extra, f"record vs aircraft_table patents differ: not in table {missing[:6]}, not in record {extra[:6]}"
order = sorted(pids, key=POS.get)
moved = sum(1 for a, b in zip(pids, order) if a != b)

appr = M[M.is_approved.astype(str) == "True"]
NAME_STR = {pid: "; ".join(g.sort_values("ua").aircraft_name.tolist()) for pid, g in appr.groupby("patent_id")}
assert all(v.strip() for v in NAME_STR.values()), "an approved aircraft has an empty name"
has_name_row = {base(r[ix["Patent_ID"]]) for r in body if s(r[ix["Section"]]) == "T1" and s(r[ix["Field"]]) == "aircraftName"}
blocks = {p: [] for p in pids}
changed, added = [], []
for row in body:
    pid = base(row[ix["Patent_ID"]]); sec, fld = s(row[ix["Section"]]), s(row[ix["Field"]])
    if sec == "T1" and fld == "aircraftName" and pid in NAME_STR and s(row[ix["Value"]]) != NAME_STR[pid]:
        changed.append((pid, row[ix["Value"]], NAME_STR[pid])); row = list(row); row[ix["Value"]] = NAME_STR[pid]
    blocks[pid].append(row)
    if sec == "T1" and fld == "isApproved" and pid in NAME_STR and pid not in has_name_row:
        new = [None] * len(hdr); new[ix["Patent_ID"]], new[ix["Section"]], new[ix["Sub_Dimension"]], new[ix["Field"]], new[ix["Value"]] = pid, "T1", "Aircraft / Prototype Name", "aircraftName", NAME_STR[pid]
        blocks[pid].append(new); added.append(pid)
out = [hdr] + [r for p in order for r in blocks[p]]
print(f"record {RECORD.name}: {len(body)} rows, {len(pids)} patents; patents that move: {moved}; names to change: {len(changed)}; name rows to add: {len(added)}")
for pid, old, new in changed[:40]: print(f"  {pid}: {old!r} -> {new!r}")
if len(changed) > 40: print(f"  ... {len(changed) - 40} more")
for pid in added[:20]: print(f"  + {pid}: {NAME_STR[pid]!r}")
print(f"  first patents in order: {[(p, COMP[p]) for p in order[:4]]}")

# ── the feed (what the wizard's Load Batch reads): same patent order, rows of a patent kept together ──
fws = load_workbook(FEED, read_only=True)
fsheet = fws["Review"] if "Review" in fws.sheetnames else fws.worksheets[0]
fit = fsheet.iter_rows(values_only=True); fhdr = list(next(fit)); frows = [list(r) for r in fit]
fpid = fhdr.index("Patent_ID")
fblocks = {}
for r in frows: fblocks.setdefault(base(r[fpid]), []).append(r)
fpids = list(fblocks)
fextra = sorted(set(fpids) - set(POS)); assert not fextra, f"feed patents not in aircraft_table: {fextra[:8]}"
forder = sorted(fpids, key=POS.get)
fmoved = sum(1 for a, b in zip(fpids, forder) if a != b)
contig = sum(1 for i in range(1, len(frows)) if base(frows[i][fpid]) != base(frows[i-1][fpid])) + 1
print(f"feed {FEED.name} ({fsheet.title}): {len(frows)} rows, {len(fpids)} patents in {contig} blocks; patents that move: {fmoved}")

if not APPLY: print("dry run — re-run with --apply"); sys.exit(0)
if not (moved or changed or added or fmoved or contig != len(fpids)): print("nothing to write"); sys.exit(0)
ts = datetime.now().strftime("%Y%m%d_%H%M%S")
bdir = RECORD.parent / "_backups"; bdir.mkdir(exist_ok=True)
if moved or changed or added:
    bak = bdir / f"{RECORD.stem}.PRE_CANON_{ts}.xlsx"; shutil.copy2(RECORD, bak)
    wb = Workbook(); ws = wb.active; ws.title = "Review"
    for r in out: ws.append(r)
    wb.save(RECORD); print(f"written {RECORD} ({len(out) - 1} rows); backup {bak}")
if fmoved or contig != len(fpids):
    fbak = bdir / f"{FEED.stem}.PRE_CANON_{ts}.xlsx"; shutil.copy2(FEED, fbak)
    wb = Workbook(); ws = wb.active; ws.title = fsheet.title; ws.append(fhdr)
    for p in forder:
        for r in fblocks[p]: ws.append(r)
    wb.save(FEED); print(f"written {FEED} ({len(frows)} rows); backup {fbak}")
subprocess.run([sys.executable, str(Path(__file__).with_name("split_wizard_all_export.py")), str(RECORD), "--apply"], check=True)
