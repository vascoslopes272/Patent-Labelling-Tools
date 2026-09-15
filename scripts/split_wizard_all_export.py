#!/usr/bin/env python3
"""Split a whole-corpus wizard export (reviewed_patents_Batch_ALL.xlsx, made by loading
joined/wizard_all/ml_predict_labels_Batch_ALL.xlsx + resuming reviewed_patents_Batch_ALL.xlsx)
back into the five labels/reviewed_patents_Batch_NN.xlsx files.

    python scripts/split_wizard_all_export.py [export.xlsx]                 # dry run: report only
    python scripts/split_wizard_all_export.py [export.xlsx] --apply         # back up, write, log
    python scripts/split_wizard_all_export.py [export.xlsx] --out DIR       # write the five files to DIR (test)
    python scripts/split_wizard_all_export.py export.xlsx --partial [...]   # SESSION export (subset of patents): replace only those

Default export path: ~/Downloads/reviewed_patents_Batch_ALL.xlsx.
Rules
  * cells are copied with their TYPES (bool / int / float / str) — the wizard's xlBool() reads the text
    "False" as true, so a text-only rewrite would corrupt every unticked box
  * a patent goes to the batch it already lives in (joined/wizard_all/patent_batch_map.csv);
    a patent absent from the map is an error, never silently placed
  * inside each batch the patents keep the order of the CURRENT labels file (new ones appended),
    so notebook 04's file-order sort is stable
  * --apply copies the current five files to labels/_backups/PRE_WIZARD_ALL_<ts>/ first and appends
    one CORRECTIONS_LOG.csv row per changed (patent, section, sub-dimension, field) cell
"""
import re, sys, csv, shutil
from collections import OrderedDict
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook, Workbook

ROOT = Path("/mnt/storage_11tb/Drive_files_to_syncronize/3 - Images DataSets & Labelling Outputs/1639_LABELLED")
LAB = ROOT / "labels"
MAP = ROOT / "joined" / "wizard_all" / "patent_batch_map.csv"
BATCHES = ["Batch_01", "Batch_02", "Batch_03", "Batch_04", "Batch_05"]

argv = sys.argv[1:]
APPLY = "--apply" in argv
PARTIAL = "--partial" in argv     # the export holds only some patents (a SESSION batch): replace those, keep the rest
OUT = Path(argv[argv.index("--out") + 1]) if "--out" in argv else None
args = [a for i, a in enumerate(argv) if not a.startswith("--") and (i == 0 or argv[i - 1] != "--out")]
src = Path(args[0]) if args else Path.home() / "Downloads" / "reviewed_patents_Batch_ALL.xlsx"
base = lambda s: re.sub(r"_arch\d+$", "", str(s))
def s(v): return "" if v is None else str(v)

def read(path):
    ws = load_workbook(path, read_only=True)["Review"]
    it = ws.iter_rows(values_only=True); hdr = list(next(it))
    return hdr, [tuple(r) for r in it]

hdr, rows = read(src)
need = ["Patent_ID", "Section", "Sub_Dimension", "Field", "Value"]
assert all(c in hdr for c in need), f"export lacks {need}: {hdr}"
ix = {c: hdr.index(c) for c in hdr}
key = lambda r: (s(r[ix["Patent_ID"]]), s(r[ix["Section"]]), s(r[ix["Sub_Dimension"]]), s(r[ix["Field"]]))
with open(MAP) as f: bmap = {r["patent_id"]: r["batch"] for r in csv.DictReader(f)}
pats = OrderedDict((base(r[ix["Patent_ID"]]), None) for r in rows)
unknown = [p for p in pats if p not in bmap]
if unknown: sys.exit(f"!! {len(unknown)} patents in the export are not in the batch map: {unknown[:10]}")
by_batch = {b: [] for b in BATCHES}
for r in rows: by_batch[bmap[base(r[ix["Patent_ID"]])]].append(r)
print(f"export {src.name}: {len(rows)} rows, {len(pats)} patents -> " + str({b: len({base(r[ix['Patent_ID']]) for r in v}) for b, v in by_batch.items()}))

ts = datetime.now().strftime("%Y%m%d_%H%M%S")
log_rows, out = [], {}
for b in BATCHES:
    chdr, cur = read(LAB / f"reviewed_patents_{b}.xlsx")
    cix = {c: chdr.index(c) for c in chdr}
    order = {p: i for i, p in enumerate(OrderedDict((base(r[cix["Patent_ID"]]), None) for r in cur))}
    new = sorted(by_batch[b], key=lambda r: order.get(base(r[ix["Patent_ID"]]), len(order)))   # stable
    new = [tuple(r[ix[c]] if c in ix else None for c in chdr) for r in new]                     # current column order
    nix = cix
    ckey = lambda r: (s(r[nix["Patent_ID"]]), s(r[nix["Section"]]), s(r[nix["Sub_Dimension"]]), s(r[nix["Field"]]))
    a = {ckey(r): r[nix["Value"]] for r in cur}; n = {ckey(r): r[nix["Value"]] for r in new}
    changed = added = removed = 0
    for k in list(dict.fromkeys(list(n) + list(a))):
        old, val = a.get(k, "<absent>"), n.get(k, "<absent>")
        if old == val and type(old) is type(val): continue
        kind = "changed" if "<absent>" not in (old, val) else ("added" if old == "<absent>" else "removed")
        changed += kind == "changed"; added += kind == "added"; removed += kind == "removed"
        log_rows.append(dict(applied_utc=datetime.now(timezone.utc).isoformat(timespec="seconds"), batch=b, patent_id=k[0], section=k[1],
                             field=k[3] if k[2] in ("", "nan") else f"{k[3]} [{k[2]}]",
                             old_value="(absent)" if old == "<absent>" else s(old), new_value="(row removed)" if val == "<absent>" else s(val),
                             rule="wizard relabel session 2026-09-15 (whole-corpus load, split by scripts/split_wizard_all_export.py)",
                             evidence=f"export {src.name}", confirmed_by="annotator"))
    lost = sorted({base(r[cix["Patent_ID"]]) for r in cur} - {base(r[nix["Patent_ID"]]) for r in new})
    if lost and PARTIAL:
        # a session export carries only some patents: keep every other patent's current rows, in place
        inx = {base(r[nix["Patent_ID"]]) for r in new}
        merged, seen = [], set()
        for r in cur:
            p = base(r[cix["Patent_ID"]])
            if p in inx:
                if p not in seen: merged += [x for x in new if base(x[nix["Patent_ID"]]) == p]; seen.add(p)
            else: merged.append(r)
        merged += [x for x in new if base(x[nix["Patent_ID"]]) not in {base(r[cix["Patent_ID"]]) for r in cur}]
        new = merged; lost = []
        # the diff counts above were computed against the full file: recount on the merged result
        n = {ckey(r): r[nix["Value"]] for r in new}
        log_rows = [l for l in log_rows if l["batch"] != b]
        changed = added = removed = 0
        for k in list(dict.fromkeys(list(n) + list(a))):
            old, val = a.get(k, "<absent>"), n.get(k, "<absent>")
            if old == val and type(old) is type(val): continue
            kind = "changed" if "<absent>" not in (old, val) else ("added" if old == "<absent>" else "removed")
            changed += kind == "changed"; added += kind == "added"; removed += kind == "removed"
            log_rows.append(dict(applied_utc=datetime.now(timezone.utc).isoformat(timespec="seconds"), batch=b, patent_id=k[0], section=k[1],
                                 field=k[3] if k[2] in ("", "nan") else f"{k[3]} [{k[2]}]",
                                 old_value="(absent)" if old == "<absent>" else s(old), new_value="(row removed)" if val == "<absent>" else s(val),
                                 rule="wizard relabel session 2026-09-15 (session batch, split by scripts/split_wizard_all_export.py --partial)",
                                 evidence=f"export {src.name}", confirmed_by="annotator"))
    print(f"  {b}: {len(cur)} -> {len(new)} rows, {len({base(r[nix['Patent_ID']]) for r in new})} patents; cells changed {changed}, added {added}, removed {removed}"
          + (f"  !! {len(lost)} patents MISSING from the export: {lost[:6]}" if lost else ""))
    if lost: sys.exit("refusing: a batch would lose patents (use --partial for a session export)")
    out[b] = (chdr, new)

def write(path, chdr, new):
    wb = Workbook(); ws = wb.active; ws.title = "Review"; ws.append(chdr)
    for r in new: ws.append(list(r))
    wb.save(path)

if OUT:
    OUT.mkdir(parents=True, exist_ok=True)
    for b in BATCHES: write(OUT / f"reviewed_patents_{b}.xlsx", *out[b])
    print(f"test output: 5 files in {OUT} ({len(log_rows)} diff rows, not logged)"); sys.exit(0)
if not APPLY:
    print(f"dry run — {len(log_rows)} log rows would be written; re-run with --apply"); sys.exit(0)
bdir = LAB / "_backups" / f"PRE_WIZARD_ALL_{ts}"; bdir.mkdir(parents=True, exist_ok=True)
for b in BATCHES:
    shutil.copy2(LAB / f"reviewed_patents_{b}.xlsx", bdir / f"reviewed_patents_{b}.xlsx")
    write(LAB / f"reviewed_patents_{b}.xlsx", *out[b])
if log_rows:
    with open(LAB / "CORRECTIONS_LOG.csv", "a", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(log_rows[0])); w.writerows(log_rows)
print(f"applied: 5 files written, backups in {bdir}, {len(log_rows)} CORRECTIONS_LOG rows appended")
